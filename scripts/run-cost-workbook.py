"""The Run-Cost Model workbook, half two: the sheet itself.

Reads the JSON that scripts/run-cost-workbook-data.ts prints and writes the
Excel file. Three sheets:

  Read me             what the tool is for, how to fill it, the colour key,
                      the rules, the rubric, the limits, the credit line
  Your model          every line blank, formulas in place, an instruction
                      beside every line
  Reference example   the same sheet with the ordering agent filled in

The formulas are the page's formulas, row for row. Percent lines take a
percentage (12 means 12%), the same as the page. One difference the Read me
sheet says out loud: on the page a blank line is unknown and keeps the result
off the screen; in Excel a blank cell is 0, so a half-filled sheet shows a
number that is not a result.

Branding follows .claude/skills/the-living-craft-design: one black band per
sheet, sun for every cell you act on, mist for what is worked out, ember for
the cohort panel with ink text on it, no borders where a fill will do.

    npm run workbook

Needs openpyxl (pip install openpyxl).
"""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NOIR = "000000"
SUN = "FFC123"
EMBER = "FD8549"
MIST = "F1F3F5"
INK = "16212E"
QUIET = "5F6B78"
WHITE = "FFFFFF"

BODY = "Helvetica Neue"

# The formulas, one per worked-out line. `x` reads an option's cell in the
# current column, `D` a shared cell in column D. Each is a function so a line
# only references rows that already exist above it.
ARM_FORMULAS = {
    "buildCost": lambda x, D: f"=({x('buildDays')}+{x('evalDays')})*{D('engDayRate')}",
    "modelCost": lambda x, D: f"={D('casesPerMonth')}*{x('modelCalls')}*{x('retryMult')}*{x('modelCallCost')}",
    "toolCost": lambda x, D: f"={D('casesPerMonth')}*{x('toolCalls')}*{x('toolCallCost')}",
    "reviewCost": lambda x, D: f"={D('casesPerMonth')}*{x('reviewShare')}/100*{x('reviewMinutes')}/60*{D('reviewerRate')}",
    "escalationCost": lambda x, D: f"={D('casesPerMonth')}*{x('escalationShare')}/100*{x('escalationMinutes')}/60*{D('engHourRate')}",
    "declinedCost": lambda x, D: f"={D('casesPerMonth')}*{x('declinedShare')}/100*{x('declinedMinutes')}/60*{D('reviewerRate')}",
    "upkeepCost": lambda x, D: f"={x('toolingCost')}+({x('evalUpkeepDays')}+{x('requalDays')}/12+{x('regressionDays')}+{x('incidentDays')})*{D('engDayRate')}",
    "outcomes": lambda x, D: f"={D('totalCases')}*{x('acceptRate')}/100",
}
TOTAL_FORMULAS = {
    "runPerMonth": lambda x, D: f"={x('modelCost')}+{x('toolCost')}+{x('reviewCost')}+{x('escalationCost')}+{x('declinedCost')}+{x('upkeepCost')}",
    "runPeriod": lambda x, D: f"={x('runPerMonth')}*{D('months')}",
    "total": lambda x, D: f"={x('buildCost')}+{x('runPeriod')}",
    "perCase": lambda x, D: f"=IFERROR({x('total')}/{D('totalCases')},\"\")",
    "perOutcome": lambda x, D: f"=IFERROR({x('total')}/{x('outcomes')},\"n/a\")",
    "net": lambda x, D: f"={D('manualBaseline')}-{x('total')}",
    "breakEven": lambda x, D: f"=IF(({D('manualBaseline')}/{D('months')}-{x('runPerMonth')})<=0,\"never\",{x('buildCost')}/({D('manualBaseline')}/{D('months')}-{x('runPerMonth')}))",
}

fill = lambda hex_: PatternFill("solid", start_color=hex_, end_color=hex_)
font = lambda **kw: Font(name=BODY, color=kw.pop("color", INK), **kw)


def band(ws, row, title, subtitle, credit, cols=7):
    """The one black area on a sheet: wordmark, title, credit."""
    for r in (row, row + 1, row + 2):
        for c in range(1, cols + 1):
            ws.cell(r, c).fill = fill(NOIR)
    ws.cell(row, 1, "● The Living Craft").font = font(bold=True, size=11, color=SUN)
    ws.cell(row + 1, 1, title).font = font(bold=True, size=18, color=WHITE)
    ws.cell(row + 2, 1, f"{subtitle}  ·  {credit}").font = font(size=9, color="A3ACB8")
    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row + 1].height = 28
    ws.row_dimensions[row + 2].height = 18
    return row + 4


def wrapped(cell, **kw):
    cell.alignment = Alignment(wrap_text=True, vertical="top", **kw)
    return cell


def build_model_sheet(wb, data, title, values, story=None):
    """One model sheet. `values` is the inputs object or None for a blank sheet."""
    ws = wb.create_sheet(title)
    arms = data["arms"]
    widths = [44, 20, 62, 16, 16, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = band(ws, 1, data["toolName"], "Four ways to do the same job, costed over one period", data["credit"])

    if story:
        ws.cell(r, 1, story["title"]).font = font(bold=True, size=12)
        r += 1
        for line in story["lines"]:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
            wrapped(ws.cell(r, 1, line)).font = font(size=10, color=QUIET)
            ws.row_dimensions[r].height = 30
            r += 1
        r += 1

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    wrapped(ws.cell(r, 1, "Fill the yellow cells. Column C says what to type and where the number comes from. Grey rows are worked out for you; do not type into them. Percent lines take a percentage: 12 means 12%.")).font = font(size=10)
    ws.row_dimensions[r].height = 30
    r += 2

    # Column headers.
    heads = ["Line", "Unit", "What to type", *[a["short"] for a in arms]]
    for c, h in enumerate(heads, start=1):
        cell = ws.cell(r, c, h)
        cell.font = font(bold=True, size=10, color=QUIET)
        cell.alignment = Alignment(horizontal="right" if c > 3 else "left")
    r += 1
    for c, a in enumerate(arms, start=4):
        wrapped(ws.cell(r, c, a["legend"]), horizontal="right").font = font(size=8, color=QUIET)
    ws.row_dimensions[r].height = 26
    header_row = r - 1
    r += 1

    rows = {}  # key -> row number, shared keys and arm keys alike
    arm_cols = {a["key"]: get_column_letter(4 + i) for i, a in enumerate(arms)}

    def section_head(sec):
        nonlocal r
        for c in range(1, 8):
            ws.cell(r, c).fill = fill(MIST)
        ws.cell(r, 1, f"{sec['num']}  {sec['name']}").font = font(bold=True, size=11)
        ws.cell(r, 3, sec["question"]).font = font(size=10, color=QUIET)
        ws.row_dimensions[r].height = 20
        r += 1

    def input_row(sec, row):
        nonlocal r
        label = ("• " if row["forgotten"] else "") + row["label"]
        wrapped(ws.cell(r, 1, label)).font = font(size=10, bold=row["forgotten"])
        wrapped(ws.cell(r, 2, row["unit"])).font = font(size=9, color=QUIET)
        wrapped(ws.cell(r, 3, row["why"])).font = font(size=9, color=QUIET)
        ws.row_dimensions[r].height = 42
        if sec["scope"] == "shared":
            cell = ws.cell(r, 4)
            cell.fill = fill(SUN)
            cell.font = font(size=10, bold=True)
            cell.alignment = Alignment(horizontal="right")
            if values:
                cell.value = values["shared"][row["key"]]
        else:
            for a in arms:
                cell = ws.cell(r, 4 + arms.index(a))
                cell.fill = fill(SUN)
                cell.font = font(size=10, bold=True)
                cell.alignment = Alignment(horizontal="right")
                if values:
                    cell.value = values["arms"][a["key"]][row["key"]]
        rows[row["key"]] = r
        r += 1

    def computed_row(label, why, formulas, fmt="#,##0.00", key=None):
        """A worked-out line: mist fill, a formula per option (or one for shared)."""
        nonlocal r
        for c in range(1, 8):
            ws.cell(r, c).fill = fill(MIST)
        wrapped(ws.cell(r, 1, label)).font = font(size=10, bold=True, color=QUIET)
        ws.cell(r, 2, "worked out").font = font(size=9, color=QUIET)
        wrapped(ws.cell(r, 3, why)).font = font(size=9, color=QUIET)
        ws.row_dimensions[r].height = 30
        for col, f in formulas.items():
            cell = ws.cell(r, col, f)
            cell.font = font(size=10, bold=True)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="right")
        if key:
            rows[key] = r
        r += 1

    D = lambda key: f"$D${rows[key]}"  # a shared cell, absolute
    X = lambda col, key: f"{col}{rows[key]}"  # an option's cell

    # Currency label first, in section 1.
    for sec in data["sections"]:
        section_head(sec)
        if sec["num"] == 1:
            wrapped(ws.cell(r, 1, "Currency label")).font = font(size=10)
            ws.cell(r, 2, "free text").font = font(size=9, color=QUIET)
            wrapped(ws.cell(r, 3, "Such as INR or USD. A label only; the formulas do not read it.")).font = font(size=9, color=QUIET)
            cell = ws.cell(r, 4, values["currency"] if values else None)
            cell.fill = fill(SUN)
            cell.font = font(size=10, bold=True)
            ws.row_dimensions[r].height = 30
            r += 1
        for row in sec["inputs"]:
            input_row(sec, row)
        for row in sec["computed"]:
            k = row["key"]
            if sec["scope"] == "shared":
                if k == "totalCases":
                    f = f"={D('casesPerMonth')}*{D('months')}"
                else:
                    f = f"={D('totalCases')}*{D('manualMinutes')}/60*{D('manualRate')}"
                computed_row(row["label"], row["why"], {4: f}, "#,##0", k)
            else:
                formulas = {}
                for a in arms:
                    col = arm_cols[a["key"]]
                    x = lambda key, col=col: X(col, key)
                    formulas[4 + arms.index(a)] = ARM_FORMULAS[k](x, D)
                computed_row(row["label"], row["why"], formulas, "#,##0", k)
        r += 1

    # Totals.
    for c in range(1, 8):
        ws.cell(r, c).fill = fill(MIST)
    ws.cell(r, 1, "Totals").font = font(bold=True, size=11)
    ws.cell(r, 3, "Every option on the same cases. The number to argue about is cost per acceptable outcome.").font = font(size=10, color=QUIET)
    r += 1
    for t in data["totals"]:
        k = t["key"]
        formulas = {}
        for a in arms:
            col = arm_cols[a["key"]]
            x = lambda key, col=col: X(col, key)
            formulas[4 + arms.index(a)] = TOTAL_FORMULAS[k](x, D)
        fmt = "#,##0.00" if k in ("perCase", "perOutcome") else "0.0" if k == "breakEven" else "#,##0"
        computed_row(t["label"], t["why"], formulas, fmt, k)

    # The rubric, as two rows: the candidates that save money, and the leader.
    cand = {}
    for a in arms:
        col = arm_cols[a["key"]]
        cand[4 + arms.index(a)] = f"=IF({X(col, 'net')}>0,{X(col, 'perOutcome')},\"\")"
    computed_row("Cost per acceptable outcome, options that save money", "Blank for an option that costs more over the period than doing it by hand.", cand, "#,##0.00", "candidates")
    cr = rows["candidates"]
    lead = f"=IFERROR(INDEX($D${header_row}:$G${header_row},MATCH(MIN(D{cr}:G{cr}),D{cr}:G{cr},0)),\"Nothing here beats doing it by hand\")"
    computed_row("Leading option", "The option with the lowest cost per acceptable outcome among those that save money. The rubric on the Read me sheet says what to do with each answer.", {4: lead}, "@", "leader")
    ws.cell(rows["leader"], 4).alignment = Alignment(horizontal="left")

    ws.freeze_panes = ws.cell(header_row + 2, 4)
    return ws


def build_readme(wb, data):
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 110
    r = band(ws, 1, data["toolName"], "Four ways to do the same job, costed over one period", data["credit"], cols=1)

    def h(text):
        nonlocal r
        ws.cell(r, 1, text).font = font(bold=True, size=12)
        r += 1

    def p(text, color=INK, size=10):
        nonlocal r
        wrapped(ws.cell(r, 1, text)).font = font(size=size, color=color)
        ws.row_dimensions[r].height = max(15, 15 * (1 + len(text) // 105))
        r += 1

    h("What this is")
    for line in data["purpose"]:
        p(line)
    r += 1

    h("How to fill it in")
    for line in [
        "1. Open the sheet called Your model. Fill the yellow cells. Column C says what to type and where the number comes from.",
        "2. Sections 1 and 2 take one value each. Sections 3 to 7 take one value per option: Rules, Rules v2, Assisted, Agent.",
        "3. Percent lines take a percentage. Type 12 for 12%, not 0.12.",
        "4. Grey rows are worked out for you. Do not type into them.",
        "5. Excel treats a blank cell as 0, so a half-filled sheet shows a number that is not a result. The page at the link below keeps the result off the screen until every line is set; here, fill every yellow cell before you read the totals.",
        "6. The Reference example sheet is the same model with every cell filled: an ordering agent across 40 sites in India, in rupees at Indian rates. Copy it, or use it to see what a finished model looks like.",
        f"7. The same model runs in the browser, with the result read against the rubric and a PDF you can build: {data['pageUrl']}",
    ]:
        p(line)
    r += 1

    h("Colour key")
    c = ws.cell(r, 1, "Yellow: a cell you fill in.")
    c.fill = fill(SUN); c.font = font(size=10, bold=True); r += 1
    c = ws.cell(r, 1, "Grey: a line the sheet works out. Do not type into it.")
    c.fill = fill(MIST); c.font = font(size=10, bold=True, color=QUIET); r += 1
    p("• before a line: one of the nine lines most business cases leave out. Fill these from a trial log, not from a guess.")
    r += 1

    h("The four options")
    for a in data["arms"]:
        p(f"{a['name']}: {a['what']}")
    p(data["armsNote"], QUIET)
    r += 1

    h("Rules for using it in a budget conversation")
    for rule in data["rules"]:
        p(f"{rule['lead']} {rule['rest']}")
    r += 1

    h("How the outcome is read")
    p("The leading option is the one with the lowest cost per acceptable outcome among the options that save money against the manual baseline over the period. The last row of each model sheet names it.", QUIET)
    for o in data["outcomes"]:
        p(f"{o['name']} — {o['when']}. {o['what']}")
    r += 1

    h("The nine lines teams leave out")
    p(" · ".join(x.lower() for x in data["forgotten"]))
    r += 1

    h("What this model does not price")
    for l in data["limits"]:
        p(l)
    r += 1

    # The cohort panel: ember with ink text, the one promotional surface.
    top = r
    ws.cell(r, 1, data["cta"]["heading"]).font = font(bold=True, size=12); r += 1
    for line in data["cta"]["lines"]:
        p(line)
    c = ws.cell(r, 1, f"Apply at {data['applyUrl']}")
    c.hyperlink = data["applyUrl"]
    c.font = font(bold=True, size=10, underline="single"); r += 1
    for rr in range(top, r):
        ws.cell(rr, 1).fill = fill(EMBER)
    r += 1
    p(f"{data['toolName']} · The Living Craft · {data['credit']} · free to use and to pass on", QUIET, 9)
    return ws


def main():
    data = json.load(sys.stdin)
    out = sys.argv[1]
    wb = Workbook()
    build_readme(wb, data)
    build_model_sheet(wb, data, "Your model", None)
    build_model_sheet(wb, data, "Reference example", data["example"], data["exampleStory"])
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
