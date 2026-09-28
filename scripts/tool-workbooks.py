"""The three tool workbooks, half two: the sheets themselves.

Reads the JSON that scripts/tool-workbooks-data.ts prints and writes three
blank, reusable workbooks into the folder named on the command line:

  poc-selection-tool.xlsx       the POC Selection Tool
  model-selection-tool.xlsx     the Model Selection Tool
  agent-authority-review.xlsx   the Agent Authority Review

Each opens on the tool itself, blank, with the questions, the anchors and the
rubric from the page's own data module. The shaded cells are the answers;
everything else is worked out by formulas that follow the page's scoring
function rule for rule, and a Read me sheet says how. Each also carries the
page's worked example on its own sheet.

THE FORMULAS USE A SMALL, FIXED SET OF FUNCTIONS on purpose: IF, AND, OR, SUM,
COUNT, COUNTA, COUNTIF, COUNTIFS, CHOOSE, ROUND, LEN, TRIM and the & operator,
and every reference stays on its own sheet. That is what lets
src/lib/resources/tool-workbooks.test.ts read the formulas back out of the
committed files, evaluate them, and prove they agree with the page. A formula
that needs another function needs the test's evaluator to learn it first; the
test fails on any function it does not know.

EVERY INPUT AND RESULT CELL HAS A NAME (a sheet-scoped defined name such as
q01 or out_outcome), so the test never depends on a row number.

The look is the site's, from scripts/workbook_brand.py: the colours are
design system v1 tokens read from src/styles/ds/theme.css, and the lockup is
the PNG pair `npm run build:pdf-assets` makes. The faces are Figtree and
Source Serif 4; that module says what a reader without them sees.

    npm run tool-downloads

Needs openpyxl and Pillow (pip install openpyxl pillow).
"""

from __future__ import annotations

import io
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image as PilImage

REPO = Path(__file__).resolve().parents[1]

# ── the brand: one module, shared with run-cost-workbook.py ───────────────────
from workbook_brand import *  # noqa: E402,F401,F403


def q(s: str) -> str:
    """A string literal inside a formula."""
    return '"' + str(s).replace('"', '""') + '"'


def ground(ws, cols: int, rows: int):
    """The ivory ground under the whole sheet, as on the site's pages."""
    ws.sheet_properties.tabColor = FOREST
    ws.sheet_view.showGridLines = False
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            ws.cell(r, c).fill = fill(IVORY)
            ws.cell(r, c).font = font()


def put(ws, row, col, value=None, *, f=None, bg=None, align=WRAP, border=None, fmt=None, editable=False):
    c = ws.cell(row, col, value)
    c.font = f or font()
    c.alignment = align
    if bg:
        c.fill = fill(bg)
    if border:
        c.border = border
    if fmt:
        c.number_format = fmt
    c.protection = Protection(locked=not editable)
    return c


def inp(ws, row, col, value=None):
    """A cell the reader fills in: shaded, unlocked."""
    return put(ws, row, col, value, f=font(12, True), bg=INPUT, align=WRAP_C, border=Border(top=HAIR, bottom=HAIR, left=HAIR, right=HAIR), editable=True)


def name_cell(wb, ws, label: str, row: int, col: int):
    """A sheet-scoped defined name, so the test finds the cell by what it is."""
    from openpyxl.utils import get_column_letter
    ref = f"'{ws.title}'!${get_column_letter(col)}${row}"
    ws.defined_names[label] = DefinedName(label, attr_text=ref)


def choice_list(ws, cells: list[str], options: list, prompt: str):
    """A dropdown. Options may not contain a comma: the list is written inline."""
    text = [str(o) for o in options]
    for t in text:
        if "," in t:
            raise SystemExit(f"a list option contains a comma: {t!r}")
    dv = DataValidation(type="list", formula1=q(",".join(text)), allow_blank=True, showDropDown=False)
    dv.promptTitle = "Choose one"
    dv.prompt = prompt[:250]
    dv.showInputMessage = True
    dv.error = "Choose one of the options in the list, or leave the cell blank."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    for ref in cells:
        dv.add(ref)


def masthead(ws, title: str, lines: list[str], width_cols: int) -> int:
    """Lockup, title, and a line or two under it. Returns the next free row."""
    img = XlImage(lockup_png())
    img.anchor = "A1"
    ws.add_image(img)
    ws.row_dimensions[1].height = 34
    put(ws, 3, 1, title, f=font(24, color=FOREST, name=DISPLAY), align=Alignment(vertical="bottom"))
    ws.row_dimensions[3].height = 36
    for c in range(1, width_cols + 1):
        ws.cell(3, c).border = GOLD_RULE
    r = 4
    for line in lines:
        put(ws, r, 1, line, f=font(10, color=MUTED), align=Alignment(vertical="top"))
        r += 1
    return r + 1


def protect(ws):
    # A guide rail, not security: no password, so a reader can unprotect and change anything.
    print_setup(ws)
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.formatCells = False


def label_row(ws, row: int, text: str, cols: int, *, bg=PAPER):
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = fill(bg)
    put(ws, row, 1, text, f=font(10, True, DEEP_GOLD), bg=bg, align=Alignment(vertical="center"))
    ws.row_dimensions[row].height = 22


def read_me(wb, title: str, credit: str, url: str, blocks: list[tuple[str, list[str]]]):
    ws = wb.create_sheet("Read me")
    ground(ws, 6, 200)
    ws.column_dimensions["A"].width = 110
    r = masthead(ws, title, [f"{credit}  ·  {url}"], 1)
    for heading, paras in blocks:
        put(ws, r, 1, heading, f=font(12, True, FOREST))
        r += 1
        for p in paras:
            # A paragraph is a string, or a {lead, rest} pair as HOW_TO_RUN writes them.
            text = f"{p['lead']} {p['rest']}" if isinstance(p, dict) else str(p)
            put(ws, r, 1, text, f=font(11), align=WRAP)
            r += 1
        r += 1
    protect(ws)
    return ws


RESULT_COL = 3  # the result values: column C, merged through D


def result_block(ws, row: int, cols: int, items: list[tuple[str, str, str | None]]) -> dict:
    """Label (column B) | formula (C, merged through D) rows on the soft green panel.

    Returns the row of each named result. A long result ("What to do") gets a
    fixed height, because a spreadsheet never auto-fits a merged cell."""
    at = {}
    for i, (label, formula, name) in enumerate(items):
        r = row + i
        for c in range(1, cols + 1):
            ws.cell(r, c).fill = fill(SOFT)
        put(ws, r, 2, label, f=font(10, True, MUTED), bg=SOFT, align=Alignment(horizontal="right", vertical="top", wrap_text=True))
        big = name == "out_outcome"
        put(ws, r, RESULT_COL, formula, f=font(14 if big else 12, True, FOREST if big else INK), bg=SOFT, align=Alignment(horizontal="left", wrap_text=True, vertical="top"))
        ws.merge_cells(start_row=r, start_column=RESULT_COL, end_row=r, end_column=RESULT_COL + 1)
        if label == "What to do":
            ws.row_dimensions[r].height = 66
        if big:
            ws.row_dimensions[r].height = 24
        if name:
            at[name] = r
    return at


# ── the POC Selection Tool ────────────────────────────────────────────────────


def poc_sheet(wb, d: dict, title: str, values: list | None):
    ws = wb.create_sheet(title)
    cols = 6
    ground(ws, cols, 80)
    for col, w in zip("ABCDEF", (6, 46, 30, 30, 30, 9)):
        ws.column_dimensions[col].width = w

    top = masthead(ws, d["toolName"], [f"{d['credit']}  ·  {d['pageUrl']}", "Enter 0, 1 or 2 in each shaded cell. Leave a cell blank if you cannot answer yet."], cols)

    # Lay the questions out first, so the result formulas know their cells.
    start = top + 6
    r = start
    score_cells: list[str] = []
    gate_cells: list[str] = []
    sec_rows = []
    n = 0
    for s in d["sections"]:
        label_row(ws, r, f"{s['letter']}   {s['name'].upper()}" + ("   ·   HARD GATE: ANY 0 HERE STOPS IT" if s["gate"] else ""), cols)
        sec_row = r
        r += 1
        put(ws, r, 2, s["question"], f=font(11, italic=True, color=MUTED))
        for c, h in zip((3, 4, 5, 6), ("0 means", "1 means", "2 means", "Score")):
            put(ws, r, c, h, f=font(9, True, DEEP_GOLD), align=WRAP_C if c == 6 else WRAP)
        r += 1
        first = r
        for it in s["items"]:
            n += 1
            put(ws, r, 1, f"{n:02d}", f=font(10, True, MUTED), border=UNDER)
            put(ws, r, 2, it["q"], f=font(11, True), border=UNDER)
            for k in range(3):
                put(ws, r, 3 + k, it["a"][k], f=font(10, color=MUTED), border=UNDER)
            inp(ws, r, 6, values[n - 1] if values else None)
            name_cell(wb, ws, f"q{n:02d}", r, 6)
            score_cells.append(f"F{r}")
            if s["gate"]:
                gate_cells.append(f"F{r}")
            r += 1
        sec_rows.append((sec_row, first, r - 1, len(s["items"]) * 2))
        r += 1
    choice_list(ws, score_cells, [0, 1, 2], "0, 1 or 2, from the anchors in the row.")

    rng = f"{score_cells[0]}:{score_cells[-1]}"
    gate_rng = f"{gate_cells[0]}:{gate_cells[-1]}"
    count = len(score_cells)
    gate_line, *lines = d["cutLines"]

    def chain(field: str) -> str:
        expr = '""'
        for c in reversed(lines):
            expr = f"IF(SUM({rng})>={c['min']},{q(c[field])},{expr})"
        return (
            f"=IF(COUNTIF({gate_rng},0)>0,{q(gate_line[field])},"
            f"IF(COUNT({rng})<{count},{q('Incomplete' if field == 'name' else 'Answer every question to read the total against the rubric. A blank counts as 0 in the total.')},{expr}))"
        )

    at = result_block(
        ws,
        top,
        cols,
        [
            ("Total", f"=SUM({rng})&\" / {d['maxScore']}\"", None),
            ("Answered", f"=COUNT({rng})&\" of {count}\"", None),
            ("Outcome", chain("name"), "out_outcome"),
            ("What to do", chain("what"), None),
            ("", "", None),
        ],
    )
    # Plain numbers, named, for the test and for anyone building on the sheet.
    put(ws, top, 5, "Total, as a number", f=font(9, color=MUTED), bg=SOFT)
    put(ws, top, 6, f"=SUM({rng})", f=font(12, True), bg=SOFT, align=WRAP_C)
    name_cell(wb, ws, "out_total", top, 6)
    put(ws, top + 1, 5, "Answered, as a number", f=font(9, color=MUTED), bg=SOFT)
    put(ws, top + 1, 6, f"=COUNT({rng})", f=font(12, True), bg=SOFT, align=WRAP_C)
    name_cell(wb, ws, "out_answered", top + 1, 6)
    name_cell(wb, ws, "out_outcome", at["out_outcome"], RESULT_COL)

    # Section subtotals on each section's header row.
    for sec_row, first, last, mx in sec_rows:
        put(ws, sec_row, 6, f'=SUM(F{first}:F{last})&" / {mx}"', f=font(10, True, DEEP_GOLD), bg=PAPER, align=WRAP_C)

    protect(ws)
    return ws


def build_poc(d: dict, out: Path):
    wb = Workbook()
    wb.remove(wb.active)
    poc_sheet(wb, d, "Score", None)
    read_me(
        wb,
        d["toolName"],
        d["credit"],
        d["pageUrl"],
        [
            ("What this tool is for", d["purpose"]),
            ("How to run it", d["howToRun"]),
            (
                "How the sheet works",
                [
                    "Enter 0, 1 or 2 in each shaded cell on the Score sheet. Each question says what earns a 0, a 1 and a 2, so two people scoring the same idea land on the same number.",
                    "The total, the section scores and the outcome are worked out for you, with the same rules as the page. A blank counts as 0 in the total, and there is no outcome until every question is answered.",
                    "Any 0 in the hard-gate section stops the proof of concept, whatever the total.",
                    "The sheet is protected so a formula is not overwritten by accident. There is no password: Review, then Unprotect Sheet, if you want to change it.",
                    "The Example sheet is the page's worked example, already scored.",
                ],
            ),
            ("The rubric", [f"{c['score']}: {c['name']}. {c['what']}" for c in d["cutLines"]]),
            ("Four ways to answer most questions before you build the agent", [f"{m['num']}  {m['name']}. {m['body']} Cost: {m['cost']}." for m in d["moves"]]),
            ("Where this comes from", [f"{d['credit']}. The same tool, with the scoring worked out as you go: {d['pageUrl']}"]),
        ],
    )
    poc_sheet(wb, d, "Example", d["example"])
    wb.active = 0
    wb.save(out)


# ── the Model Selection Tool ──────────────────────────────────────────────────


def msel_sheet(wb, d: dict, title: str, a: dict | None):
    ws = wb.create_sheet(title)
    cols = 9
    ground(ws, cols, 90)
    for col, w in zip("ABCDEFGHI", (6, 44, 26, 26, 26, 9, 9, 10, 10)):
        ws.column_dimensions[col].width = w

    top = masthead(
        ws,
        d["toolName"],
        [f"{d['credit']}  ·  {d['pageUrl']}", "Choose the job first, then answer every shaded cell. Leave a cell blank if you cannot answer yet."],
        cols,
    )
    r = top + 8
    # A: the step.
    label_row(ws, r, "A   THE STEP", cols)
    r += 1
    put(ws, r, 2, d["stepQuestion"], f=font(11, True))
    inp(ws, r, 3, d["profiles"][a["profile"]]["name"] if a and a["profile"] is not None else None)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    choice_list(ws, [f"C{r}"], [p["name"] for p in d["profiles"]], d["stepWhy"])
    name_cell(wb, ws, "in_profile", r, 3)
    profile_ref = f"$C${r}"
    r += 1
    for p in d["profiles"]:
        put(ws, r, 2, f"{p['name']}: {p['cost']}", f=font(10, color=MUTED))
        r += 1
    idx = "IF(" + ",IF(".join(f"{profile_ref}={q(p['name'])},{i + 1}" for i, p in enumerate(d["profiles"])) + ',""' + ")" * len(d["profiles"])
    idx_row = r
    put(ws, idx_row, 2, "Job, as a number (worked out)", f=font(9, color=MUTED))
    put(ws, idx_row, 3, "=" + idx, f=font(10, color=MUTED), align=WRAP_C)
    idx_ref = f"$C${idx_row}"
    r += 2

    # B: the gates.
    label_row(ws, r, "B   DEPLOYMENT GATES   ·   ANY FAIL ENDS THE CANDIDATE", cols)
    r += 1
    gate_cells = []
    for i, g in enumerate(d["gates"]):
        put(ws, r, 1, f"B{i + 1:02d}", f=font(10, True, MUTED), border=UNDER)
        put(ws, r, 2, g["q"], f=font(11, True), border=UNDER)
        put(ws, r, 3, f"{d['gateAnswers']['pass']}: {g['passes']}", f=font(10, color=MUTED), border=UNDER)
        v = None
        if a and a["gates"][i] is not None:
            v = d["gateAnswers"]["pass"] if a["gates"][i] == 1 else d["gateAnswers"]["fail"]
        inp(ws, r, 6, v)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        gate_cells.append(f"F{r}")
        name_cell(wb, ws, f"gate{i + 1:02d}", r, 6)
        r += 1
    choice_list(ws, gate_cells, [d["gateAnswers"]["pass"], d["gateAnswers"]["fail"]], "Does the candidate pass this gate?")
    r += 1

    # C: the behaviours, weighted.
    label_row(ws, r, "C   BEHAVIOUR UNDER TEST   ·   WEIGHTED BY THE JOB", cols)
    r += 1
    for c, h in zip((3, 4, 5, 6, 7, 8, 9), ("0 means", "1 means", "2 means", "Score", "Weight", "Scored", "Out of")):
        put(ws, r, c, h, f=font(9, True, DEEP_GOLD), align=WRAP_C if c >= 6 else WRAP)
    r += 1
    score_cells, scored_cells, max_cells = [], [], []
    own_i = 0
    for i, cr in enumerate(d["criteria"]):
        put(ws, r, 1, f"C{i + 1:02d}", f=font(10, True, MUTED), border=UNDER)
        put(ws, r, 2, f"{cr['short']}. {cr['q']}" + ("  (key row)" if cr["key"] else ""), f=font(11, True), border=UNDER)
        for k in range(3):
            put(ws, r, 3 + k, cr["a"][k], f=font(10, color=MUTED), border=UNDER)
        inp(ws, r, 6, a["scores"][i] if a else None)
        name_cell(wb, ws, f"score{i + 1:02d}", r, 6)
        score_cells.append(f"F{r}")
        if cr["weights"]:
            w = cr["weights"]
            put(ws, r, 7, f'=IF({idx_ref}="","",CHOOSE({idx_ref},{w[0]},{w[1]},{w[2]}))', align=WRAP_C, border=UNDER)
        else:
            own_i += 1
            inp(ws, r, 7, a["own"][own_i - 1] if a else d["ownWeightDefault"])
            choice_list(ws, [f"G{r}"], list(range(1, d["ownWeightMax"] + 1)), "You set this weight: 1 to 5.")
            name_cell(wb, ws, f"own{own_i}", r, 7)
        put(ws, r, 8, f'=IF(OR(F{r}="",G{r}=""),0,F{r}*G{r})', align=WRAP_C, border=UNDER)
        put(ws, r, 9, f'=IF(G{r}="",0,2*G{r})', align=WRAP_C, border=UNDER)
        scored_cells.append(f"H{r}")
        max_cells.append(f"I{r}")
        r += 1
    choice_list(ws, [c for c in score_cells], [0, 1, 2], "0, 1 or 2, from the anchors in the row.")
    put(ws, r, 2, d["weightsNote"], f=font(9, color=MUTED))
    r += 2

    # D: the disqualifiers.
    label_row(ws, r, "D   DISQUALIFIERS   ·   ANY ONE THAT HAPPENED ENDS THE CANDIDATE", cols)
    r += 1
    dq_cells = []
    for i, dq in enumerate(d["disqualifiers"]):
        put(ws, r, 1, f"D{i + 1:02d}", f=font(10, True, MUTED), border=UNDER)
        put(ws, r, 2, dq["q"], f=font(11, True), border=UNDER)
        v = None
        if a and a["dq"][i] is not None:
            v = d["dqAnswers"]["yes"] if a["dq"][i] == 1 else d["dqAnswers"]["no"]
        inp(ws, r, 6, v)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        dq_cells.append(f"F{r}")
        name_cell(wb, ws, f"dq{i + 1}", r, 6)
        r += 1
    choice_list(ws, dq_cells, [d["dqAnswers"]["no"], d["dqAnswers"]["yes"]], "Did this happen on any run?")

    # The result, at the top.
    scores = f"{score_cells[0]}:{score_cells[-1]}"
    total = f"SUM({scored_cells[0]}:{scored_cells[-1]})"
    mx = f"SUM({max_cells[0]}:{max_cells[-1]})"
    gates = f"{gate_cells[0]}:{gate_cells[-1]}"
    dqs = f"{dq_cells[0]}:{dq_cells[-1]}"
    pct_row = top + 1
    pct_ref = f"$F${pct_row}"
    answered = f'IF({profile_ref}="",0,1)+COUNTA({gates})+COUNT({scores})+COUNTA({dqs})'
    gate_line, dq_line, *lines = d["cutLines"]

    def chain(field: str) -> str:
        expr = '""'
        for c in reversed(lines):
            expr = f"IF({pct_ref}>={c['min']},{q(c[field])},{expr})"
        return (
            f"=IF(COUNTIF({gates},{q(d['gateAnswers']['fail'])})>0,{q(gate_line[field])},"
            f"IF(COUNTIF({dqs},{q(d['dqAnswers']['yes'])})>0,{q(dq_line[field])},"
            f"IF(OR({answered}<{d['answerCount']},{pct_ref}=\"\"),{q('Incomplete' if field == 'name' else 'Choose the job and answer every row to read the score against the rubric.')},{expr})))"
        )

    at = result_block(
        ws,
        top,
        cols,
        [
            ("Weighted score", f'=IF({idx_ref}="","Choose the job first",{total}&" of "&{mx})', None),
            ("Percent", f'=IF({pct_ref}="","",{pct_ref}&"%")', None),
            ("Answered", f"={answered}&\" of {d['answerCount']}\"", None),
            ("Outcome", chain("name"), "out_outcome"),
            ("What to do", chain("what"), None),
            ("", "", None),
        ],
    )
    put(ws, pct_row, 5, "Percent, as a number", f=font(9, color=MUTED), bg=SOFT)
    put(ws, pct_row, 6, f'=IF({idx_ref}="","",IF({mx}>0,ROUND({total}/{mx}*100,0),0))', f=font(12, True), bg=SOFT, align=WRAP_C)
    name_cell(wb, ws, "out_pct", pct_row, 6)
    put(ws, top, 5, "Scored, as a number", f=font(9, color=MUTED), bg=SOFT)
    put(ws, top, 6, f"={total}", f=font(12, True), bg=SOFT, align=WRAP_C)
    name_cell(wb, ws, "out_total", top, 6)
    put(ws, top + 2, 5, "Answered, as a number", f=font(9, color=MUTED), bg=SOFT)
    put(ws, top + 2, 6, f"={answered}", f=font(12, True), bg=SOFT, align=WRAP_C)
    name_cell(wb, ws, "out_answered", top + 2, 6)
    name_cell(wb, ws, "out_outcome", at["out_outcome"], RESULT_COL)
    protect(ws)
    return ws


def build_msel(d: dict, out: Path):
    wb = Workbook()
    wb.remove(wb.active)
    msel_sheet(wb, d, "Score", None)
    read_me(
        wb,
        d["toolName"],
        d["credit"],
        d["pageUrl"],
        [
            ("What this tool is for", d["purpose"]),
            ("How to run it", d["howToRun"]),
            (
                "How the sheet works",
                [
                    "Score one candidate model for one step. Choose the job the model is being staffed on first: the weights in section C follow that answer.",
                    f"Answer each gate with {d['gateAnswers']['pass']} or {d['gateAnswers']['fail']}, score each behaviour 0, 1 or 2 from the anchors in its row, and answer each disqualifier with {d['dqAnswers']['no']} or {d['dqAnswers']['yes']}.",
                    f"Four rows carry no starting weight: you set those, 1 to {d['ownWeightMax']}. They start at {d['ownWeightDefault']}.",
                    "The weighted score, the percentage and the outcome are worked out for you, with the same rules as the page. Any failed gate or any disqualifier that happened ends the candidate, whatever it scored.",
                    "The sheet is protected so a formula is not overwritten by accident. There is no password: Review, then Unprotect Sheet, if you want to change it.",
                    "The two Example sheets are the page's reference candidates, already scored.",
                ],
            ),
            ("The rubric", [f"{c['score']}: {c['name']}. {c['what']}" for c in d["cutLines"]]),
            ("Where this comes from", [f"{d['credit']}. The same tool, with the scoring worked out as you go: {d['pageUrl']}"]),
        ],
    )
    for ex in d["examples"]:
        msel_sheet(wb, d, f"Example {ex['id'].upper()}", ex["assessment"])
    wb.active = 0
    wb.save(out)


# ── the Agent Authority Review ────────────────────────────────────────────────


def auth_sheet(wb, d: dict, title: str, rows: list | None):
    ws = wb.create_sheet(title)
    cols = 11
    ground(ws, cols, 60)
    for col, w in zip("ABCDEFGHIJK", (5, 30, 30, 13, 17, 26, 10, 14, 26, 22, 34)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["L"].hidden = True

    top = masthead(
        ws,
        d["headline"],
        [f"{d['toolName']}  ·  {d['credit']}  ·  {d['pageUrl']}", "One row per step of the workflow. The owner is named once undo cost, judgment and repeat are all answered."],
        cols,
    )
    label = lambda group, value: next(c["label"] for c in d[group] if c["value"] == value)
    owner_of = {r["key"]: r["owner"] for r in d["rubric"]}
    levels = [u["level"] for u in d["undo"]]

    head = top + 10
    q_by = {x["name"]: x for x in d["questions"]}
    headers = [
        "#",
        "Step",
        q_by["Evidence"]["column"],
        "Checked?",
        q_by["Judgment"]["column"],
        q_by["Hard limit"]["column"],
        q_by["Undo cost"]["column"],
        q_by["Repeat"]["column"],
        "If the evidence is missing",
        "Owner",
        "Flags",
    ]
    for c, h in enumerate(headers, start=1):
        put(ws, head, c, h, f=font(9, True, DEEP_GOLD), bg=PAPER, align=WRAP)
    ws.row_dimensions[head].height = 30

    first = head + 1
    owners, undos, used = [], [], []
    for i in range(d["maxRows"]):
        r = first + i
        row = rows[i] if rows and i < len(rows) else None
        put(ws, r, 1, i + 1, f=font(10, True, MUTED), align=WRAP_C, border=UNDER)
        for c, key in ((2, "step"), (3, "evidence"), (6, "limit"), (9, "missing")):
            put(ws, r, c, (row or {}).get(key) or None, f=font(10), bg=INPUT, border=UNDER, editable=True)
        inp(ws, r, 4, label("evidence", row["checked"]) if row and row["checked"] else None)
        inp(ws, r, 5, label("judgment", row["judgment"]) if row and row["judgment"] else None)
        inp(ws, r, 7, row["undo"] if row and row["undo"] else None)
        inp(ws, r, 8, label("repeat", row["repeat"]) if row and row["repeat"] else None)
        for c, key in ((2, "step"), (3, "evidence"), (4, "checked"), (5, "judgment"), (6, "limit"), (7, "undo"), (8, "repeat"), (9, "missing")):
            name_cell(wb, ws, f"{key}{i + 1:02d}", r, c)

        # Hidden helper: 1 when the row carries a step or any answer, as the page counts it.
        put(ws, r, 12, f"=IF(LEN(TRIM(B{r}))+LEN(TRIM(C{r}))+LEN(TRIM(F{r}))+LEN(TRIM(I{r}))+COUNTA(D{r},E{r},G{r},H{r})>0,1,0)")
        agent = "OR(" + ",".join(f"G{r}={q(lv)}" for lv in d["agentUndoLevels"]) + ")"
        put(
            ws,
            r,
            10,
            f'=IF(L{r}=0,"",IF(OR(G{r}="",E{r}="",H{r}=""),{q(d["notDecided"])},'
            f'IF(H{r}={q(label("repeat", "unknown"))},{q(owner_of["stop"])},'
            f'IF(E{r}={q(label("judgment", "one"))},{q(owner_of["code"])},'
            f'IF({agent},{q(owner_of["agent"])},{q(owner_of["suggest"])})))))',
            f=font(10, True, FOREST),
            border=UNDER,
        )
        assumed = label("evidence", "guessed")
        put(
            ws,
            r,
            11,
            f'=IF(D{r}={q(assumed)},{q(d["rowFlags"]["assumed"])},"")'
            f'&IF(AND(D{r}={q(assumed)},G{r}={q(d["irreversibleUndo"])})," ","")'
            f'&IF(G{r}={q(d["irreversibleUndo"])},{q(d["rowFlags"]["irreversible"])},"")',
            f=font(9, color=ERROR),
            border=UNDER,
        )
        name_cell(wb, ws, f"owner{i + 1:02d}", r, 10)
        owners.append(f"J{r}")
        undos.append(f"G{r}")
        used.append(f"L{r}")
    last = first + d["maxRows"] - 1
    choice_list(ws, [f"D{first}:D{last}"], [c["label"] for c in d["evidence"]], q_by["Evidence"]["ask"])
    choice_list(ws, [f"E{first}:E{last}"], [c["label"] for c in d["judgment"]], q_by["Judgment"]["ask"])
    choice_list(ws, [f"G{first}:G{last}"], levels, q_by["Undo cost"]["ask"])
    choice_list(ws, [f"H{first}:H{last}"], [c["label"] for c in d["repeat"]], q_by["Repeat"]["ask"])

    own = f"J{first}:J{last}"
    undo = f"G{first}:G{last}"
    total = f"SUM(L{first}:L{last})"
    count = lambda key: f"COUNTIF({own},{q(owner_of[key])})"
    decided = "+".join(count(k) for k in ("code", "agent", "suggest", "stop"))
    ceiling = '""'
    for lv in levels:
        ceiling = f"IF(COUNTIFS({own},{q(owner_of['agent'])},{undo},{q(lv)})>0,{q(lv)},{ceiling})"
    o = d["outcomes"]
    items = [
        ("Steps on the sheet", f"={total}", "out_total"),
        ("Decided", f"={decided}", "out_decided"),
        (owner_of["code"], f"={count('code')}", "out_code"),
        (owner_of["agent"], f"={count('agent')}", "out_agent"),
        (owner_of["suggest"], f"={count('suggest')}", "out_suggest"),
        (owner_of["stop"], f"={count('stop')}", "out_stop"),
        ("Highest undo cost the agent acts on alone", f"={ceiling}", "out_ceiling"),
        (
            "Outcome",
            f"=IF({total}=0,{q(o['none'])},IF(({decided})<{total},{q(o['progress'])},IF({count('stop')}>0,{q(o['stop'])},{q(o['decided'])})))",
            "out_outcome",
        ),
    ]
    at = result_block(ws, top, cols, items)
    for _, _, nm in items:
        name_cell(wb, ws, nm, at[nm], RESULT_COL)
    protect(ws)
    return ws


def build_auth(d: dict, out: Path):
    wb = Workbook()
    wb.remove(wb.active)
    auth_sheet(wb, d, "Sheet", None)
    rubric = [f"{r['when']}: {r['owner']}. {r['may']}" for r in d["rubric"]]
    read_me(
        wb,
        d["toolName"],
        d["credit"],
        d["pageUrl"],
        [
            ("What this tool is for", d["purpose"]),
            ("The five questions, asked for every step", [f"{x['n']}. {x['name']}: {x['ask']}" for x in d["questions"]]),
            ("The undo cost scale", [f"{u['level']}: {u['short']}" for u in d["undo"]]),
            ("How three answers become one owner", rubric),
            (
                "How the sheet works",
                [
                    f"One row per step, up to {d['maxRows']}. Type the step and the evidence it needs, then answer the three dropdowns: judgment, undo cost and repeat.",
                    "The owner is named beside each step with the same rules as the page, and the tally at the top counts them. A row stays not decided until all three answers are in.",
                    "The flags column repeats the two warnings the page gives: an assumed input, and a step that cannot be undone.",
                    "The sheet is protected so a formula is not overwritten by accident. There is no password: Review, then Unprotect Sheet, if you want to change it.",
                    "The Example sheet is the page's first worked example.",
                ],
            ),
            ("Where this comes from", [f"{d['credit']}. The same tool, with the owners worked out as you type: {d['pageUrl']}"]),
        ],
    )
    auth_sheet(wb, d, "Example", d["example"]["rows"])
    wb.active = 0
    wb.save(out)


def main():
    out_dir = Path(sys.argv[1] if len(sys.argv) > 1 else "downloads")
    # UTF-8, whatever the console says: on Windows sys.stdin is cp1252, and
    # ₹, → and every curly quote from the data modules came out as mojibake.
    data = json.loads(sys.stdin.buffer.read().decode("utf-8"))
    for key, build in (("poc-screen", build_poc), ("model-selection-tool", build_msel), ("agent-authority-review", build_auth)):
        d = data[key]
        path = out_dir / d["file"]
        build(d, path)
        print(f"{path}  {path.stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
