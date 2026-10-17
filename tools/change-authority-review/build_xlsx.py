#!/usr/bin/env python3
"""Build downloads/change-authority-review.xlsx (behind the download gate; see CLAUDE.md).

Five tabs: Start here, 1 · Action Inventory, 2 · Authority Requirements,
3 · Checklist, and Worked Example. The three working tabs are blank. The
Worked Example tab stacks the same three sections on one sheet, filled in and
locked, and its formulas point at its own rows, so it computes the brief's
acceptance figures by itself.

WHERE THE CONTENT COMES FROM. Not from this file. `dump_content.ts` prints the
words, the ten checks and the worked example out of
src/data/change-authority-review.ts and src/lib/changeAuthority.ts, and this
script renders them. The workbook therefore says what the page says.

WHERE THE BRAND COMES FROM. tools/rework-cost-check/brand.json, which that
tool's extract_brand.py reads out of src/styles/ds/theme.css. Reused rather
than copied, as the Citation Fit Check and the Gate Routing Map do. The three
authority colours are the brand's own quiet fills, each with its matching ink,
listed in AUTHORITY_STYLE. None of them is the success green, so an authority
cell never reads as a "Ready".

FORMULA COMPATIBILITY. Only SUM, SUMPRODUCT, IF, AND, OR, NOT, MIN, COUNTIF,
COUNTIFS, COUNTA, INDEX and MATCH. IFS is avoided: openpyxl writes it without
the `_xlfn.` prefix and LibreOffice then returns #NAME?. A test across two
columns with a "" criterion is a SUMPRODUCT, because Excel and LibreOffice
disagree about whether COUNTIF(range,"") counts a formula's "".

Every rule below mirrors a function in src/lib/changeAuthority.ts, named beside
it. verify_xlsx.py recalculates the file in LibreOffice and asserts the figures.

Run:  python3 tools/change-authority-review/build_xlsx.py
"""
from __future__ import annotations

import json
import math
import shutil
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
BRAND = json.loads((REPO / "tools" / "rework-cost-check" / "brand.json").read_text(encoding="utf-8"))
BRAND_DIR = REPO / "public" / "brand"
BUILD = HERE / ".build"
OUT = REPO / "downloads" / "change-authority-review.xlsx"
REFS = HERE / "refs.json"

FOOTER = "© The Living Craft · thelivingcraft.ai · Content CC BY 4.0"

K = BRAND["colors"]
DISPLAY = BRAND["fonts"]["display"]["family"]
BODY = BRAND["fonts"]["body"]["family"]

AA, PC, NA = "Agent acts", "Preview and consent", "Not allowed"


def rgb(name: str) -> str:
    return "FF" + K[name].lstrip("#").upper()


def solid(name: str) -> PatternFill:
    return PatternFill("solid", fgColor=rgb(name))


FILL_HEAD = solid("noir")
FILL_INPUT = solid("accent_quiet")
FILL_OK, FILL_WARN, FILL_BAD = solid("success_quiet"), solid("warning_quiet"), solid("danger_quiet")
# A Tab 2 column that does not apply to the row's authority. The hairline
# colour, darker than every authority fill, so a greyed cell never reads as one.
FILL_OFF = solid("border_hair")

# The three authorities: a quiet brand fill and the ink that passes AA on it.
# Every authority cell also carries its word, so colour is never the only signal.
AUTHORITY_STYLE = {
    AA: ("info_quiet", "info"),
    PC: ("accent_quiet", "accent_ink"),
    NA: ("surface_hover", "text_quiet"),
}

EDGE_INPUT = Border(*[Side(style="thin", color=rgb("accent_ink"))] * 4)
HAIR = Border(bottom=Side(style="thin", color=rgb("border_hair")))
GOLD_RULE = Border(bottom=Side(style="medium", color=rgb("ember")))

# --weight-display is 400 in this system: the serif is set at its regular weight.
F_TITLE = Font(name=DISPLAY, size=26, color=rgb("text_loud"))
F_H1 = Font(name=DISPLAY, size=18, color=rgb("text_loud"))
F_H2 = Font(name=DISPLAY, size=14, color=rgb("text_loud"))
F_LEAD = Font(name=BODY, size=12, color=rgb("text_quiet"))
F_LABEL = Font(name=BODY, size=9, bold=True, color=rgb("accent_ink"))
F_HEAD = Font(name=BODY, size=10, bold=True, color=rgb("text_on_invert"))
F_BODY = Font(name=BODY, size=11, color=rgb("text_body"))
F_BODY_B = Font(name=BODY, size=11, bold=True, color=rgb("text_body"))
F_NOTE = Font(name=BODY, size=9, color=rgb("text_quiet"))
F_BIG = Font(name=BODY, size=16, bold=True, color=rgb("text_loud"))
F_BANNER = Font(name=BODY, size=11, bold=True, color=rgb("warning"))

WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
WRAP_MID = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
INPUT_TEXT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
HEAD_ALIGN = Alignment(horizontal="left", vertical="bottom", wrap_text=True)

# Column widths per section. The Worked Example stacks all three sections on
# one sheet, so it takes the widest of each column.
INV_COLS = {"A": 5, "B": 30, "C": 32, "D": 14, "E": 20, "F": 12, "G": 20, "H": 20, "I": 20, "J": 13, "K": 36}
REQ_COLS = {"A": 5, "B": 30, "C": 20, "D": 11, "E": 22, "F": 14, "G": 16, "H": 20, "I": 17, "J": 20,
            "K": 20, "L": 12, "M": 12, "N": 36}
CHK_COLS = {"A": 7, "B": 34, "C": 26, "D": 24, "E": 13, "F": 10, "G": 12, "H": 52}
LIST_COL = "Z"    # the dropdown options, hidden
PRIO_COL = "Y"    # the checklist's next-step priority, hidden


def content() -> dict:
    """The words and the worked example, from the TypeScript that owns them."""
    npx = shutil.which("npx") or "npx"
    out = subprocess.run(
        [npx, "-y", "tsx", str(HERE / "dump_content.ts")],
        capture_output=True, text=True, encoding="utf-8", cwd=REPO, check=False,
    )
    if out.returncode != 0:
        raise SystemExit(f"dump_content.ts failed:\n{out.stderr[-2000:]}")
    return json.loads(out.stdout)


# ---------------------------------------------------------------------------
# Cell helpers (the Gate Routing Map's, unchanged where they could be)
# ---------------------------------------------------------------------------


def put(ws, ref, value=None, *, font=F_BODY, fill=None, align=LEFT, fmt=None, editable=False, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font
    c.alignment = align
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if border is not None:
        c.border = border
    # Sheet protection is a guide rail, not security: calculated cells stay
    # locked so a reader does not type over a formula by accident.
    c.protection = Protection(locked=not editable)
    return c


def inp(ws, ref, value=None, *, align=INPUT_TEXT, editable=True):
    return put(ws, ref, value if value not in ("", None) else None, fill=FILL_INPUT, align=align,
               editable=editable, border=EDGE_INPUT)


def calc(ws, ref, formula, *, fmt=None, font=F_BODY_B, align=LEFT):
    return put(ws, ref, formula, font=font, align=align, fmt=fmt)


def merged(ws, a, b, text, *, font=F_BODY, align=WRAP):
    ws.merge_cells(f"{a}:{b}")
    return put(ws, a, text, font=font, align=align)


def q(s: str) -> str:
    """A string literal inside a formula."""
    return '"' + s.replace('"', '""') + '"'


def fit(ws, row: int, texts: list[tuple[str, float]], *, floor: float = 20, line: float = 14.5):
    """Set a row height from its longest wrapped text.

    Excel does not grow a row to fit wrapped text in a file it did not write, so
    a row left at the default height shows one line of a three-line cell. The
    estimate is generous on purpose: on a machine without Figtree, Excel draws a
    wider substitute.
    """
    lines = 1
    for text, width in texts:
        if not text:
            continue
        per_line = max(1.0, width * 1.05)
        lines = max(lines, sum(math.ceil(max(1, len(part)) / per_line) for part in str(text).split("\n")))
    ws.row_dimensions[row].height = max(floor, lines * line + 6)


class Lists:
    """Dropdown options written into a hidden column, one block per list."""

    def __init__(self, ws: Worksheet):
        self.ws, self.next, self.made = ws, 1, {}
        ws.column_dimensions[LIST_COL].hidden = True

    def validation(self, options: list[str]) -> DataValidation:
        key = tuple(options)
        if key not in self.made:
            top = self.next
            for i, option in enumerate(options):
                put(self.ws, f"{LIST_COL}{top + i}", option)
            self.next = top + len(options) + 1
            dv = DataValidation(type="list", formula1=f"${LIST_COL}${top}:${LIST_COL}${top + len(options) - 1}",
                                allow_blank=True, showDropDown=False)
            dv.error, dv.errorTitle = "Pick one of the options in the list.", "Not an option"
            self.ws.add_data_validation(dv)
            self.made[key] = dv
        return self.made[key]


def header_row(ws, row: int, heads: list[str], *, height: float = 32):
    for i, h in enumerate(heads):
        put(ws, f"{chr(65 + i)}{row}", h, font=F_HEAD, fill=FILL_HEAD, align=HEAD_ALIGN)
    ws.row_dimensions[row].height = height


def section_head(ws, top: int, tab: dict, note: str, last_col: str):
    put(ws, f"A{top}", f"{int(tab['n'])} · {tab['name']}", font=F_H1)
    ws.row_dimensions[top].height = 28
    merged(ws, f"A{top + 1}", f"{last_col}{top + 1}", tab["ask"], font=F_LEAD, align=WRAP_MID)
    ws.row_dimensions[top + 1].height = 20
    merged(ws, f"A{top + 2}", f"{last_col}{top + 2}", note, font=F_NOTE, align=WRAP_MID)
    ws.row_dimensions[top + 2].height = 30


def summary(ws, row: int, label: str, formula: str, *, note: str | None = None, note_to: str = "H"):
    merged(ws, f"A{row}", f"B{row}", label, font=F_BODY, align=WRAP_MID)
    calc(ws, f"C{row}", formula, font=F_BIG)
    ws[f"C{row}"].border = HAIR
    if note:
        merged(ws, f"D{row}", f"{note_to}{row}", note, font=F_NOTE, align=WRAP_MID)
    ws.row_dimensions[row].height = 30 if len(label) > 34 else 24


def danger_when_positive(ws, cell: str):
    ws.conditional_formatting.add(cell, FormulaRule(
        formula=[f"{cell}>0"],   # the guarded cells are counts, never blank
        fill=FILL_BAD, font=Font(name=BODY, size=16, bold=True, color=rgb("danger"))))


def authority_colours(ws, cells: str, first: str):
    for authority, (fill, ink) in AUTHORITY_STYLE.items():
        ws.conditional_formatting.add(cells, FormulaRule(
            formula=[f'{first}="{authority}"'], fill=solid(fill), font=Font(name=BODY, bold=True, color=rgb(ink))))


def paint_authority(ws, ref: str, authority: str):
    """A fixed authority label, painted the way the conditional rule paints a cell."""
    fill, ink = AUTHORITY_STYLE[authority]
    ws[ref].fill = solid(fill)
    ws[ref].font = Font(name=BODY, size=10, bold=True, color=rgb(ink))


def protect(ws: Worksheet):
    ws.protection.sheet = True          # no password: a guide rail, not security
    ws.protection.formatColumns = False  # a reader may still widen a column
    ws.protection.formatRows = False


def grid(ws: Worksheet, *widths: dict):
    cols: dict[str, float] = {}
    for w in widths:
        for col, v in w.items():
            cols[col] = max(cols.get(col, 0), v)
    for col, v in cols.items():
        ws.column_dimensions[col].width = v
    ws.column_dimensions[PRIO_COL].hidden = True


# ---------------------------------------------------------------------------
# The three sections. Each writes at row `top` and returns the cells it owns.
# ---------------------------------------------------------------------------

INV_KEYS = {"B": "name", "C": "changes", "D": "undo", "E": "seenBy", "F": "tested"}


def inventory(ws, top: int, d: dict, lists: Lists, rows: list[dict] | None, *, editable: bool = True) -> dict:
    tab = d["tabs"][0]
    section_head(ws, top, tab, d["inventoryNote"], "K")
    hdr = top + 4
    first, last = hdr + 1, hdr + d["rows"]
    header_row(ws, hdr, ["#", "Action", "What it changes", "Can it be undone?",
                         "Who sees the result before it can be undone?", "Undo tested?",
                         "Suggested authority", "Authority chosen", "Final authority",
                         "Differs from suggestion?", "Reason for the difference"], height=46)
    o = d["options"]
    dv = {"D": lists.validation(o["undo"]), "E": lists.validation(o["seenBy"]),
          "F": lists.validation(o["yesNo"]), "H": lists.validation(o["authorities"])}
    for i in range(d["rows"]):
        r = first + i
        row = rows[i] if rows and i < len(rows) else {}
        put(ws, f"A{r}", i + 1, font=F_NOTE, align=CENTER, border=HAIR)
        for col, key in INV_KEYS.items():
            inp(ws, f"{col}{r}", row.get(key), editable=editable, align=INPUT_TEXT if col in "BC" else CENTER)
        inp(ws, f"H{r}", row.get("chosen"), align=CENTER, editable=editable)
        inp(ws, f"K{r}", row.get("reason"), editable=editable)
        for col, v in dv.items():
            v.add(f"{col}{r}")
        # suggestedAuthority() in src/lib/changeAuthority.ts. Never "Not allowed".
        calc(ws, f"G{r}",
             f'=IF(OR(B{r}="",D{r}="",E{r}=""),"",'
             f'IF(AND(D{r}="Yes",OR(E{r}="Nobody",E{r}="The user only")),"{AA}","{PC}"))', align=CENTER)
        calc(ws, f"I{r}", f'=IF(H{r}<>"",H{r},G{r})', align=CENTER)                 # finalAuthority()
        calc(ws, f"J{r}", f'=IF(AND(H{r}<>"",H{r}<>G{r}),"Yes","")', align=CENTER)  # differs()
        for col in "GIJ":
            ws[f"{col}{r}"].border = HAIR
        fit(ws, r, [(row.get("name", ""), INV_COLS["B"]), (row.get("changes", ""), INV_COLS["C"]),
                    (row.get("reason", ""), INV_COLS["K"])], floor=30)
    authority_colours(ws, f"G{first}:G{last}", f"G{first}")
    authority_colours(ws, f"H{first}:H{last}", f"H{first}")
    authority_colours(ws, f"I{first}:I{last}", f"I{first}")

    # inventorySummary() in src/lib/changeAuthority.ts. Below the table, as the brief asks.
    B, F, G, I_, J, Kc = (f"{c}{first}:{c}{last}" for c in "BFGIJK")
    s = last + 2
    put(ws, f"A{s}", "Summary", font=F_H2)
    ws.row_dimensions[s].height = 26
    summary(ws, s + 1, "Actions listed", f"=COUNTA({B})", note_to="K")
    put(ws, f"A{s + 3}", "ACTIONS BY AUTHORITY", font=F_LABEL)
    for col, h in zip("CD", ["Final", "Suggested"]):
        put(ws, f"{col}{s + 3}", h, font=F_LABEL, align=CENTER)
    by, suggested = {}, {}
    for j, authority in enumerate(d["options"]["authorities"]):
        r = s + 4 + j
        put(ws, f"B{r}", authority, border=HAIR)
        paint_authority(ws, f"B{r}", authority)
        calc(ws, f"C{r}", f'=COUNTIF({I_},"{authority}")', align=CENTER)
        calc(ws, f"D{r}", f'=COUNTIF({G},"{authority}")', font=F_NOTE, align=CENTER)
        by[authority], suggested[authority] = f"C{r}", f"D{r}"
        ws.row_dimensions[r].height = 22
    a = s + 4 + len(d["options"]["authorities"]) + 1
    summary(ws, a, "Actions with no authority", f'=SUMPRODUCT(({B}<>"")*({I_}=""))',
            note="Actions with a name and no authority. Fill in D and E, or choose an authority.", note_to="K")
    danger_when_positive(ws, f"C{a}")
    summary(ws, a + 1, "Undo not tested", f'=SUMPRODUCT(({B}<>"")*({F}<>"Yes"))',
            note="Actions where nobody has actually tried the undo.", note_to="K")
    summary(ws, a + 2, "Overrides", f'=COUNTIF({J},"Yes")',
            note="Actions where the authority chosen differs from the suggestion.", note_to="K")
    summary(ws, a + 3, "Overrides with no reason", f'=SUMPRODUCT(({J}="Yes")*({Kc}=""))',
            note="Write the reason in column K.", note_to="K")
    danger_when_positive(ws, f"C{a + 3}")
    summary(ws, a + 4, "Actions that need consent but the agent takes alone",
            f'=COUNTIFS({G},"{PC}",{I_},"{AA}")',
            note="Suggested Preview and consent, chosen Agent acts: an override that removed consent.", note_to="K")
    danger_when_positive(ws, f"C{a + 4}")
    return {"hdr": hdr, "first": first, "last": last, "end": a + 4,
            "listed": f"C{s + 1}", "byAuthority": by, "suggested": suggested,
            "noAuthority": f"C{a}", "undoNotTested": f"C{a + 1}", "overrides": f"C{a + 2}",
            "overridesNoReason": f"C{a + 3}", "consentRemoved": f"C{a + 4}",
            "suggestedCells": [f"G{first + i}" for i in range(d["rows"])],
            "finalCells": [f"I{first + i}" for i in range(d["rows"])],
            "differsCells": [f"J{first + i}" for i in range(d["rows"])]}


# Tab 2: which columns belong to which authority (AUTHORITY_FIELDS in changeAuthority.ts).
REQ_GROUPS = [
    (AA, "E", ["Undo path"]),
    (PC, "FGHIJK", ["Preview shows the full change?", "Preview says what can’t be undone?", "Who consents",
                    "Consent is given per change (not a standing “always allow”)?",
                    "Agent acts as that person, and the log records who approved?",
                    "If consent is refused, the agent…"]),
    (NA, "L", ["Access removed?"]),
]
REQ_KEYS = {"E": "undoPath", "F": "fullChange", "G": "saysIrreversible", "H": "consenter", "I": "perChange",
            "J": "actsAs", "K": "onRefusal", "L": "accessRemoved"}


def requirements(ws, top: int, d: dict, lists: Lists, rows: list[dict] | None, link, *,
                 editable: bool = True) -> dict:
    """`link(col, i)` is the Tab 1 cell for column `col` of action i (0-based)."""
    tab = d["tabs"][1]
    section_head(ws, top, tab, "Rows mirror the Action Inventory. Fill in only the columns for the row’s final "
                               "authority; the others are greyed out. Each row names the first thing missing.", "N")
    grp = top + 4
    hdr = grp + 1
    first, last = hdr + 1, hdr + d["rows"]
    for authority, cols, _ in REQ_GROUPS:
        a, b = f"{cols[0]}{grp}", f"{cols[-1]}{grp}"
        if a != b:
            ws.merge_cells(f"{a}:{b}")
        put(ws, a, authority)
        paint_authority(ws, a, authority)
    ws.row_dimensions[grp].height = 20
    heads = ["#", "Action", "Final authority", "Undo tested?"]
    for _, _, h in REQ_GROUPS:
        heads += h
    header_row(ws, hdr, heads + ["Authority ready?", "Why not ready"], height=72)

    o = d["options"]
    dv_yn, dv_ypn, dv_ref = (lists.validation(o["yesNo"]), lists.validation(o["answers"]),
                             lists.validation(o["onRefusal"]))
    # A text column that does not apply to the row's authority refuses input.
    # The brief asks for those columns locked; protection cannot depend on a
    # formula, so validation does the refusing. A dropdown cell already holds
    # its list validation, and a cell takes only one, so those are greyed only.
    refuse = {}
    for authority in (AA, PC):
        v = DataValidation(type="custom", formula1=f'$C{first}="{authority}"', allow_blank=True)
        v.error = f"This column is for {authority} actions. Change the authority on the Action Inventory first."
        v.errorTitle = "Not this authority"
        ws.add_data_validation(v)
        refuse[authority] = v

    W = d["why"]
    for i in range(d["rows"]):
        r = first + i
        row = rows[i] if rows and i < len(rows) else {}
        put(ws, f"A{r}", i + 1, font=F_NOTE, align=CENTER, border=HAIR)
        for col, src_col in zip("BCD", "BIF"):
            src = link(src_col, i)
            calc(ws, f"{col}{r}", f'=IF({src}="","",{src})', font=F_BODY if col == "B" else F_BODY_B,
                 align=WRAP_MID if col == "B" else CENTER)
            ws[f"{col}{r}"].border = HAIR
        for authority, cols, _ in REQ_GROUPS:
            for col in cols:
                inp(ws, f"{col}{r}", row.get(REQ_KEYS[col]), editable=editable,
                    align=INPUT_TEXT if col in "EH" else CENTER)
                if col == "F":
                    dv_ypn.add(f"{col}{r}")
                elif col == "K":
                    dv_ref.add(f"{col}{r}")
                elif col in "GIJL":
                    dv_yn.add(f"{col}{r}")
                else:
                    refuse[authority].add(f"{col}{r}")

        # authorityReadiness() in src/lib/changeAuthority.ts. "Why not ready"
        # names the first missing item; "Authority ready?" is Ready exactly when
        # it is blank, so the two columns cannot disagree.
        C, D, E, F, G, H, I_, J, Kc, L = (f"{c}{r}" for c in "CDEFGHIJKL")
        agent = f'IF({E}="",{q(W["noUndoPath"])},IF({D}<>"Yes",{q(W["undoNotTested"])},""))'
        consent = (f'IF({F}<>"Yes",{q(W["previewPartial"])},IF({G}<>"Yes",{q(W["previewHidesIrreversible"])},'
                   f'IF({H}="",{q(W["noConsenter"])},IF({I_}<>"Yes",{q(W["standingConsent"])},'
                   f'IF({J}<>"Yes",{q(W["notActingAs"])},'
                   f'IF(AND({Kc}<>"Changes nothing",{Kc}<>"Revises and asks again"),{q(W["unsafeRefusal"])},""))))))')
        denied = f'IF({L}<>"Yes",{q(W["accessPresent"])},"")'
        calc(ws, f"N{r}", f'=IF({C}="","",IF({C}="{AA}",{agent},IF({C}="{PC}",{consent},'
                          f'IF({C}="{NA}",{denied},""))))', font=F_BODY, align=WRAP_MID)
        calc(ws, f"M{r}", f'=IF({C}="","",IF(N{r}="","Ready","Not ready"))', align=CENTER)
        for col in "MN":
            ws[f"{col}{r}"].border = HAIR
        fit(ws, r, [(row.get("_name", ""), REQ_COLS["B"]), (row.get("undoPath", ""), REQ_COLS["E"]),
                    (row.get("onRefusal", ""), REQ_COLS["K"])], floor=30)

    for authority, cols, _ in REQ_GROUPS:
        rng = f"{cols[0]}{first}:{cols[-1]}{last}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'$C{first}<>"{authority}"'], fill=FILL_OFF, font=Font(name=BODY, color=rgb("text_quiet"))))
    authority_colours(ws, f"C{first}:C{last}", f"C{first}")
    for word, fill, colour in [("Ready", FILL_OK, "success"), ("Not ready", FILL_BAD, "danger")]:
        ws.conditional_formatting.add(f"M{first}:M{last}", FormulaRule(
            formula=[f'M{first}="{word}"'], fill=fill, font=Font(name=BODY, bold=True, color=rgb(colour))))

    # requirementsSummary() in src/lib/changeAuthority.ts
    Cr, Fr, Gr, Hr, Ir, Jr, Kr, Lr, M = (f"{c}{first}:{c}{last}" for c in "CFGHIJKLM")
    s = last + 2
    put(ws, f"A{s}", "Summary", font=F_H2)
    ws.row_dimensions[s].height = 26
    put(ws, f"A{s + 1}", "BY AUTHORITY", font=F_LABEL)
    for col, h in zip("CD", ["Actions", "Ready"]):
        put(ws, f"{col}{s + 1}", h, font=F_LABEL, align=CENTER)
    actions, ready = {}, {}
    for j, authority in enumerate(d["options"]["authorities"]):
        r = s + 2 + j
        put(ws, f"B{r}", authority, border=HAIR)
        paint_authority(ws, f"B{r}", authority)
        calc(ws, f"C{r}", f'=COUNTIF({Cr},"{authority}")', align=CENTER)
        calc(ws, f"D{r}", f'=COUNTIFS({Cr},"{authority}",{M},"Ready")', align=CENTER)
        actions[authority], ready[authority] = f"C{r}", f"D{r}"
        ws.row_dimensions[r].height = 22
    a = s + 2 + len(d["options"]["authorities"]) + 1
    pc = f'({Cr}="{PC}")'
    summary(ws, a, "Authority not ready", f'=COUNTIF({M},"Not ready")',
            note="Column N names the first thing missing on each.", note_to="J")
    danger_when_positive(ws, f"C{a}")
    summary(ws, a + 1, "Previews that hide part of the change",
            f'=SUMPRODUCT({pc}*((({Fr}<>"Yes")+({Gr}<>"Yes"))>0))',
            note="The preview does not show the full change, or does not say what can’t be undone.", note_to="J")
    summary(ws, a + 2, "Consent that isn’t per change or has no named person",
            f'=SUMPRODUCT({pc}*((({Ir}<>"Yes")+({Hr}=""))>0))', note_to="J")
    summary(ws, a + 3, "Agent not acting as the consenting person", f'=SUMPRODUCT({pc}*({Jr}<>"Yes"))',
            note_to="J")
    summary(ws, a + 4, "No safe outcome on refusal",
            f'=SUMPRODUCT({pc}*({Kr}<>"Changes nothing")*({Kr}<>"Revises and asks again"))',
            note="Proceeds anyway, Not decided, or left blank.", note_to="J")
    summary(ws, a + 5, "Not allowed, access still present", f'=SUMPRODUCT(({Cr}="{NA}")*({Lr}<>"Yes"))',
            note_to="J")
    return {"hdr": hdr, "first": first, "last": last, "end": a + 5,
            "actions": actions, "ready": ready, "notReady": f"C{a}", "previewHides": f"C{a + 1}",
            "consentWeak": f"C{a + 2}", "notActingAs": f"C{a + 3}", "unsafeRefusal": f"C{a + 4}",
            "accessPresent": f"C{a + 5}",
            "readyCells": [f"M{first + i}" for i in range(d["rows"])],
            "whyCells": [f"N{first + i}" for i in range(d["rows"])],
            "inputs": {k: [f"{c}{first + i}" for i in range(d["rows"])] for c, k in REQ_KEYS.items()}}


# Where each piece of evidence lives, and its label beside a check (the brief's hints).
EVIDENCE_LABEL = {
    "listed": ("Actions listed", "Tab 1"),
    "noAuthority": ("Actions with no authority", "Tab 1"),
    "undoNotTested": ("Undo not tested", "Tab 1"),
    "consentRemoved": ("Consent removed by override", "Tab 1"),
    "overridesNoReason": ("Overrides with no reason", "Tab 1"),
    "previewHides": ("Previews that hide part of the change", "Tab 2"),
    "consentWeak": ("Consent that isn’t per change or has no named person", "Tab 2"),
    "notActingAs": ("Agent not acting as the consenting person", "Tab 2"),
    "unsafeRefusal": ("No safe outcome on refusal", "Tab 2"),
    "accessPresent": ("Not allowed, access still present", "Tab 2"),
}


def checklist(ws, top: int, d: dict, lists: Lists, answers: list[str] | None, ev: dict, where: dict, *,
              editable: bool = True) -> dict:
    """`ev` maps an evidence key to the cell holding it; `where` names the tab or section it came from."""
    tab = d["tabs"][2]
    section_head(ws, top, tab, "Answer each check from evidence, not from memory. The four marked ★ are "
                               "critical: a No on any one of them is a Hold.", "H")
    r_dec, r_total, r_answered, r_next = top + 4, top + 5, top + 6, top + 7
    hdr = top + 9
    first, last = hdr + 1, hdr + len(d["checks"])
    E, Fp, Y = f"E{first}:E{last}", f"F{first}:F{last}", f"{PRIO_COL}{first}:{PRIO_COL}{last}"

    dv = lists.validation(d["options"]["answers"])
    row_of = {}
    for i, c in enumerate(d["checks"]):
        r = first + i
        row_of[c["n"]] = r
        given = answers[c["n"] - 1] if answers else None
        put(ws, f"A{r}", f"{c['n']} ★" if c["critical"] else c["n"], font=F_BODY_B if c["critical"] else F_NOTE,
            align=CENTER, border=HAIR)
        merged(ws, f"B{r}", f"D{r}", c["text"], align=WRAP_MID)
        for col in "BCD":
            ws[f"{col}{r}"].border = HAIR
        inp(ws, f"E{r}", given, align=CENTER, editable=editable)
        dv.add(f"E{r}")
        # POINTS and STATUS in src/lib/changeAuthority.ts
        calc(ws, f"F{r}", f'=IF(E{r}="Yes",2,IF(E{r}="Partial",1,IF(E{r}="No",0,"")))', align=CENTER)
        calc(ws, f"G{r}", f'=IF(E{r}="Yes","Pass",IF(E{r}="Partial","Partial",IF(E{r}="No","Fail","")))',
             align=CENTER)
        if c["evidence"]:
            parts = [f'"{EVIDENCE_LABEL[k][0]} ({where[k]}): "&{ev[k]}' for k in c["evidence"]]
            calc(ws, f"H{r}", "=" + '&" · "&'.join(parts), font=F_NOTE, align=WRAP_MID)
        # priority() in changeAuthority.ts: 1 critical No, 2 No, 3 critical Partial, 4 Partial, 9 nothing.
        calc(ws, f"{PRIO_COL}{r}",
             f'=IF(E{r}="No",{1 if c["critical"] else 2},IF(E{r}="Partial",{3 if c["critical"] else 4},9))',
             font=F_NOTE)
        for col in "FGH":
            ws[f"{col}{r}"].border = HAIR
        fit(ws, r, [(c["text"], CHK_COLS["B"] + CHK_COLS["C"] + CHK_COLS["D"])],
            floor=36 if c["evidence"] else 26)

    for word, fill, colour in [("Pass", FILL_OK, "success"), ("Partial", FILL_WARN, "warning"), ("Fail", FILL_BAD, "danger")]:
        ws.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(
            formula=[f'G{first}="{word}"'], fill=fill, font=Font(name=BODY, bold=True, color=rgb(colour))))

    # checklistResult() in src/lib/changeAuthority.ts
    crit = [c for c in d["checks"] if c["critical"]]
    crit_no = ",".join(f'E{row_of[c["n"]]}="No"' for c in crit)
    crit_partial = ",".join(f'E{row_of[c["n"]]}="Partial"' for c in crit)
    n = len(d["checks"])
    dec = f"C{r_dec}"
    merged(ws, f"A{r_dec}", f"B{r_dec}", "Decision", font=F_BODY_B, align=WRAP_MID)
    calc(ws, dec, f'=IF(OR({crit_no}),"Hold",IF(COUNTA({E})<{n},"",'
                  f'IF(AND(SUM({Fp})>={d["readyAt"]},NOT(OR({crit_partial}))),"Ready","Fix first")))',
         font=F_BIG, align=CENTER)
    m = d["decisionMeaning"]
    merged(ws, f"D{r_dec}", f"H{r_dec}",
           f'=IF({dec}="Hold",{q(m["Hold"])},IF({dec}="Fix first",{q(m["Fix first"])},'
           f'IF({dec}="Ready",{q(m["Ready"])},IF(COUNTA({E})=0,"","Answer all {n} checks to see the decision."))))',
           font=F_BODY, align=WRAP_MID)
    ws[f"D{r_dec}"].protection = Protection(locked=True)
    ws.row_dimensions[r_dec].height = 34
    for word, fill, colour in [("Ready", FILL_OK, "success"), ("Fix first", FILL_WARN, "warning"), ("Hold", FILL_BAD, "danger")]:
        ws.conditional_formatting.add(dec, FormulaRule(
            formula=[f'{dec}="{word}"'], fill=fill, font=Font(name=BODY, size=16, bold=True, color=rgb(colour))))

    summary(ws, r_total, "Total", f"=SUM({Fp})",
            note=f"out of {d['maxTotal']}. Yes scores 2, Partial 1, No 0. Ready needs {d['readyAt']} or more "
                 f"and every critical check answered Yes.")
    summary(ws, r_answered, "Checks answered", f"=COUNTA({E})", note=f"out of {n}")
    merged(ws, f"A{r_next}", f"B{r_next}", "Your next step", font=F_BODY_B, align=WRAP_MID)
    merged(ws, f"C{r_next}", f"H{r_next}",
           f'=IF(MIN({Y})<9,INDEX(B{first}:B{last},MATCH(MIN({Y}),{Y},0)),'
           f'IF(COUNTA({E})={n},{q(d["noOpenItems"])},""))', font=F_BODY, align=WRAP_MID)
    ws.row_dimensions[r_next].height = 40
    ws[f"C{r_next}"].border = HAIR

    header_row(ws, hdr, ["#", "Check", "", "", "Answer", "Points", "Status", "Evidence from the other tabs"])
    ws.merge_cells(f"B{hdr}:D{hdr}")
    return {"hdr": hdr, "first": first, "last": last, "end": last, "decision": dec, "total": f"C{r_total}",
            "answered": f"C{r_answered}", "nextStep": f"C{r_next}",
            "answers": {str(k): f"E{v}" for k, v in row_of.items()},
            "status": {str(k): f"G{v}" for k, v in row_of.items()},
            "evidence": {str(k): f"H{v}" for k, v in row_of.items()}}


# ---------------------------------------------------------------------------
# Start here
# ---------------------------------------------------------------------------


def brand_lockup(mark_h: int = 44, name_h: int = 22, gap: int = 12):
    """The mark beside the lettering, composited into one PNG at twice its displayed size.

    Excel cannot embed WebP and the repository's brand files are WebP, so they
    are converted at build time rather than a PNG copy being committed and left
    to drift. Neither image is enlarged past its native resolution. The mark
    alone is also written at 2x, as the brief asks, to .build/lc-mark@2x.png.
    """
    from PIL import Image as PILImage

    def load(name: str, h: int):
        src = BRAND_DIR / f"{name}.webp"
        if not src.exists():
            raise SystemExit(f"missing brand asset {src.relative_to(REPO)}")
        im = PILImage.open(src).convert("RGBA")
        target = min(h * 2, im.height)
        return im.resize((max(1, round(im.width * target / im.height)), target), PILImage.LANCZOS)

    mark, name = load("lc-mark-keyed", mark_h), load("lc-name", name_h)
    pad = gap * 2
    w, h = mark.width + pad + name.width, max(mark.height, name.height)
    canvas = PILImage.new("RGBA", (w, h), (0, 0, 0, 0))
    canvas.paste(mark, (0, (h - mark.height) // 2), mark)
    canvas.paste(name, (mark.width + pad, (h - name.height) // 2), name)
    BUILD.mkdir(parents=True, exist_ok=True)
    mark.save(BUILD / "lc-mark@2x.png")
    out = BUILD / "lc-lockup@2x.png"
    canvas.save(out)
    display_h = max(mark_h, name_h)
    return out, round(w * display_h / h), display_h


def cover_table(ws, r: int, label: str, heads: list[str], rows: list[list[str]], *, paint: bool) -> int:
    put(ws, f"B{r}", label, font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    r += 1
    for col, h in zip("BCD", heads):
        put(ws, f"{col}{r}", h, font=F_HEAD, fill=FILL_HEAD, align=HEAD_ALIGN)
    ws.row_dimensions[r].height = 22
    r += 1
    for row in rows:
        for col, text in zip("BCD", row):
            put(ws, f"{col}{r}", text, align=WRAP_MID, border=HAIR)
        if paint:
            paint_authority(ws, f"B{r}", row[0])
        else:
            ws[f"B{r}"].font = F_BODY_B
        fit(ws, r, [(row[1], 46), (row[2], 64)], floor=24, line=15)
        r += 1
    return r + 1


def start_tab(ws: Worksheet, d: dict):
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 3, "B": 30, "C": 46, "D": 64}.items():
        ws.column_dimensions[col].width = w

    png, w, h = brand_lockup()
    img = XlImage(str(png))
    img.width, img.height = w, h
    ws.add_image(img, "B2")
    ws.row_dimensions[2].height = 36

    put(ws, "B4", d["fullTitle"], font=F_TITLE)
    ws.row_dimensions[4].height = 40
    for col in "BCD":
        ws[f"{col}4"].border = GOLD_RULE
    merged(ws, "B5", "D5", d["tagline"], font=F_LEAD, align=WRAP_MID)
    ws.row_dimensions[5].height = 26

    r = 7
    for label, text in [("What to bring", d["bring"]), ("What to do", d["do"]),
                        ("What you leave with", d["leave"])]:
        put(ws, f"B{r}", label.upper(), font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
        merged(ws, f"C{r}", f"D{r}", text, align=WRAP)
        fit(ws, r, [(text, 100)], floor=22, line=15)
        r += 2

    r = cover_table(ws, r, "THE THREE AUTHORITY LEVELS",
                    ["Authority", "Use it when", "What it needs before the agent may act"],
                    [[x["authority"], x["when"], x["needs"]] for x in d["authorityTable"]], paint=True)
    r = cover_table(ws, r, "CAN IT BE UNDONE?", ["Answer", "Meaning", "Example"],
                    [[x["answer"], x["meaning"], x["example"]] for x in d["undoTable"]], paint=False)

    put(ws, f"B{r}", "HOW TO USE IT", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    merged(ws, f"C{r}", f"D{r}", d["howToUse"], align=WRAP)
    r += 1
    for t in d["tabs"]:
        put(ws, f"B{r}", t["sheet"], font=F_BODY_B)
        merged(ws, f"C{r}", f"D{r}", t["ask"], align=WRAP_MID)
        ws.row_dimensions[r].height = 22
        r += 1
    r += 1

    put(ws, f"B{r}", "CELL LEGEND", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    put(ws, f"C{r}", d["legend"])
    r += 1
    inp(ws, f"C{r}", "Yours to fill", align=LEFT, editable=False)
    put(ws, f"D{r}", "Shaded, with a thin border.", font=F_NOTE)
    ws.row_dimensions[r].height = 22
    r += 1
    put(ws, f"C{r}", "Calculates for you", font=F_BODY_B)
    put(ws, f"D{r}", "No fill, and locked.", font=F_NOTE)
    ws.row_dimensions[r].height = 22
    r += 1
    put(ws, f"C{r}", "Does not apply", fill=FILL_OFF, font=Font(name=BODY, color=rgb("text_quiet")))
    put(ws, f"D{r}", "On Tab 2, a column that is not for this row’s authority.", font=F_NOTE)
    ws.row_dimensions[r].height = 22
    r += 2

    merged(ws, f"B{r}", f"D{r}", FOOTER, font=F_NOTE, align=LEFT)
    r += 1
    merged(ws, f"B{r}", f"D{r}", d["credit"], font=F_NOTE, align=LEFT)
    r += 2
    # The cohort invitation and a working application address, which the
    # outreach readiness handoff (28 September 2026) asks every file a reader
    # keeps to carry. The words come from src/data/resource-cohort-copy.ts.
    put(ws, f"B{r}", "THE LIVING CRAFT COHORT", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    merged(ws, f"C{r}", f"D{r}", d["cohort"], align=WRAP)
    fit(ws, r, [(d["cohort"], 100)], floor=22, line=15)
    r += 1
    ws.merge_cells(f"C{r}:D{r}")
    link = put(ws, f"C{r}", f"Apply at {d['applyUrl']}", font=F_BODY_B)
    link.hyperlink = d["applyUrl"]
    protect(ws)


# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------


def sheet_ref(sheet: str, cell: str) -> str:
    return f"'{sheet}'!{cell}"


def build() -> dict:
    d = content()
    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    tabs = {t["n"]: t["sheet"] for t in d["tabs"]}
    refs: dict = {"blank": {}, "example": {}}
    on_tab1 = {k for k, v in EVIDENCE_LABEL.items() if v[1] == "Tab 1"}

    # ---- the three blank working tabs -------------------------------------
    ws1 = wb.create_sheet(tabs["01"]); grid(ws1, INV_COLS)
    inv = inventory(ws1, 1, d, Lists(ws1), None)
    ws2 = wb.create_sheet(tabs["02"]); grid(ws2, REQ_COLS)
    req = requirements(ws2, 1, d, Lists(ws2), None,
                       lambda col, i: sheet_ref(tabs["01"], f"{col}{inv['first'] + i}"))
    ws3 = wb.create_sheet(tabs["03"]); grid(ws3, CHK_COLS)
    ws3.sheet_view.showGridLines = False
    ev = {k: sheet_ref(tabs["01"] if k in on_tab1 else tabs["02"], (inv if k in on_tab1 else req)[k])
          for k in EVIDENCE_LABEL}
    where = {k: v[1] for k, v in EVIDENCE_LABEL.items()}
    ck = checklist(ws3, 1, d, Lists(ws3), None, ev, where)
    for ws, sec in [(ws1, inv), (ws2, req), (ws3, ck)]:
        ws.freeze_panes = f"A{sec['hdr'] + 1}"
        protect(ws)
    refs["blank"] = {"sheets": {"inventory": tabs["01"], "requirements": tabs["02"], "checklist": tabs["03"]},
                     "inventory": inv, "requirements": req, "checklist": ck}

    # ---- the worked example, all three sections on one sheet, locked ------
    wx = wb.create_sheet("Worked Example"); grid(wx, INV_COLS, REQ_COLS, CHK_COLS)
    lists = Lists(wx)
    merged(wx, "A1", "K1", d["exampleLabel"], font=F_BANNER, align=WRAP_MID)
    for col in "ABCDEFGHIJK":
        wx[f"{col}1"].fill = FILL_WARN
    wx.row_dimensions[1].height = 30
    ex = d["example"]
    xinv = inventory(wx, 3, d, lists, ex["actions"], editable=False)
    reqs = [{**r, "_name": a["name"]} for r, a in zip(ex["requirements"], ex["actions"])]
    xreq = requirements(wx, xinv["end"] + 3, d, lists, reqs, lambda col, i: f"{col}{xinv['first'] + i}",
                        editable=False)
    xev = {k: (xinv if k in on_tab1 else xreq)[k] for k in EVIDENCE_LABEL}
    xwhere = {k: ("section 1" if v[1] == "Tab 1" else "section 2") for k, v in EVIDENCE_LABEL.items()}
    xck = checklist(wx, xreq["end"] + 3, d, lists, ex["answers"], xev, xwhere, editable=False)
    wx.freeze_panes = "A2"
    protect(wx)
    refs["example"] = {"sheet": "Worked Example", "inventory": xinv, "requirements": xreq, "checklist": xck}

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = rgb("noir")[2:] if ws.title != "Worked Example" else rgb("ember")[2:]
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
    start_tab(start, d)

    wb.properties.title = d["fullTitle"]
    wb.properties.creator = d["credit"].removeprefix("Built by ")
    wb.properties.subject = d["tagline"]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    refs["expected"] = d["expected"]
    refs["checks"] = [{"n": c["n"], "text": c["text"], "critical": c["critical"]} for c in d["checks"]]
    refs["why"] = d["why"]
    refs["noOpenItems"] = d["noOpenItems"]
    REFS.write_text(json.dumps(refs, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return refs


if __name__ == "__main__":
    build()
    print(f"wrote {OUT.relative_to(REPO)} and {REFS.relative_to(REPO)}")
