#!/usr/bin/env python3
"""Build downloads/citation-fit-check.xlsx (behind the download gate; see CLAUDE.md).

Six tabs: Start here, 1 · Condition Map, 2 · Spot Check, 3 · Checklist,
4 · Test Cases, and Worked Example. The four working tabs are blank. The
Worked Example tab stacks the same four sections on one sheet, filled in, and
its formulas point at its own rows, so it computes the brief's acceptance
figures by itself and a reader can change an answer and watch the decision move.

WHERE THE CONTENT COMES FROM. Not from this file. `dump_content.ts` prints the
words, the twelve checks and the worked example out of
src/data/citation-fit-check.ts and src/lib/citationFit.ts, and this script
renders them. The workbook therefore says what the page says.

WHERE THE BRAND COMES FROM. tools/rework-cost-check/brand.json, which that
tool's extract_brand.py reads out of src/styles/ds/theme.css. Reused rather
than copied, so there is one file of workbook colours to keep current.

FORMULA COMPATIBILITY. Only SUM, IF, AND, OR, MIN, COUNTIF, COUNTIFS, COUNTA,
INDEX, MATCH and TEXT. IFS is avoided for the reason the rework tool found:
openpyxl writes it without the `_xlfn.` prefix and LibreOffice then returns
#NAME?. Nested IF needs no prefix.

Every rule below mirrors a function in src/lib/citationFit.ts, named beside it.
verify_xlsx.py recalculates the file in LibreOffice and asserts the figures.

Run:  python3 tools/citation-fit-check/build_xlsx.py
"""
from __future__ import annotations

import json
import math
import shutil
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
BRAND = json.loads((REPO / "tools" / "rework-cost-check" / "brand.json").read_text(encoding="utf-8"))
BRAND_DIR = REPO / "public" / "brand"
BUILD = HERE / ".build"
OUT = REPO / "downloads" / "citation-fit-check.xlsx"
REFS = HERE / "refs.json"

FOOTER = "© The Living Craft · thelivingcraft.ai · Content CC BY 4.0"

K = BRAND["colors"]
DISPLAY = BRAND["fonts"]["display"]["family"]
BODY = BRAND["fonts"]["body"]["family"]


def rgb(name: str) -> str:
    return "FF" + K[name].lstrip("#").upper()


def solid(name: str) -> PatternFill:
    return PatternFill("solid", fgColor=rgb(name))


FILL_HEAD = solid("noir")
FILL_INPUT = solid("accent_quiet")
FILL_PANEL = solid("surface")
FILL_OK, FILL_WARN, FILL_BAD = solid("success_quiet"), solid("warning_quiet"), solid("danger_quiet")

EDGE_INPUT = Border(*[Side(style="thin", color=rgb("accent_ink"))] * 4)
HAIR = Border(bottom=Side(style="thin", color=rgb("border_hair")))
GOLD_RULE = Border(bottom=Side(style="medium", color=rgb("ember")))

# --weight-display is 400 in this system: the serif is set at its regular weight.
F_TITLE = Font(name=DISPLAY, size=26, color=rgb("text_loud"))
F_H1 = Font(name=DISPLAY, size=18, color=rgb("text_loud"))
F_LEAD = Font(name=BODY, size=12, color=rgb("text_quiet"))
F_LABEL = Font(name=BODY, size=9, bold=True, color=rgb("accent_ink"))
F_HEAD = Font(name=BODY, size=10, bold=True, color=rgb("text_on_invert"))
F_BODY = Font(name=BODY, size=11, color=rgb("text_body"))
F_BODY_B = Font(name=BODY, size=11, bold=True, color=rgb("text_body"))
F_NOTE = Font(name=BODY, size=9, color=rgb("text_quiet"))
F_GHOST = Font(name=BODY, size=10, italic=True, color=rgb("text_quiet"))
F_BIG = Font(name=BODY, size=16, bold=True, color=rgb("text_loud"))
F_GROUP = Font(name=BODY, size=10, bold=True, color=rgb("accent_ink_strong"))
F_BANNER = Font(name=BODY, size=11, bold=True, color=rgb("warning"))

WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
WRAP_MID = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")
INPUT_TEXT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
HEAD_ALIGN = Alignment(horizontal="left", vertical="bottom", wrap_text=True)

PCT = "0.0%"

# One column grid for every working sheet, so the Worked Example can stack all
# four sections and each still lines up.
COLS = {"A": 6, "B": 38, "C": 26, "D": 24, "E": 24, "F": 28, "G": 26, "H": 26}
LIST_COL = "Z"    # the dropdown options, hidden
PRIO_COL = "Y"    # the checklist's next-step priority, hidden


def content() -> dict:
    """The words and the worked example, from the TypeScript that owns them."""
    npx = shutil.which("npx") or "npx"
    out = subprocess.run(
        [npx, "-y", "tsx", str(HERE / "dump_content.ts")],
        capture_output=True, text=True, encoding="utf-8", cwd=REPO, check=False,
    )
    if out.returncode != 0:
        raise SystemExit(f"dump_content.ts failed:\n{out.stderr[-2000:]}")
    return json.loads(out.stdout)


# ---------------------------------------------------------------------------
# Cell helpers
# ---------------------------------------------------------------------------


def put(ws, ref, value=None, *, font=F_BODY, fill=None, align=LEFT, fmt=None, editable=False, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font
    c.alignment = align
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if border is not None:
        c.border = border
    # Sheet protection is a guide rail, not security: calculated cells stay
    # locked so a reader does not type over a formula by accident.
    c.protection = Protection(locked=not editable)
    return c


def inp(ws, ref, value=None, *, align=INPUT_TEXT):
    return put(ws, ref, value if value not in ("", None) else None, fill=FILL_INPUT, align=align,
               editable=True, border=EDGE_INPUT)


def calc(ws, ref, formula, *, fmt=None, font=F_BODY_B, align=LEFT):
    return put(ws, ref, formula, font=font, align=align, fmt=fmt)


def merged(ws, a, b, text, *, font=F_BODY, align=WRAP):
    ws.merge_cells(f"{a}:{b}")
    return put(ws, a, text, font=font, align=align)


def q(s: str) -> str:
    """A string literal inside a formula."""
    return '"' + s.replace('"', '""') + '"'


def fit(ws, row: int, texts: list[tuple[str, float]], *, floor: float = 20, line: float = 14.5):
    """Set a row height from its longest wrapped text.

    Excel does not grow a row to fit wrapped text in a file it did not write, so
    a row left at the default height shows one line of a three-line cell. The
    estimate is generous on purpose: the faces are Figtree and Source Serif 4,
    and on a machine without them Excel draws a wider substitute.
    """
    lines = 1
    for text, width in texts:
        if not text:
            continue
        per_line = max(1.0, width * 1.05)
        lines = max(lines, sum(math.ceil(max(1, len(part)) / per_line) for part in str(text).split("\n")))
    ws.row_dimensions[row].height = max(floor, lines * line + 6)


class Lists:
    """Dropdown options written into a hidden column, one block per list.

    An inline list cannot hold "Same document, other section": its comma is the
    inline list's own separator, so the option would silently split in two. A
    range has no such rule and survives the Google Sheets and LibreOffice
    conversions.
    """

    def __init__(self, ws: Worksheet):
        self.ws, self.next, self.made = ws, 1, {}
        ws.column_dimensions[LIST_COL].hidden = True

    def validation(self, options: list[str]) -> DataValidation:
        key = tuple(options)
        if key not in self.made:
            top = self.next
            for i, option in enumerate(options):
                put(self.ws, f"{LIST_COL}{top + i}", option)
            self.next = top + len(options) + 1
            dv = DataValidation(type="list", formula1=f"${LIST_COL}${top}:${LIST_COL}${top + len(options) - 1}",
                                allow_blank=True, showDropDown=False)
            dv.error, dv.errorTitle = "Pick one of the options in the list.", "Not an option"
            self.ws.add_data_validation(dv)
            self.made[key] = dv
        return self.made[key]


def header_row(ws, row: int, heads: list[str]):
    for i, h in enumerate(heads):
        put(ws, f"{chr(65 + i)}{row}", h, font=F_HEAD, fill=FILL_HEAD, align=HEAD_ALIGN)
    ws.row_dimensions[row].height = 32


def section_head(ws, top: int, tab: dict, note: str):
    put(ws, f"A{top}", f"{int(tab['n'])} · {tab['name']}", font=F_H1)
    ws.row_dimensions[top].height = 28
    merged(ws, f"A{top + 1}", f"H{top + 1}", tab["ask"], font=F_LEAD, align=WRAP_MID)
    ws.row_dimensions[top + 1].height = 20
    merged(ws, f"A{top + 2}", f"H{top + 2}", note, font=F_NOTE, align=WRAP_MID)
    ws.row_dimensions[top + 2].height = 28


def summary(ws, row: int, label: str, formula: str, *, fmt=None, note: str | None = None):
    merged(ws, f"A{row}", f"B{row}", label, font=F_BODY, align=WRAP_MID)
    calc(ws, f"C{row}", formula, fmt=fmt, font=F_BIG)
    ws[f"C{row}"].border = HAIR
    if note:
        merged(ws, f"D{row}", f"H{row}", note, font=F_NOTE, align=WRAP_MID)
    ws.row_dimensions[row].height = 24


def danger_when_positive(ws, cells: str, first: str):
    ws.conditional_formatting.add(cells, FormulaRule(
        formula=[f"{first}>0"],   # the guarded cells are counts, never blank
        fill=FILL_BAD, font=Font(name=BODY, bold=True, color=rgb("danger"))))


def protect(ws: Worksheet):
    ws.protection.sheet = True          # no password: a guide rail, not security
    ws.protection.formatColumns = False  # a reader may still widen a column
    ws.protection.formatRows = False


def grid(ws: Worksheet):
    for col, w in COLS.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions[PRIO_COL].hidden = True


# ---------------------------------------------------------------------------
# The four sections. Each writes at row `top` and returns the cells it owns.
# ---------------------------------------------------------------------------


def condition_map(ws, top: int, d: dict, lists: Lists, rows: list[dict] | None) -> dict:
    tab = d["tabs"][0]
    section_head(ws, top, tab, "One row per rule. Write the condition that narrows it, where that condition "
                               "is stored, and which fact about the asker decides it.")
    r_rules, r_away, r_cant = top + 4, top + 5, top + 6
    hdr = top + 8
    first, last = hdr + 1, hdr + d["rows"]["conditionMap"]

    # conditionMapSummary() in src/lib/citationFit.ts
    summary(ws, r_rules, "Rules mapped", f"=COUNTA(B{first}:B{last})")
    summary(ws, r_away, "Conditions stored away from their rule",
            f'=COUNTA(F{first}:F{last})-COUNTIF(F{first}:F{last},"Same passage")',
            note="Anything not in the rule's own passage. Retrieval can return the rule without it.")
    summary(ws, r_cant, "Conditions the assistant can't check",
            f'=COUNTIF(H{first}:H{last},"Not available")',
            note="The assistant has no way to learn this fact about the asker, so it cannot apply the condition.")
    danger_when_positive(ws, f"C{r_cant}", f"C{r_cant}")

    header_row(ws, hdr, ["#", "Rule", "Source and section", "Condition", "Condition type",
                         "Where the condition lives", "Asker fact needed", "Where the assistant gets that fact"])
    dv_type = lists.validation(d["options"]["conditionTypes"])
    dv_lives = lists.validation(d["options"]["conditionLives"])
    dv_src = lists.validation(d["options"]["factSources"])
    for i in range(d["rows"]["conditionMap"]):
        r = first + i
        row = rows[i] if rows and i < len(rows) else {}
        put(ws, f"A{r}", i + 1, font=F_NOTE, align=CENTER, border=HAIR)
        for col, key in zip("BCDEFGH", ["rule", "source", "condition", "conditionType", "lives",
                                        "askerFact", "factSource"]):
            inp(ws, f"{col}{r}", row.get(key))
        dv_type.add(f"E{r}"); dv_lives.add(f"F{r}"); dv_src.add(f"H{r}")
        fit(ws, r, [(row.get(k, ""), COLS[c]) for c, k in
                    zip("BCDEFGH", ["rule", "source", "condition", "conditionType", "lives", "askerFact",
                                    "factSource"])], floor=30)
    return {"hdr": hdr, "end": last, "rulesMapped": f"C{r_rules}", "storedAway": f"C{r_away}",
            "cannotCheck": f"C{r_cant}"}


def spot_check(ws, top: int, d: dict, lists: Lists, rows: list[dict] | None) -> dict:
    tab = d["tabs"][1]
    section_head(ws, top, tab, "Ten questions of one type, each with the facts about the asker that decide it. "
                               "Read the retrieved passages and the answer from a trace, or run the question.")
    r_recall, r_wrong, r_acc = top + 4, top + 5, top + 6
    hdr = top + 8
    first, last = hdr + 1, hdr + d["rows"]["spotCheck"]
    D, E, F, G = (f"{c}{first}:{c}{last}" for c in "DEFG")

    # spotCheckSummary() in src/lib/citationFit.ts. A ratio with nothing under
    # it is blank, never 0%: "recall 0%" and "nobody measured it" differ.
    summary(ws, r_recall, "Condition recall",
            f'=IF(COUNTIF({D},"Yes")=0,"",COUNTIFS({D},"Yes",{E},"Yes")/COUNTIF({D},"Yes"))', fmt=PCT,
            note="Of the questions where the rule came back, the share where its condition came back too.")
    summary(ws, r_wrong, "Cited but wrong", f'=COUNTIFS({G},"Yes",{F},"No")',
            note="Answers a reader would trust and act on.")
    danger_when_positive(ws, f"C{r_wrong}:H{r_wrong}", f"$C${r_wrong}")
    summary(ws, r_acc, "Answer accuracy for the asker",
            f'=IF(COUNTA({F})=0,"",COUNTIF({F},"Yes")/COUNTA({F}))', fmt=PCT,
            note="Correct for this asker, out of the rows where you recorded whether it was correct.")

    header_row(ws, hdr, ["#", "Question", "Asker facts", "Rule retrieved?", "Condition retrieved?",
                         "Answer correct for this asker?", "Answer cited a source?"])
    dv = lists.validation(d["options"]["yesNo"])
    for i in range(d["rows"]["spotCheck"]):
        r = first + i
        row = rows[i] if rows and i < len(rows) else {}
        put(ws, f"A{r}", i + 1, font=F_NOTE, align=CENTER, border=HAIR)
        inp(ws, f"B{r}", row.get("question"))
        inp(ws, f"C{r}", row.get("askerFacts"))
        for col, key in zip("DEFG", ["ruleRetrieved", "conditionRetrieved", "correctForAsker", "cited"]):
            inp(ws, f"{col}{r}", row.get(key), align=CENTER)
            dv.add(f"{col}{r}")
        fit(ws, r, [(row.get("question", ""), COLS["B"]), (row.get("askerFacts", ""), COLS["C"])], floor=30)
    return {"hdr": hdr, "end": last, "conditionRecall": f"C{r_recall}", "citedButWrong": f"C{r_wrong}",
            "accuracy": f"C{r_acc}"}


def checklist(ws, top: int, d: dict, lists: Lists, answers: list[str] | None, ev: dict, where: dict) -> dict:
    """`ev` maps an evidence key to the cell holding it; `where` names the tab it came from."""
    tab = d["tabs"][2]
    section_head(ws, top, tab, "Answer each check from evidence, not from memory. The four marked ★ are "
                               "critical: a No on any one of them is a Hold.")
    r_dec, r_total, r_answered, r_next = top + 4, top + 5, top + 6, top + 7
    hdr = top + 9

    # Lay out the rows: a group heading, then its checks.
    row_of: dict[int, int] = {}
    r = hdr + 1
    for group in d["checkGroups"]:
        merged(ws, f"A{r}", f"H{r}", group, font=F_GROUP, align=LEFT)
        for col in "ABCDEFGH":
            ws[f"{col}{r}"].fill = FILL_PANEL
        ws.row_dimensions[r].height = 22
        r += 1
        for c in [c for c in d["checks"] if c["group"] == group]:
            row_of[c["n"]] = r
            r += 1
    first, last = hdr + 1, r - 1
    # The check text spans B:D; the answer, points, status and evidence sit in E to H.
    C, Dp, Y = f"E{first}:E{last}", f"F{first}:F{last}", f"{PRIO_COL}{first}:{PRIO_COL}{last}"

    dv = lists.validation(d["options"]["answers"])
    hint = {
        "rulesMapped": lambda ref: f'="Rules mapped ({where["rulesMapped"]}): "&{ref}',
        "conditionRecall": lambda ref: f'="Condition recall ({where["conditionRecall"]}): "&IF({ref}="","not measured yet",TEXT({ref},"0.0%"))',
        "cannotCheck": lambda ref: f'="Conditions the assistant can\'t check ({where["cannotCheck"]}): "&{ref}',
        "testCasesWritten": lambda ref: f'="Test cases written ({where["testCasesWritten"]}): "&{ref}',
    }
    for c in d["checks"]:
        r = row_of[c["n"]]
        given = answers[c["n"] - 1] if answers else None
        put(ws, f"A{r}", f"{c['n']} ★" if c["critical"] else c["n"], font=F_BODY_B if c["critical"] else F_NOTE,
            align=CENTER, border=HAIR)
        merged(ws, f"B{r}", f"D{r}", c["text"], align=WRAP_MID)
        for col in "BCD":
            ws[f"{col}{r}"].border = HAIR
        inp(ws, f"E{r}", given, align=CENTER)
        dv.add(f"E{r}")
        # POINTS and STATUS in src/lib/citationFit.ts
        calc(ws, f"F{r}", f'=IF(E{r}="Yes",2,IF(E{r}="Partial",1,IF(E{r}="No",0,"")))', align=CENTER)
        calc(ws, f"G{r}", f'=IF(E{r}="Yes","Pass",IF(E{r}="Partial","Partial",IF(E{r}="No","Fail","")))',
             align=CENTER)
        if c.get("evidence"):
            calc(ws, f"H{r}", hint[c["evidence"]](ev[c["evidence"]]), font=F_NOTE, align=WRAP_MID)
        # The next-step priority: 1 critical No, 2 No, 3 Partial, 9 nothing to do.
        calc(ws, f"{PRIO_COL}{r}",
             f'=IF(E{r}="No",{1 if c["critical"] else 2},IF(E{r}="Partial",3,9))', font=F_NOTE)
        for col in "FGH":
            ws[f"{col}{r}"].border = HAIR
        # A row with an evidence hint needs two lines in column H.
        fit(ws, r, [(c["text"], COLS["B"] + COLS["C"] + COLS["D"])], floor=36 if c.get("evidence") else 24)

    status_rules = [("Pass", FILL_OK, "success"), ("Partial", FILL_WARN, "warning"), ("Fail", FILL_BAD, "danger")]
    for word, fill, colour in status_rules:
        ws.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(
            formula=[f'G{first}="{word}"'], fill=fill, font=Font(name=BODY, bold=True, color=rgb(colour))))

    # checklistResult() in src/lib/citationFit.ts
    crit = ",".join(f'E{row_of[c["n"]]}="No"' for c in d["checks"] if c["critical"])
    n = len(d["checks"])
    dec = f"C{r_dec}"
    merged(ws, f"A{r_dec}", f"B{r_dec}", "Decision", font=F_BODY_B, align=WRAP_MID)
    calc(ws, dec, f'=IF(OR({crit}),"Hold",IF(COUNTA({C})<{n},"",IF(SUM({Dp})>={d["readyAt"]},"Ready","Fix first")))',
         font=F_BIG, align=CENTER)
    m = d["decisionMeaning"]
    merged(ws, f"D{r_dec}", f"H{r_dec}",
           f'=IF({dec}="Hold",{q(m["Hold"])},IF({dec}="Fix first",{q(m["Fix first"])},'
           f'IF({dec}="Ready",{q(m["Ready"])},IF(COUNTA({C})=0,"","Answer all {n} checks to see the decision."))))',
           font=F_BODY, align=WRAP_MID)
    ws[f"D{r_dec}"].protection = Protection(locked=True)
    ws.row_dimensions[r_dec].height = 34
    for word, fill, colour in [("Ready", FILL_OK, "success"), ("Fix first", FILL_WARN, "warning"), ("Hold", FILL_BAD, "danger")]:
        ws.conditional_formatting.add(dec, FormulaRule(
            formula=[f'{dec}="{word}"'], fill=fill, font=Font(name=BODY, size=16, bold=True, color=rgb(colour))))

    summary(ws, r_total, "Total", f"=SUM({Dp})", note=f"out of {d['maxTotal']}. Yes scores 2, Partial 1, No 0. "
                                                     f"Ready needs {d['readyAt']} or more and no critical No.")
    summary(ws, r_answered, "Checks answered", f"=COUNTA({C})", note=f"out of {n}")
    merged(ws, f"A{r_next}", f"B{r_next}", "Your next step", font=F_BODY_B, align=WRAP_MID)
    merged(ws, f"C{r_next}", f"H{r_next}",
           f'=IF(MIN({Y})>3,"",INDEX(B{first}:B{last},MATCH(MIN({Y}),{Y},0)))', font=F_BODY, align=WRAP_MID)
    ws.row_dimensions[r_next].height = 34
    ws[f"C{r_next}"].border = HAIR

    header_row(ws, hdr, ["#", "Check", "", "", "Answer", "Points", "Status", "Evidence from the other tabs"])
    ws.merge_cells(f"B{hdr}:D{hdr}")
    return {"hdr": hdr, "end": last, "decision": dec, "total": f"C{r_total}", "answered": f"C{r_answered}",
            "nextStep": f"C{r_next}", "answers": {str(k): f"E{v}" for k, v in row_of.items()},
            "status": {str(k): f"G{v}" for k, v in row_of.items()}}


def test_cases(ws, top: int, d: dict, lists: Lists, rows: list[dict] | None, *, guidance: bool) -> dict:
    tab = d["tabs"][3]
    section_head(ws, top, tab, d["testCasesNote"])
    ws[f"A{top + 2}"].font = F_BODY_B
    r_count = top + 4
    keys = ["question", "askerProfile", "clause", "correctAnswer", "whyMisleads", "expected"]
    heads = ["#", "Question", "Asker profile", "Clause the assistant will cite", "Correct answer for this asker",
             "Why the citation alone misleads", "Expected behaviour"]

    hdr = top + 6
    if guidance:
        put(ws, f"A{hdr}", "Two examples, for guidance. They are not counted.", font=F_LABEL)
        header_row(ws, hdr + 1, heads)
        for i, g in enumerate(d["guidanceCases"]):
            r = hdr + 2 + i
            put(ws, f"A{r}", "e.g.", font=F_GHOST, align=CENTER, border=HAIR)
            for col, key in zip("BCDEFG", keys):
                put(ws, f"{col}{r}", g[key], font=F_GHOST, align=WRAP_MID, border=HAIR)
            fit(ws, r, [(g[k], COLS[c]) for c, k in zip("BCDEFG", keys)], line=13.5)
        hdr = hdr + 2 + len(d["guidanceCases"]) + 1
        put(ws, f"A{hdr - 1}", "Your test cases", font=F_LABEL)
        ws.row_dimensions[hdr - 1].height = 26

    first, last = hdr + 1, hdr + d["rows"]["testCases"]
    summary(ws, r_count, "Test cases written", f"=COUNTA(B{first}:B{last})")   # testCasesWritten()
    header_row(ws, hdr, heads)
    dv = lists.validation(d["options"]["expectedBehaviours"])
    for i in range(d["rows"]["testCases"]):
        r = first + i
        row = rows[i] if rows and i < len(rows) else {}
        put(ws, f"A{r}", i + 1, font=F_NOTE, align=CENTER, border=HAIR)
        for col, key in zip("BCDEFG", keys):
            inp(ws, f"{col}{r}", row.get(key))
        dv.add(f"G{r}")
        fit(ws, r, [(row.get(k, ""), COLS[c]) for c, k in zip("BCDEFG", keys)], floor=30)
    return {"hdr": hdr, "end": last, "testCasesWritten": f"C{r_count}"}


# ---------------------------------------------------------------------------
# Start here
# ---------------------------------------------------------------------------


def brand_lockup(mark_h: int = 44, name_h: int = 22, gap: int = 12):
    """The mark beside the lettering, composited into one PNG at twice its displayed size.

    Excel cannot embed WebP and the repository's brand files are WebP, so they
    are converted at build time rather than a PNG copy being committed and left
    to drift. Neither image is enlarged past its native resolution.
    """
    from PIL import Image as PILImage

    def load(name: str, h: int):
        src = BRAND_DIR / f"{name}.webp"
        if not src.exists():
            raise SystemExit(f"missing brand asset {src.relative_to(REPO)}")
        im = PILImage.open(src).convert("RGBA")
        target = min(h * 2, im.height)
        return im.resize((max(1, round(im.width * target / im.height)), target), PILImage.LANCZOS)

    mark, name = load("lc-mark-keyed", mark_h), load("lc-name", name_h)
    pad = gap * 2
    w, h = mark.width + pad + name.width, max(mark.height, name.height)
    canvas = PILImage.new("RGBA", (w, h), (0, 0, 0, 0))
    canvas.paste(mark, (0, (h - mark.height) // 2), mark)
    canvas.paste(name, (mark.width + pad, (h - name.height) // 2), name)
    BUILD.mkdir(parents=True, exist_ok=True)
    mark.save(BUILD / "lc-mark@2x.png")
    out = BUILD / "lc-lockup@2x.png"
    canvas.save(out)
    display_h = max(mark_h, name_h)
    return out, round(w * display_h / h), display_h


def start_tab(ws: Worksheet, d: dict):
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 3, "B": 24, "C": 62, "D": 34}.items():
        ws.column_dimensions[col].width = w

    png, w, h = brand_lockup()
    img = XlImage(str(png))
    img.width, img.height = w, h
    ws.add_image(img, "B2")
    ws.row_dimensions[2].height = 36

    put(ws, "B4", d["title"], font=F_TITLE)
    ws.row_dimensions[4].height = 40
    for col in "BCD":
        ws[f"{col}4"].border = GOLD_RULE
    merged(ws, "B5", "D5", d["tagline"], font=F_LEAD, align=WRAP_MID)
    ws.row_dimensions[5].height = 36

    r = 7
    for label, text in [("What to bring", d["bring"]), ("What to do", d["do"]),
                        ("What you leave with", d["leave"]), ("How to use it", d["howToUse"])]:
        put(ws, f"B{r}", label.upper(), font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
        merged(ws, f"C{r}", f"D{r}", text, align=WRAP)
        fit(ws, r, [(text, 90)], floor=22, line=15)
        r += 2

    put(ws, f"B{r}", "THE FOUR TABS", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    r += 1
    for t in d["tabs"]:
        put(ws, f"B{r}", t["sheet"], font=F_BODY_B)
        merged(ws, f"C{r}", f"D{r}", t["ask"], align=WRAP_MID)
        ws.row_dimensions[r].height = 22
        r += 1
    r += 1

    put(ws, f"B{r}", "CELL LEGEND", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    put(ws, f"C{r}", d["legend"])
    r += 1
    inp(ws, f"C{r}", "Yours to fill", align=LEFT)
    put(ws, f"D{r}", "Shaded, with a thin border.", font=F_NOTE)
    ws.row_dimensions[r].height = 22
    r += 1
    put(ws, f"C{r}", "Calculates for you", font=F_BODY_B)
    put(ws, f"D{r}", "No fill, and locked.", font=F_NOTE)
    ws.row_dimensions[r].height = 22
    r += 2

    merged(ws, f"B{r}", f"D{r}", FOOTER, font=F_NOTE, align=LEFT)
    r += 1
    merged(ws, f"B{r}", f"D{r}", d["credit"], font=F_NOTE, align=LEFT)
    r += 2
    # The cohort invitation and a working application address, which the
    # outreach readiness handoff (28 September 2026) asks every file a reader
    # keeps to carry. The words come from src/data/resource-cohort-copy.ts.
    put(ws, f"B{r}", "THE LIVING CRAFT COHORT", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    merged(ws, f"C{r}", f"D{r}", d["cohort"], align=WRAP)
    fit(ws, r, [(d["cohort"], 90)], floor=22, line=15)
    r += 1
    ws.merge_cells(f"C{r}:D{r}")
    link = put(ws, f"C{r}", f"Apply at {d['applyUrl']}", font=F_BODY_B)
    link.hyperlink = d["applyUrl"]
    protect(ws)



# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------


def sheet_ref(sheet: str, cell: str) -> str:
    return f"'{sheet}'!{cell}"


def build() -> dict:
    d = content()
    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    tabs = {t["n"]: t["sheet"] for t in d["tabs"]}
    refs: dict = {"blank": {}, "example": {}}

    # ---- the four blank working tabs --------------------------------------
    ws1 = wb.create_sheet(tabs["01"]); grid(ws1)
    cm = condition_map(ws1, 1, d, Lists(ws1), None)
    ws2 = wb.create_sheet(tabs["02"]); grid(ws2)
    sc = spot_check(ws2, 1, d, Lists(ws2), None)
    ws4 = wb.create_sheet(tabs["04"]); grid(ws4)
    tc = test_cases(ws4, 1, d, Lists(ws4), None, guidance=True)
    ws3 = wb.create_sheet(tabs["03"], index=3); grid(ws3)
    ws3.sheet_view.showGridLines = False
    ev = {"rulesMapped": sheet_ref(tabs["01"], cm["rulesMapped"]),
          "cannotCheck": sheet_ref(tabs["01"], cm["cannotCheck"]),
          "conditionRecall": sheet_ref(tabs["02"], sc["conditionRecall"]),
          "testCasesWritten": sheet_ref(tabs["04"], tc["testCasesWritten"])}
    where = {"rulesMapped": "Tab 1", "cannotCheck": "Tab 1", "conditionRecall": "Tab 2", "testCasesWritten": "Tab 4"}
    ck = checklist(ws3, 1, d, Lists(ws3), None, ev, where)
    for ws, sec in [(ws1, cm), (ws2, sc), (ws3, ck), (ws4, tc)]:
        ws.freeze_panes = f"A{sec['hdr'] + 1}"
        protect(ws)
    refs["blank"] = {tabs["01"]: cm, tabs["02"]: sc, tabs["03"]: ck, tabs["04"]: tc}

    # ---- the worked example, all four sections on one sheet ----------------
    wx = wb.create_sheet("Worked Example"); grid(wx)
    lists = Lists(wx)
    merged(wx, "A1", "H1", d["exampleLabel"], font=F_BANNER, align=WRAP_MID)
    for col in "ABCDEFGH":
        wx[f"{col}1"].fill = FILL_WARN
    wx.row_dimensions[1].height = 28
    ex = d["example"]
    xcm = condition_map(wx, 3, d, lists, ex["conditions"])
    xsc = spot_check(wx, xcm["end"] + 3, d, lists, ex["spot"])
    # Section 4 is written before section 3 so the checklist can point at its count.
    xtc_top = xsc["end"] + 3 + 30
    xtc = test_cases(wx, xtc_top, d, lists, ex["tests"], guidance=False)
    xev = {"rulesMapped": xcm["rulesMapped"], "cannotCheck": xcm["cannotCheck"],
           "conditionRecall": xsc["conditionRecall"], "testCasesWritten": xtc["testCasesWritten"]}
    xwhere = {k: f"section {v}" for k, v in
              {"rulesMapped": 1, "cannotCheck": 1, "conditionRecall": 2, "testCasesWritten": 4}.items()}
    xck = checklist(wx, xsc["end"] + 3, d, lists, ex["answers"], xev, xwhere)
    if xck["end"] + 3 > xtc_top:
        raise SystemExit("the worked example's checklist runs into its test cases; raise the gap")
    wx.freeze_panes = "A2"
    protect(wx)
    refs["example"] = {"sheet": "Worked Example", "conditionMap": xcm, "spotCheck": xsc, "checklist": xck,
                       "testCases": xtc}

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = rgb("noir")[2:] if ws.title != "Worked Example" else rgb("ember")[2:]
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
    start_tab(start, d)

    wb.properties.title = d["title"]
    wb.properties.creator = d["credit"].removeprefix("Built by ")
    wb.properties.subject = d["tagline"]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    refs["expected"] = d["expected"]
    refs["checks"] = [{"n": c["n"], "text": c["text"], "critical": c["critical"]} for c in d["checks"]]
    REFS.write_text(json.dumps(refs, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return refs


if __name__ == "__main__":
    build()
    print(f"wrote {OUT.relative_to(REPO)} and {REFS.relative_to(REPO)}")
