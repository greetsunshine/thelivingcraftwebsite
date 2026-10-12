#!/usr/bin/env python3
"""Build downloads/gate-routing-map.xlsx (behind the download gate; see CLAUDE.md).

Five tabs: Start here, 1 · Gate Inventory, 2 · Route Requirements,
3 · Checklist, and Worked Example. The three working tabs are blank. The
Worked Example tab stacks the same three sections on one sheet, filled in and
locked, and its formulas point at its own rows, so it computes the brief's
acceptance figures by itself.

WHERE THE CONTENT COMES FROM. Not from this file. `dump_content.ts` prints the
words, the ten checks and the worked example out of src/data/gate-routing-map.ts
and src/lib/gateRouting.ts, and this script renders them. The workbook
therefore says what the page says.

WHERE THE BRAND COMES FROM. tools/rework-cost-check/brand.json, which that
tool's extract_brand.py reads out of src/styles/ds/theme.css. Reused rather
than copied, as the Citation Fit Check does. The four route colours are the
brand's own quiet fills, each with its matching ink, listed in ROUTE_STYLE.

FORMULA COMPATIBILITY. Only SUM, SUMPRODUCT, IF, AND, OR, NOT, MIN, COUNTIF,
COUNTIFS, COUNTA, INDEX and MATCH. IFS is avoided for the reason the rework
tool found: openpyxl writes it without the `_xlfn.` prefix and LibreOffice then
returns #NAME?. Nested IF needs no prefix. A test across two columns that
COUNTIFS cannot express with a "" criterion is a SUMPRODUCT, because Excel and
LibreOffice disagree about whether COUNTIF(range,"") counts a formula's "".

Every rule below mirrors a function in src/lib/gateRouting.ts, named beside it.
verify_xlsx.py recalculates the file in LibreOffice and asserts the figures.

Run:  python3 tools/gate-routing-map/build_xlsx.py
"""
from __future__ import annotations

import json
import math
import shutil
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
BRAND = json.loads((REPO / "tools" / "rework-cost-check" / "brand.json").read_text(encoding="utf-8"))
BRAND_DIR = REPO / "public" / "brand"
BUILD = HERE / ".build"
OUT = REPO / "downloads" / "gate-routing-map.xlsx"
REFS = HERE / "refs.json"

FOOTER = "© The Living Craft · thelivingcraft.ai · Content CC BY 4.0"

K = BRAND["colors"]
DISPLAY = BRAND["fonts"]["display"]["family"]
BODY = BRAND["fonts"]["body"]["family"]


def rgb(name: str) -> str:
    return "FF" + K[name].lstrip("#").upper()


def solid(name: str) -> PatternFill:
    return PatternFill("solid", fgColor=rgb(name))


FILL_HEAD = solid("noir")
FILL_INPUT = solid("accent_quiet")
FILL_PANEL = solid("surface")
FILL_OK, FILL_WARN, FILL_BAD = solid("success_quiet"), solid("warning_quiet"), solid("danger_quiet")
# A Tab 2 column that does not apply to the row's route. The hairline colour,
# darker than every route fill, so a greyed cell never reads as a route.
FILL_OFF = solid("border_hair")

# The four routes: a quiet brand fill and the ink that passes AA on it. Every
# route cell also carries its word, so the colour is never the only signal.
ROUTE_STYLE = {
    "Block": ("danger_quiet", "danger"),
    "Human decides": ("accent_quiet", "accent_ink"),
    "Check downstream": ("info_quiet", "info"),
    "Track": ("surface_hover", "text_quiet"),
}

EDGE_INPUT = Border(*[Side(style="thin", color=rgb("accent_ink"))] * 4)
HAIR = Border(bottom=Side(style="thin", color=rgb("border_hair")))
GOLD_RULE = Border(bottom=Side(style="medium", color=rgb("ember")))

# --weight-display is 400 in this system: the serif is set at its regular weight.
F_TITLE = Font(name=DISPLAY, size=26, color=rgb("text_loud"))
F_H1 = Font(name=DISPLAY, size=18, color=rgb("text_loud"))
F_H2 = Font(name=DISPLAY, size=14, color=rgb("text_loud"))
F_LEAD = Font(name=BODY, size=12, color=rgb("text_quiet"))
F_LABEL = Font(name=BODY, size=9, bold=True, color=rgb("accent_ink"))
F_HEAD = Font(name=BODY, size=10, bold=True, color=rgb("text_on_invert"))
F_BODY = Font(name=BODY, size=11, color=rgb("text_body"))
F_BODY_B = Font(name=BODY, size=11, bold=True, color=rgb("text_body"))
F_NOTE = Font(name=BODY, size=9, color=rgb("text_quiet"))
F_BIG = Font(name=BODY, size=16, bold=True, color=rgb("text_loud"))
F_BANNER = Font(name=BODY, size=11, bold=True, color=rgb("warning"))

WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
WRAP_MID = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
INPUT_TEXT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
HEAD_ALIGN = Alignment(horizontal="left", vertical="bottom", wrap_text=True)

# Column widths per section. The Worked Example stacks all three sections on
# one sheet, so it takes the widest of each column.
INV_COLS = {"A": 5, "B": 30, "C": 32, "D": 14, "E": 13, "F": 24, "G": 16, "H": 14, "I": 17, "J": 17,
            "K": 17, "L": 12, "M": 32}
REQ_COLS = {"A": 5, "B": 30, "C": 17, "D": 14, "E": 10, "F": 10, "G": 22, "H": 18, "I": 14, "J": 20,
            "K": 16, "L": 13, "M": 18, "N": 16, "O": 18, "P": 22, "Q": 12, "R": 36}
CHK_COLS = {"A": 7, "B": 34, "C": 26, "D": 24, "E": 13, "F": 10, "G": 12, "H": 48}
LIST_COL = "Z"    # the dropdown options, hidden
PRIO_COL = "Y"    # the checklist's next-step priority, hidden


def content() -> dict:
    """The words and the worked example, from the TypeScript that owns them."""
    npx = shutil.which("npx") or "npx"
    out = subprocess.run(
        [npx, "-y", "tsx", str(HERE / "dump_content.ts")],
        capture_output=True, text=True, encoding="utf-8", cwd=REPO, check=False,
    )
    if out.returncode != 0:
        raise SystemExit(f"dump_content.ts failed:\n{out.stderr[-2000:]}")
    return json.loads(out.stdout)


# ---------------------------------------------------------------------------
# Cell helpers (the Citation Fit Check's, unchanged where they could be)
# ---------------------------------------------------------------------------


def put(ws, ref, value=None, *, font=F_BODY, fill=None, align=LEFT, fmt=None, editable=False, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font
    c.alignment = align
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if border is not None:
        c.border = border
    # Sheet protection is a guide rail, not security: calculated cells stay
    # locked so a reader does not type over a formula by accident.
    c.protection = Protection(locked=not editable)
    return c


def inp(ws, ref, value=None, *, align=INPUT_TEXT, editable=True):
    return put(ws, ref, value if value not in ("", None) else None, fill=FILL_INPUT, align=align,
               editable=editable, border=EDGE_INPUT)


def calc(ws, ref, formula, *, fmt=None, font=F_BODY_B, align=LEFT):
    return put(ws, ref, formula, font=font, align=align, fmt=fmt)


def merged(ws, a, b, text, *, font=F_BODY, align=WRAP):
    ws.merge_cells(f"{a}:{b}")
    return put(ws, a, text, font=font, align=align)


def q(s: str) -> str:
    """A string literal inside a formula."""
    return '"' + s.replace('"', '""') + '"'


def fit(ws, row: int, texts: list[tuple[str, float]], *, floor: float = 20, line: float = 14.5):
    """Set a row height from its longest wrapped text.

    Excel does not grow a row to fit wrapped text in a file it did not write, so
    a row left at the default height shows one line of a three-line cell. The
    estimate is generous on purpose: the faces are Figtree and Source Serif 4,
    and on a machine without them Excel draws a wider substitute.
    """
    lines = 1
    for text, width in texts:
        if not text:
            continue
        per_line = max(1.0, width * 1.05)
        lines = max(lines, sum(math.ceil(max(1, len(part)) / per_line) for part in str(text).split("\n")))
    ws.row_dimensions[row].height = max(floor, lines * line + 6)


class Lists:
    """Dropdown options written into a hidden column, one block per list."""

    def __init__(self, ws: Worksheet):
        self.ws, self.next, self.made = ws, 1, {}
        ws.column_dimensions[LIST_COL].hidden = True

    def validation(self, options: list[str]) -> DataValidation:
        key = tuple(options)
        if key not in self.made:
            top = self.next
            for i, option in enumerate(options):
                put(self.ws, f"{LIST_COL}{top + i}", option)
            self.next = top + len(options) + 1
            dv = DataValidation(type="list", formula1=f"${LIST_COL}${top}:${LIST_COL}${top + len(options) - 1}",
                                allow_blank=True, showDropDown=False)
            dv.error, dv.errorTitle = "Pick one of the options in the list.", "Not an option"
            self.ws.add_data_validation(dv)
            self.made[key] = dv
        return self.made[key]


def header_row(ws, row: int, heads: list[str], *, height: float = 32):
    for i, h in enumerate(heads):
        put(ws, f"{chr(65 + i)}{row}", h, font=F_HEAD, fill=FILL_HEAD, align=HEAD_ALIGN)
    ws.row_dimensions[row].height = height


def section_head(ws, top: int, tab: dict, note: str, last_col: str):
    put(ws, f"A{top}", f"{int(tab['n'])} · {tab['name']}", font=F_H1)
    ws.row_dimensions[top].height = 28
    merged(ws, f"A{top + 1}", f"{last_col}{top + 1}", tab["ask"], font=F_LEAD, align=WRAP_MID)
    ws.row_dimensions[top + 1].height = 20
    merged(ws, f"A{top + 2}", f"{last_col}{top + 2}", note, font=F_NOTE, align=WRAP_MID)
    ws.row_dimensions[top + 2].height = 28


def summary(ws, row: int, label: str, formula: str, *, note: str | None = None, note_to: str = "H"):
    merged(ws, f"A{row}", f"B{row}", label, font=F_BODY, align=WRAP_MID)
    calc(ws, f"C{row}", formula, font=F_BIG)
    ws[f"C{row}"].border = HAIR
    if note:
        merged(ws, f"D{row}", f"{note_to}{row}", note, font=F_NOTE, align=WRAP_MID)
    ws.row_dimensions[row].height = 24


def danger_when_positive(ws, cell: str):
    ws.conditional_formatting.add(cell, FormulaRule(
        formula=[f"{cell}>0"],   # the guarded cells are counts, never blank
        fill=FILL_BAD, font=Font(name=BODY, size=16, bold=True, color=rgb("danger"))))


def route_colours(ws, cells: str, first: str):
    for route, (fill, ink) in ROUTE_STYLE.items():
        ws.conditional_formatting.add(cells, FormulaRule(
            formula=[f'{first}="{route}"'], fill=solid(fill), font=Font(name=BODY, bold=True, color=rgb(ink))))


def paint_route(ws, ref: str, route: str):
    """A fixed route label (headings, the cover's table), painted the way the conditional rule paints a cell."""
    fill, ink = ROUTE_STYLE[route]
    ws[ref].fill = solid(fill)
    ws[ref].font = Font(name=BODY, size=10, bold=True, color=rgb(ink))


def protect(ws: Worksheet):
    ws.protection.sheet = True          # no password: a guide rail, not security
    ws.protection.formatColumns = False  # a reader may still widen a column
    ws.protection.formatRows = False


def grid(ws: Worksheet, *widths: dict):
    cols: dict[str, float] = {}
    for w in widths:
        for col, v in w.items():
            cols[col] = max(cols.get(col, 0), v)
    for col, v in cols.items():
        ws.column_dimensions[col].width = v
    ws.column_dimensions[PRIO_COL].hidden = True


# ---------------------------------------------------------------------------
# The three sections. Each writes at row `top` and returns the cells it owns.
# ---------------------------------------------------------------------------

INV_KEYS = ["name", "checks", "purpose", "severity", "actor", "checkDecides", "deterministic"]


def inventory(ws, top: int, d: dict, lists: Lists, rows: list[dict] | None, *, editable: bool = True) -> dict:
    tab = d["tabs"][0]
    section_head(ws, top, tab, "One row per gate case or eval suite. Columns D to G suggest a route. "
                               "Leave “Route chosen” blank to accept the suggestion, or pick another and "
                               "write down why.", "M")
    hdr = top + 4
    first, last = hdr + 1, hdr + d["rows"]
    header_row(ws, hdr, ["#", "Case or suite", "What it checks", "Purpose", "Severity if it ships broken",
                         "Who acts on the output", "Can a check decide it unambiguously?", "Deterministic?",
                         "Suggested route", "Route chosen", "Final route", "Differs from suggestion?",
                         "Reason for the difference"], height=46)
    o = d["options"]
    dv = {"D": lists.validation(o["purposes"]), "E": lists.validation(o["severities"]),
          "F": lists.validation(o["actors"]), "G": lists.validation(o["yesNo"]),
          "H": lists.validation(o["yesNo"]), "J": lists.validation(o["routes"])}
    for i in range(d["rows"]):
        r = first + i
        row = rows[i] if rows and i < len(rows) else {}
        put(ws, f"A{r}", i + 1, font=F_NOTE, align=CENTER, border=HAIR)
        for col, key in zip("BCDEFGH", INV_KEYS):
            inp(ws, f"{col}{r}", row.get(key), editable=editable,
                align=INPUT_TEXT if col in "BCF" else CENTER)
        inp(ws, f"J{r}", row.get("chosen"), align=CENTER, editable=editable)
        inp(ws, f"M{r}", row.get("reason"), editable=editable)
        for col, v in dv.items():
            v.add(f"{col}{r}")
        # suggestedRoute() in src/lib/gateRouting.ts, in the brief's order.
        calc(ws, f"I{r}",
             f'=IF(OR(B{r}="",D{r}="",E{r}="",F{r}="",G{r}=""),"",'
             f'IF(D{r}="Capability","Track",'
             f'IF(F{r}="Another system with its own pipeline","Check downstream",'
             f'IF(E{r}="Low","Track",IF(G{r}="No","Human decides","Block")))))', align=CENTER)
        calc(ws, f"K{r}", f'=IF(J{r}<>"",J{r},I{r})', align=CENTER)                 # finalRoute()
        calc(ws, f"L{r}", f'=IF(AND(J{r}<>"",J{r}<>I{r}),"Yes","")', align=CENTER)  # differs()
        for col in "IKL":
            ws[f"{col}{r}"].border = HAIR
        fit(ws, r, [(row.get(k, ""), INV_COLS[c]) for c, k in zip("BCF", ["name", "checks", "actor"])]
            + [(row.get("reason", ""), INV_COLS["M"])], floor=30)
    route_colours(ws, f"I{first}:I{last}", f"I{first}")
    route_colours(ws, f"K{first}:K{last}", f"K{first}")

    # inventorySummary() in src/lib/gateRouting.ts. Below the table, as the brief asks.
    B, I_, Kc, L, M = (f"{c}{first}:{c}{last}" for c in "BIKLM")
    s = last + 2
    put(ws, f"A{s}", "Summary", font=F_H2)
    ws.row_dimensions[s].height = 26
    summary(ws, s + 1, "Cases listed", f"=COUNTA({B})", note_to="M")
    put(ws, f"A{s + 3}", "CASES BY ROUTE", font=F_LABEL)
    for col, h in zip("BCD", ["Route", "Final", "Suggested"]):
        put(ws, f"{col}{s + 3}", h if col != "B" else None, font=F_LABEL, align=CENTER if col != "B" else LEFT)
    by_route, suggested = {}, {}
    for j, route in enumerate(d["options"]["routes"]):
        r = s + 4 + j
        put(ws, f"B{r}", route, border=HAIR)
        paint_route(ws, f"B{r}", route)
        calc(ws, f"C{r}", f'=COUNTIF({Kc},"{route}")', align=CENTER)
        calc(ws, f"D{r}", f'=COUNTIF({I_},"{route}")', font=F_NOTE, align=CENTER)
        by_route[route], suggested[route] = f"C{r}", f"D{r}"
        ws.row_dimensions[r].height = 22
    after = s + 4 + len(d["options"]["routes"]) + 1
    summary(ws, after, "Unrouted cases", f'=SUMPRODUCT(({B}<>"")*({Kc}=""))',
            note="Cases with a name and no route. Fill in D to G, or pick a route.", note_to="M")
    danger_when_positive(ws, f"C{after}")
    summary(ws, after + 1, "Overrides", f'=COUNTIF({L},"Yes")',
            note="Cases where the route chosen differs from the suggestion.", note_to="M")
    summary(ws, after + 2, "Overrides with no reason", f'=SUMPRODUCT(({L}="Yes")*({M}=""))',
            note="Write the reason in column M.", note_to="M")
    danger_when_positive(ws, f"C{after + 2}")
    return {"hdr": hdr, "first": first, "last": last, "end": after + 2,
            "listed": f"C{s + 1}", "byRoute": by_route, "suggested": suggested,
            "unrouted": f"C{after}", "overrides": f"C{after + 1}", "overridesNoReason": f"C{after + 2}",
            "suggestedCells": [f"I{first + i}" for i in range(d["rows"])],
            "finalCells": [f"K{first + i}" for i in range(d["rows"])]}


# Tab 2: which columns belong to which route (ROUTE_FIELDS in gateRouting.ts).
REQ_GROUPS = [
    ("Block", "EF", ["Trials per release", "Trials that must pass"]),
    ("Human decides", "GHI", ["Owner", "Backup owner", "Halt criteria written before the release?"]),
    ("Check downstream", "JKLMNO", ["Downstream system", "Team that owns its pipeline", "Output contract exists?",
                                    "Downstream verifies the contract before either side deploys?",
                                    "Both sides can switch the new behaviour on together?",
                                    "Who can halt the release"]),
    ("Track", "P", ["Owner who reviews the trend"]),
]
REQ_KEYS = {"E": "trials", "F": "mustPass", "G": "owner", "H": "backup", "I": "haltCriteria", "J": "system",
            "K": "team", "L": "contract", "M": "verifies", "N": "switchTogether", "O": "halter", "P": "trackOwner"}


def requirements(ws, top: int, d: dict, lists: Lists, rows: list[dict] | None, link, *,
                 min_trials, editable: bool = True) -> dict:
    """`link(col, i)` is the Tab 1 cell for column `col` of case i (0-based)."""
    tab = d["tabs"][1]
    section_head(ws, top, tab, "Rows mirror the Gate Inventory. Fill in only the columns for the row’s final "
                               "route; the others are greyed out. Each row names the first thing missing.", "R")
    r_min = top + 4
    merged(ws, f"A{r_min}", f"B{r_min}", d["minTrialsLabel"], font=F_BODY_B, align=WRAP_MID)
    inp(ws, f"C{r_min}", min_trials, align=CENTER, editable=editable)
    mt = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=False)
    mt.error, mt.errorTitle = "Enter a whole number of 1 or more.", "Minimum trials"
    ws.add_data_validation(mt)
    mt.add(f"C{r_min}")
    merged(ws, f"D{r_min}", f"I{r_min}", d["minTrialsNote"], font=F_NOTE, align=WRAP_MID)
    ws.row_dimensions[r_min].height = 34
    # Delete clears a cell without passing its validation. A blank minimum
    # then reads as the default, not as "" (which every number is less than).
    MIN = f'IF($C${r_min}="",{d["defaultMinTrials"]},$C${r_min})'

    grp = top + 6
    hdr = grp + 1
    first, last = hdr + 1, hdr + d["rows"]
    for route, cols, _ in REQ_GROUPS:
        a, b = f"{cols[0]}{grp}", f"{cols[-1]}{grp}"
        if a != b:
            ws.merge_cells(f"{a}:{b}")
        put(ws, a, route)
        paint_route(ws, a, route)
    ws.row_dimensions[grp].height = 20
    heads = ["#", "Case or suite", "Final route", "Deterministic?"]
    for _, _, h in REQ_GROUPS:
        heads += h
    header_row(ws, hdr, heads + ["Route ready?", "Why not ready"], height=58)

    o = d["options"]
    dv_yn, dv_ypn = lists.validation(o["yesNo"]), lists.validation(o["answers"])
    whole = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    whole.error, whole.errorTitle = "Enter a whole number of trials.", "Trials"
    ws.add_data_validation(whole)
    # A text column that does not apply to the row's route refuses input. The
    # brief asks for those columns locked; protection cannot depend on a
    # formula, so validation does the refusing. A dropdown cell already holds
    # its list validation, and a cell takes only one, so those are greyed only.
    refuse = {}
    for route, cols, _ in REQ_GROUPS:
        if route == "Block":
            continue
        v = DataValidation(type="custom", formula1=f'$C{first}="{route}"', allow_blank=True)
        v.error = f"This column is for {route} cases. Change the route on the Gate Inventory first."
        v.errorTitle = "Not this route"
        ws.add_data_validation(v)
        refuse[route] = v

    W = d["why"]
    for i in range(d["rows"]):
        r = first + i
        row = rows[i] if rows and i < len(rows) else {}
        put(ws, f"A{r}", i + 1, font=F_NOTE, align=CENTER, border=HAIR)
        for col, key in zip("BCD", ["name", "final", "deterministic"]):
            src = link({"name": "B", "final": "K", "deterministic": "H"}[key], i)
            calc(ws, f"{col}{r}", f'=IF({src}="","",{src})', font=F_BODY if col == "B" else F_BODY_B,
                 align=WRAP_MID if col == "B" else CENTER)
            ws[f"{col}{r}"].border = HAIR
        for route, cols, _ in REQ_GROUPS:
            for col in cols:
                key = REQ_KEYS[col]
                val = row.get(key)
                inp(ws, f"{col}{r}", val, editable=editable,
                    align=CENTER if col in "EFILMN" else INPUT_TEXT)
                if col in "EF":
                    whole.add(f"{col}{r}")
                elif col == "I" or col in "MN":
                    dv_yn.add(f"{col}{r}")
                elif col == "L":
                    dv_ypn.add(f"{col}{r}")
                else:
                    refuse[route].add(f"{col}{r}")

        # routeReadiness() in src/lib/gateRouting.ts. "Why not ready" names the
        # first missing item; "Route ready?" is Ready exactly when it is blank,
        # so the two columns cannot disagree.
        C, D, E, F, G, I_, J, Kc, L, M, N, O, P = (f"{c}{r}" for c in "CDEFGIJKLMNOP")
        block = (f'IF(OR({E}="",{E}<1),{q(W["noTrials"])},IF({F}="",{q(W["mustPassBlank"])},'
                 f'IF({F}<{E},{q(W["mustPassLower"])},IF({F}>{E},{q(W["mustPassHigher"])},'
                 f'IF(AND({D}<>"Yes",{E}<{MIN}),{q(W["belowMinimum"])},"")))))')
        human = f'IF({G}="",{q(W["noOwner"])},IF({I_}<>"Yes",{q(W["noHalt"])},""))'
        down = (f'IF({J}="",{q(W["noSystem"])},IF({Kc}="",{q(W["noTeam"])},'
                f'IF({L}="Partial",{q(W["partialContract"])},IF({L}<>"Yes",{q(W["noContract"])},'
                f'IF({M}<>"Yes",{q(W["notVerified"])},IF({N}<>"Yes",{q(W["noSwitch"])},'
                f'IF({O}="",{q(W["noHalter"])},"")))))))')
        track = f'IF({P}="",{q(W["noOwner"])},"")'
        calc(ws, f"R{r}", f'=IF({C}="","",IF({C}="Block",{block},IF({C}="Human decides",{human},'
                          f'IF({C}="Check downstream",{down},IF({C}="Track",{track},"")))))',
             font=F_BODY, align=WRAP_MID)
        calc(ws, f"Q{r}", f'=IF({C}="","",IF(R{r}="","Ready","Not ready"))', align=CENTER)
        for col in "QR":
            ws[f"{col}{r}"].border = HAIR
        fit(ws, r, [(row.get("_name", ""), REQ_COLS["B"])]
            + [(str(row.get(REQ_KEYS[c], "") or ""), REQ_COLS[c]) for c in "GHJKOP"], floor=30)

    for route, cols, _ in REQ_GROUPS:
        rng = f"{cols[0]}{first}:{cols[-1]}{last}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'$C{first}<>"{route}"'], fill=FILL_OFF, font=Font(name=BODY, color=rgb("text_quiet"))))
    route_colours(ws, f"C{first}:C{last}", f"C{first}")
    for word, fill, colour in [("Ready", FILL_OK, "success"), ("Not ready", FILL_BAD, "danger")]:
        ws.conditional_formatting.add(f"Q{first}:Q{last}", FormulaRule(
            formula=[f'Q{first}="{word}"'], fill=fill, font=Font(name=BODY, bold=True, color=rgb(colour))))

    # requirementsSummary() in src/lib/gateRouting.ts
    Cr, H, Lr, Mr, Q = (f"{c}{first}:{c}{last}" for c in "CHLMQ")
    s = last + 2
    put(ws, f"A{s}", "Summary", font=F_H2)
    ws.row_dimensions[s].height = 26
    for col, h in zip("BCD", ["Route", "Cases", "Ready"]):
        put(ws, f"{col}{s + 1}", h if col != "B" else None, font=F_LABEL, align=CENTER)
    put(ws, f"A{s + 1}", "BY ROUTE", font=F_LABEL)
    cases, ready = {}, {}
    for j, route in enumerate(d["options"]["routes"]):
        r = s + 2 + j
        put(ws, f"B{r}", route, border=HAIR)
        paint_route(ws, f"B{r}", route)
        calc(ws, f"C{r}", f'=COUNTIF({Cr},"{route}")', align=CENTER)
        calc(ws, f"D{r}", f'=COUNTIFS({Cr},"{route}",{Q},"Ready")', align=CENTER)
        cases[route], ready[route] = f"C{r}", f"D{r}"
        ws.row_dimensions[r].height = 22
    a = s + 2 + len(d["options"]["routes"]) + 1
    summary(ws, a, "Routes not ready", f'=COUNTIF({Q},"Not ready")',
            note="Column R names the first thing missing on each.", note_to="I")
    danger_when_positive(ws, f"C{a}")
    summary(ws, a + 1, "Block cases not ready", f'=COUNTIFS({Cr},"Block",{Q},"Not ready")', note_to="I")
    summary(ws, a + 2, "Human cases not ready", f'=COUNTIFS({Cr},"Human decides",{Q},"Not ready")', note_to="I")
    summary(ws, a + 3, "Human cases with no backup owner", f'=SUMPRODUCT(({Cr}="Human decides")*({H}=""))',
            note="Scored by check 6, not by Route ready.", note_to="I")
    summary(ws, a + 4, "Downstream cases without a verified contract",
            f'=SUMPRODUCT(({Cr}="Check downstream")*((({Lr}<>"Yes")+({Mr}<>"Yes"))>0))',
            note="The contract is not Yes, or the downstream side does not verify it before deploying.", note_to="I")
    return {"hdr": hdr, "first": first, "last": last, "end": a + 4, "minTrials": f"C{r_min}",
            "cases": cases, "ready": ready, "notReady": f"C{a}", "blockNotReady": f"C{a + 1}",
            "humanNotReady": f"C{a + 2}", "humanNoBackup": f"C{a + 3}", "downstreamUnverified": f"C{a + 4}",
            "readyCells": [f"Q{first + i}" for i in range(d["rows"])],
            "whyCells": [f"R{first + i}" for i in range(d["rows"])],
            "inputs": {k: [f"{c}{first + i}" for i in range(d["rows"])] for c, k in REQ_KEYS.items()}}


EVIDENCE_LABEL = {
    "unrouted": ("Unrouted cases", "Tab 1"),
    "overridesNoReason": ("Overrides with no reason", "Tab 1"),
    "blockNotReady": ("Block cases not ready", "Tab 2"),
    "humanNotReady": ("Human cases not ready", "Tab 2"),
    "humanNoBackup": ("Human cases with no backup owner", "Tab 2"),
    "downstreamUnverified": ("Downstream cases without a verified contract", "Tab 2"),
}


def checklist(ws, top: int, d: dict, lists: Lists, answers: list[str] | None, ev: dict, where: dict, *,
              editable: bool = True) -> dict:
    """`ev` maps an evidence key to the cell holding it; `where` names the tab or section it came from."""
    tab = d["tabs"][2]
    section_head(ws, top, tab, "Answer each check from evidence, not from memory. The four marked ★ are "
                               "critical: a No on any one of them is a Hold.", "H")
    r_dec, r_total, r_answered, r_next = top + 4, top + 5, top + 6, top + 7
    hdr = top + 9
    first, last = hdr + 1, hdr + len(d["checks"])
    E, Fp, Y = f"E{first}:E{last}", f"F{first}:F{last}", f"{PRIO_COL}{first}:{PRIO_COL}{last}"

    dv = lists.validation(d["options"]["answers"])
    row_of = {}
    for i, c in enumerate(d["checks"]):
        r = first + i
        row_of[c["n"]] = r
        given = answers[c["n"] - 1] if answers else None
        put(ws, f"A{r}", f"{c['n']} ★" if c["critical"] else c["n"], font=F_BODY_B if c["critical"] else F_NOTE,
            align=CENTER, border=HAIR)
        merged(ws, f"B{r}", f"D{r}", c["text"], align=WRAP_MID)
        for col in "BCD":
            ws[f"{col}{r}"].border = HAIR
        inp(ws, f"E{r}", given, align=CENTER, editable=editable)
        dv.add(f"E{r}")
        # POINTS and STATUS in src/lib/gateRouting.ts
        calc(ws, f"F{r}", f'=IF(E{r}="Yes",2,IF(E{r}="Partial",1,IF(E{r}="No",0,"")))', align=CENTER)
        calc(ws, f"G{r}", f'=IF(E{r}="Yes","Pass",IF(E{r}="Partial","Partial",IF(E{r}="No","Fail","")))',
             align=CENTER)
        if c.get("evidence"):
            label, _ = EVIDENCE_LABEL[c["evidence"]]
            calc(ws, f"H{r}", f'="{label} ({where[c["evidence"]]}): "&{ev[c["evidence"]]}',
                 font=F_NOTE, align=WRAP_MID)
        # priority() in gateRouting.ts: 1 critical No, 2 No, 3 critical Partial, 4 Partial, 9 nothing.
        calc(ws, f"{PRIO_COL}{r}",
             f'=IF(E{r}="No",{1 if c["critical"] else 2},IF(E{r}="Partial",{3 if c["critical"] else 4},9))',
             font=F_NOTE)
        for col in "FGH":
            ws[f"{col}{r}"].border = HAIR
        fit(ws, r, [(c["text"], CHK_COLS["B"] + CHK_COLS["C"] + CHK_COLS["D"])], floor=36 if c.get("evidence") else 26)

    for word, fill, colour in [("Pass", FILL_OK, "success"), ("Partial", FILL_WARN, "warning"), ("Fail", FILL_BAD, "danger")]:
        ws.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(
            formula=[f'G{first}="{word}"'], fill=fill, font=Font(name=BODY, bold=True, color=rgb(colour))))

    # checklistResult() in src/lib/gateRouting.ts
    crit = [c for c in d["checks"] if c["critical"]]
    crit_no = ",".join(f'E{row_of[c["n"]]}="No"' for c in crit)
    crit_partial = ",".join(f'E{row_of[c["n"]]}="Partial"' for c in crit)
    n = len(d["checks"])
    dec = f"C{r_dec}"
    merged(ws, f"A{r_dec}", f"B{r_dec}", "Decision", font=F_BODY_B, align=WRAP_MID)
    calc(ws, dec, f'=IF(OR({crit_no}),"Hold",IF(COUNTA({E})<{n},"",'
                  f'IF(AND(SUM({Fp})>={d["readyAt"]},NOT(OR({crit_partial}))),"Ready","Fix first")))',
         font=F_BIG, align=CENTER)
    m = d["decisionMeaning"]
    merged(ws, f"D{r_dec}", f"H{r_dec}",
           f'=IF({dec}="Hold",{q(m["Hold"])},IF({dec}="Fix first",{q(m["Fix first"])},'
           f'IF({dec}="Ready",{q(m["Ready"])},IF(COUNTA({E})=0,"","Answer all {n} checks to see the decision."))))',
           font=F_BODY, align=WRAP_MID)
    ws[f"D{r_dec}"].protection = Protection(locked=True)
    ws.row_dimensions[r_dec].height = 34
    for word, fill, colour in [("Ready", FILL_OK, "success"), ("Fix first", FILL_WARN, "warning"), ("Hold", FILL_BAD, "danger")]:
        ws.conditional_formatting.add(dec, FormulaRule(
            formula=[f'{dec}="{word}"'], fill=fill, font=Font(name=BODY, size=16, bold=True, color=rgb(colour))))

    summary(ws, r_total, "Total", f"=SUM({Fp})",
            note=f"out of {d['maxTotal']}. Yes scores 2, Partial 1, No 0. Ready needs {d['readyAt']} or more "
                 f"and every critical check answered Yes.")
    summary(ws, r_answered, "Checks answered", f"=COUNTA({E})", note=f"out of {n}")
    merged(ws, f"A{r_next}", f"B{r_next}", "Your next step", font=F_BODY_B, align=WRAP_MID)
    merged(ws, f"C{r_next}", f"H{r_next}",
           f'=IF(MIN({Y})<9,INDEX(B{first}:B{last},MATCH(MIN({Y}),{Y},0)),'
           f'IF(COUNTA({E})={n},{q(d["noOpenItems"])},""))', font=F_BODY, align=WRAP_MID)
    ws.row_dimensions[r_next].height = 40
    ws[f"C{r_next}"].border = HAIR

    header_row(ws, hdr, ["#", "Check", "", "", "Answer", "Points", "Status", "Evidence from the other tabs"])
    ws.merge_cells(f"B{hdr}:D{hdr}")
    return {"hdr": hdr, "first": first, "last": last, "end": last, "decision": dec, "total": f"C{r_total}",
            "answered": f"C{r_answered}", "nextStep": f"C{r_next}",
            "answers": {str(k): f"E{v}" for k, v in row_of.items()},
            "status": {str(k): f"G{v}" for k, v in row_of.items()},
            "evidence": {str(k): f"H{v}" for k, v in row_of.items()}}


# ---------------------------------------------------------------------------
# Start here
# ---------------------------------------------------------------------------


def brand_lockup(mark_h: int = 44, name_h: int = 22, gap: int = 12):
    """The mark beside the lettering, composited into one PNG at twice its displayed size.

    Excel cannot embed WebP and the repository's brand files are WebP, so they
    are converted at build time rather than a PNG copy being committed and left
    to drift. Neither image is enlarged past its native resolution.
    """
    from PIL import Image as PILImage

    def load(name: str, h: int):
        src = BRAND_DIR / f"{name}.webp"
        if not src.exists():
            raise SystemExit(f"missing brand asset {src.relative_to(REPO)}")
        im = PILImage.open(src).convert("RGBA")
        target = min(h * 2, im.height)
        return im.resize((max(1, round(im.width * target / im.height)), target), PILImage.LANCZOS)

    mark, name = load("lc-mark-keyed", mark_h), load("lc-name", name_h)
    pad = gap * 2
    w, h = mark.width + pad + name.width, max(mark.height, name.height)
    canvas = PILImage.new("RGBA", (w, h), (0, 0, 0, 0))
    canvas.paste(mark, (0, (h - mark.height) // 2), mark)
    canvas.paste(name, (mark.width + pad, (h - name.height) // 2), name)
    BUILD.mkdir(parents=True, exist_ok=True)
    mark.save(BUILD / "lc-mark@2x.png")
    out = BUILD / "lc-lockup@2x.png"
    canvas.save(out)
    display_h = max(mark_h, name_h)
    return out, round(w * display_h / h), display_h


def start_tab(ws: Worksheet, d: dict):
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 3, "B": 22, "C": 46, "D": 64}.items():
        ws.column_dimensions[col].width = w

    png, w, h = brand_lockup()
    img = XlImage(str(png))
    img.width, img.height = w, h
    ws.add_image(img, "B2")
    ws.row_dimensions[2].height = 36

    put(ws, "B4", d["title"], font=F_TITLE)
    ws.row_dimensions[4].height = 40
    for col in "BCD":
        ws[f"{col}4"].border = GOLD_RULE
    merged(ws, "B5", "D5", d["tagline"], font=F_LEAD, align=WRAP_MID)
    ws.row_dimensions[5].height = 26

    r = 7
    for label, text in [("What to bring", d["bring"]), ("What to do", d["do"]),
                        ("What you leave with", d["leave"])]:
        put(ws, f"B{r}", label.upper(), font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
        merged(ws, f"C{r}", f"D{r}", text, align=WRAP)
        fit(ws, r, [(text, 100)], floor=22, line=15)
        r += 2

    put(ws, f"B{r}", "THE FOUR ROUTES", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    r += 1
    for col, h in zip("BCD", ["Route", "Use it when", "What the route needs before the release"]):
        put(ws, f"{col}{r}", h, font=F_HEAD, fill=FILL_HEAD, align=HEAD_ALIGN)
    ws.row_dimensions[r].height = 22
    r += 1
    for row in d["routeTable"]:
        put(ws, f"B{r}", row["route"], align=WRAP_MID, border=HAIR)
        paint_route(ws, f"B{r}", row["route"])
        put(ws, f"C{r}", row["when"], align=WRAP_MID, border=HAIR)
        put(ws, f"D{r}", row["needs"], align=WRAP_MID, border=HAIR)
        fit(ws, r, [(row["when"], 46), (row["needs"], 64)], floor=24, line=15)
        r += 1
    r += 1

    put(ws, f"B{r}", "HOW TO USE IT", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    merged(ws, f"C{r}", f"D{r}", d["howToUse"], align=WRAP)
    r += 1
    for t in d["tabs"]:
        put(ws, f"B{r}", t["sheet"], font=F_BODY_B)
        merged(ws, f"C{r}", f"D{r}", t["ask"], align=WRAP_MID)
        ws.row_dimensions[r].height = 22
        r += 1
    r += 1

    put(ws, f"B{r}", "CELL LEGEND", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    put(ws, f"C{r}", d["legend"])
    r += 1
    inp(ws, f"C{r}", "Yours to fill", align=LEFT, editable=False)
    put(ws, f"D{r}", "Shaded, with a thin border.", font=F_NOTE)
    ws.row_dimensions[r].height = 22
    r += 1
    put(ws, f"C{r}", "Calculates for you", font=F_BODY_B)
    put(ws, f"D{r}", "No fill, and locked.", font=F_NOTE)
    ws.row_dimensions[r].height = 22
    r += 1
    put(ws, f"C{r}", "Does not apply", fill=FILL_OFF, font=Font(name=BODY, color=rgb("text_quiet")))
    put(ws, f"D{r}", "On Tab 2, a column that is not for this row’s route.", font=F_NOTE)
    ws.row_dimensions[r].height = 22
    r += 2

    merged(ws, f"B{r}", f"D{r}", FOOTER, font=F_NOTE, align=LEFT)
    r += 1
    merged(ws, f"B{r}", f"D{r}", d["credit"], font=F_NOTE, align=LEFT)
    r += 2
    # The cohort invitation and a working application address, which the
    # outreach readiness handoff (28 September 2026) asks every file a reader
    # keeps to carry. The words come from src/data/resource-cohort-copy.ts.
    put(ws, f"B{r}", "THE LIVING CRAFT COHORT", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    merged(ws, f"C{r}", f"D{r}", d["cohort"], align=WRAP)
    fit(ws, r, [(d["cohort"], 100)], floor=22, line=15)
    r += 1
    ws.merge_cells(f"C{r}:D{r}")
    link = put(ws, f"C{r}", f"Apply at {d['applyUrl']}", font=F_BODY_B)
    link.hyperlink = d["applyUrl"]
    protect(ws)


# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------


def sheet_ref(sheet: str, cell: str) -> str:
    return f"'{sheet}'!{cell}"


def build() -> dict:
    d = content()
    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    tabs = {t["n"]: t["sheet"] for t in d["tabs"]}
    refs: dict = {"blank": {}, "example": {}}

    # ---- the three blank working tabs -------------------------------------
    ws1 = wb.create_sheet(tabs["01"]); grid(ws1, INV_COLS)
    inv = inventory(ws1, 1, d, Lists(ws1), None)
    ws2 = wb.create_sheet(tabs["02"]); grid(ws2, REQ_COLS)
    req = requirements(ws2, 1, d, Lists(ws2), None,
                       lambda col, i: sheet_ref(tabs["01"], f"{col}{inv['first'] + i}"),
                       min_trials=d["defaultMinTrials"])
    ws3 = wb.create_sheet(tabs["03"]); grid(ws3, CHK_COLS)
    ws3.sheet_view.showGridLines = False
    ev = {k: sheet_ref(tabs["01"] if k in ("unrouted", "overridesNoReason") else tabs["02"],
                       (inv if k in ("unrouted", "overridesNoReason") else req)[k])
          for k in EVIDENCE_LABEL}
    where = {k: v[1] for k, v in EVIDENCE_LABEL.items()}
    ck = checklist(ws3, 1, d, Lists(ws3), None, ev, where)
    for ws, sec in [(ws1, inv), (ws2, req), (ws3, ck)]:
        ws.freeze_panes = f"A{sec['hdr'] + 1}"
        protect(ws)
    refs["blank"] = {"sheets": {"inventory": tabs["01"], "requirements": tabs["02"], "checklist": tabs["03"]},
                     "inventory": inv, "requirements": req, "checklist": ck}

    # ---- the worked example, all three sections on one sheet, locked ------
    wx = wb.create_sheet("Worked Example"); grid(wx, INV_COLS, REQ_COLS, CHK_COLS)
    lists = Lists(wx)
    merged(wx, "A1", "M1", d["exampleLabel"], font=F_BANNER, align=WRAP_MID)
    for col in "ABCDEFGHIJKLM":
        wx[f"{col}1"].fill = FILL_WARN
    wx.row_dimensions[1].height = 30
    ex = d["example"]
    xinv = inventory(wx, 3, d, lists, ex["cases"], editable=False)
    reqs = [{**r, "_name": c["name"]} for r, c in zip(ex["requirements"], ex["cases"])]
    xreq = requirements(wx, xinv["end"] + 3, d, lists, reqs, lambda col, i: f"{col}{xinv['first'] + i}",
                        min_trials=ex["minTrials"], editable=False)
    xev = {k: (xinv if k in ("unrouted", "overridesNoReason") else xreq)[k] for k in EVIDENCE_LABEL}
    xwhere = {k: ("section 1" if v[1] == "Tab 1" else "section 2") for k, v in EVIDENCE_LABEL.items()}
    xck = checklist(wx, xreq["end"] + 3, d, lists, ex["answers"], xev, xwhere, editable=False)
    wx.freeze_panes = "A2"
    protect(wx)
    refs["example"] = {"sheet": "Worked Example", "inventory": xinv, "requirements": xreq, "checklist": xck}

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = rgb("noir")[2:] if ws.title != "Worked Example" else rgb("ember")[2:]
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
    start_tab(start, d)

    wb.properties.title = d["title"]
    wb.properties.creator = d["credit"].removeprefix("Built by ")
    wb.properties.subject = d["tagline"]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    refs["expected"] = d["expected"]
    refs["checks"] = [{"n": c["n"], "text": c["text"], "critical": c["critical"]} for c in d["checks"]]
    refs["why"] = d["why"]
    refs["noOpenItems"] = d["noOpenItems"]
    REFS.write_text(json.dumps(refs, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return refs


if __name__ == "__main__":
    build()
    print(f"wrote {OUT.relative_to(REPO)} and {REFS.relative_to(REPO)}")
