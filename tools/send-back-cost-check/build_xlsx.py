#!/usr/bin/env python3
"""Build public/downloads/send-back-cost-check.xlsx.

Four tabs: Start here, Check, How to find these numbers, Example.

WHERE THE CONTENT COMES FROM. Not from this file. `dump_content.ts` prints the
words and the worked example out of src/data/send-back-cost-check.ts and
src/lib/sendBackCost.ts, and this script renders them. Same pattern as
`npm run workbook`. The workbook therefore cannot drift from the page, and
verify_xlsx.py asserts that the spreadsheet's own formulas reproduce the
figures the TypeScript module computes.

FORMULA COMPATIBILITY. Only SUM SUMPRODUCT IF IFS AND OR MIN MAX ROUND COUNTIF
COUNTIFS INDEX MATCH IFERROR TEXT are used, and IFS is written as nested IF:
openpyxl writes it bare while the file format needs an `_xlfn.` prefix for
anything newer than 2007, and without the prefix LibreOffice returns #NAME?
before IFERROR can catch it. Nested IF needs no prefix and no converter can
mangle it.

Run:  python3 tools/send-back-cost-check/build_xlsx.py
"""
from __future__ import annotations

import json
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
BRAND = json.loads((HERE / "brand.json").read_text(encoding="utf-8"))
BRAND_DIR = REPO / "public" / "brand"
BUILD = HERE / ".build"
OUT = REPO / "public" / "downloads" / "send-back-cost-check.xlsx"

TITLE = "Send-Back Cost Check"
TAGLINE = "What does one task really cost once limits, judges, validators and reviewers send work back?"
FOOTER = "© The Living Craft · thelivingcraft.ai · Content CC BY 4.0 · Code MIT"
LEGEND = "Shaded cells are yours to fill. Everything else calculates."

K = BRAND["colors"]
DISPLAY = BRAND["fonts"]["display"]["family"]
BODY = BRAND["fonts"]["body"]["family"]


def content() -> dict:
    """The words and the worked example, from the TypeScript that owns them."""
    out = subprocess.run(
        ["node", "--experimental-strip-types", str(HERE / "dump_content.ts")],
        capture_output=True, text=True, cwd=REPO, check=False,
    )
    if out.returncode != 0:
        raise SystemExit(f"dump_content.ts failed:\n{out.stderr[-2000:]}")
    return json.loads(out.stdout)


def rgb(name: str) -> str:
    return "FF" + K[name].lstrip("#").upper()


FILL_NOIR = PatternFill("solid", fgColor=rgb("noir"))
FILL_INPUT = PatternFill("solid", fgColor=rgb("accent_quiet"))
FILL_MIST = PatternFill("solid", fgColor=rgb("background"))
FILL_OK = PatternFill("solid", fgColor=rgb("success_quiet"))
FILL_WARN = PatternFill("solid", fgColor=rgb("warning_quiet"))
FILL_BAD = PatternFill("solid", fgColor=rgb("danger_quiet"))

EDGE_INPUT = Border(*[Side(style="thin", color=rgb("accent_ink"))] * 4)
HAIR = Border(bottom=Side(style="thin", color=rgb("border_hair")))

# --weight-display is 400 in this system: the serif is set at its regular weight.
F_TITLE = Font(name=DISPLAY, size=26, color=rgb("text_loud"))
F_H1 = Font(name=DISPLAY, size=18, color=rgb("text_loud"))
F_H2 = Font(name=BODY, size=12, bold=True, color=rgb("text_loud"))
F_LABEL = Font(name=BODY, size=9, bold=True, color=rgb("accent_ink"))
F_HEAD = Font(name=BODY, size=10, bold=True, color=rgb("text_on_invert"))
F_BODY = Font(name=BODY, size=11, color=rgb("text_body"))
F_BODY_B = Font(name=BODY, size=11, bold=True, color=rgb("text_body"))
F_NOTE = Font(name=BODY, size=9, color=rgb("text_quiet"))
F_LEAD = Font(name=BODY, size=12, color=rgb("text_quiet"))
F_NUM = Font(name=BODY, size=11, color=rgb("text_body"))
F_BIG = Font(name=BODY, size=16, bold=True, color=rgb("text_loud"))
F_GHOST = Font(name=BODY, size=11, italic=True, color=rgb("text_quiet"))

WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
WRAP_MID = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")

# Pinned to thousands grouping: a bare #,##0 follows the reader's locale, and
# an India locale renders 10,000,000 as 1,00,00,000. These are token counts.
NUMFMT = "[$-409]#,##0"
PCTFMT = "0.0%"
MULFMT = '0.00"×"'
MONEYFMT = "[$-409]#,##0.00"


def put(ws, ref, value=None, *, font=F_BODY, fill=None, align=LEFT, fmt=None,
        editable=False, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font
    c.alignment = align
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if border is not None:
        c.border = border
    # Sheet protection is a guide rail, not security: calculated cells stay
    # locked so a reader does not type over a formula by accident.
    c.protection = Protection(locked=not editable)
    return c


def inp(ws, ref, value=None, *, fmt=None, align=RIGHT, font=F_NUM):
    return put(ws, ref, value, font=font, fill=FILL_INPUT, align=align, fmt=fmt,
               editable=True, border=EDGE_INPUT)


def calc(ws, ref, formula, *, fmt=None, font=F_NUM, align=RIGHT, border=None):
    return put(ws, ref, formula, font=font, align=align, fmt=fmt, editable=False, border=border)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def hide(ws, *cols):
    for col in cols:
        ws.column_dimensions[col].hidden = True


# Options are written into a hidden column and the dropdown points at that
# range. An inline list cannot be used for all of them: "Per task, across all
# rounds" contains a comma, which is the inline list's own separator, so the
# option would silently split into two. A range has no such rule, survives the
# Google Sheets and LibreOffice conversions, and keeps the wording the module
# and its tests use.
LIST_COL = "T"


def validate(ws, options, cells, *, slot: list[int] = [0]):
    col = LIST_COL
    top = slot[0] + 1
    for i, option in enumerate(options):
        cell = ws[f"{col}{top + i}"]
        cell.value = option
        cell.protection = Protection(locked=True)
    slot[0] = top + len(options)   # next list starts below this one
    ws.column_dimensions[col].hidden = True
    dv = DataValidation(
        type="list",
        formula1=f"${col}${top}:${col}${top + len(options) - 1}",
        allow_blank=True,
        showDropDown=False,
    )
    ws.add_data_validation(dv)
    dv.add(cells)


def merge_text(ws, ref_from, ref_to, text, *, font=F_BODY, align=WRAP, height=None):
    ws.merge_cells(f"{ref_from}:{ref_to}")
    put(ws, ref_from, text, font=font, align=align)
    if height:
        ws.row_dimensions[int("".join(ch for ch in ref_from if ch.isdigit()))].height = height


def brand_lockup(mark_h: int = 44, name_h: int = 22, gap: int = 12):
    """The mark beside the lettering, composited into one image.

    Excel cannot embed WebP and the repository's brand files are WebP, so they
    are converted at build time rather than a PNG copy being committed and left
    to drift. Both keyed assets carry an alpha channel. Composited at twice the
    displayed size, and neither is enlarged past its native resolution.
    """
    from PIL import Image as PILImage

    def load(name: str, h: int):
        src = BRAND_DIR / f"{name}.webp"
        if not src.exists():
            raise SystemExit(f"missing brand asset {src.relative_to(REPO)}")
        im = PILImage.open(src)
        if im.mode != "RGBA":
            im = im.convert("RGBA")
        target = min(h * 2, im.height)
        return im.resize((max(1, round(im.width * target / im.height)), target), PILImage.LANCZOS)

    mark, name = load("lc-mark-keyed", mark_h), load("lc-name", name_h)
    pad = gap * 2
    w, h = mark.width + pad + name.width, max(mark.height, name.height)
    canvas = PILImage.new("RGBA", (w, h), (0, 0, 0, 0))
    canvas.paste(mark, (0, (h - mark.height) // 2), mark)
    canvas.paste(name, (mark.width + pad, (h - name.height) // 2), name)
    BUILD.mkdir(parents=True, exist_ok=True)
    out = BUILD / "lc-lockup@2x.png"
    canvas.save(out)
    display_h = max(mark_h, name_h)
    return out, round(w * display_h / h), display_h


def ifchain(pairs, default='""') -> str:
    """A nested IF chain, written where IFS would read better. See the header."""
    out = default
    for test, value in reversed(pairs):
        out = f"IF({test},{value},{out})"
    return out


def q(s: str) -> str:
    """A string literal inside a formula."""
    return '"' + s.replace('"', '""') + '"'


# ---------------------------------------------------------------------------
# The Check tab. Built once and used twice: blank, and filled as the Example.
# ---------------------------------------------------------------------------
#
# Blank is unknown, not zero. Excel has no null, so each total carries a hidden
# "blocked" count beside it: a present path with a missing figure adds 1, and a
# total with any blocker shows "" rather than a number that is quietly too low.
#
# Unbounded is not a big number. A present path with no cap sets a flag, and
# every worst-case figure reads Unbounded rather than inventing a ceiling.

ROW0 = 14                    # first path row
NPATHS = 7
ROW1 = ROW0 + NPATHS - 1     # last path row
# Helper columns, hidden. One per rule, so a reader who unhides them can follow
# the arithmetic a step at a time.
H_TBLOCK, H_TCONTRIB, H_WBLOCK, H_WCONTRIB, H_UNCAP = "N", "O", "P", "Q", "R"


def check_tab(ws: Worksheet, data: dict, *, filled: bool) -> dict:
    slot = [0]   # per-sheet cursor for the hidden option lists
    paths = data["paths"]
    ex = data["example"]
    ex_by_kind = {p["kind"]: p for p in ex["paths"]}

    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 26, "B": 11, "C": 10, "D": 13, "E": 14, "F": 13, "G": 15,
                "H": 13, "I": 12, "J": 24, "K": 18})
    hide(ws, H_TBLOCK, H_TCONTRIB, H_WBLOCK, H_WCONTRIB, H_UNCAP, LIST_COL)

    put(ws, "A1", "Check" if not filled else "Example", font=F_H1)
    ws.row_dimensions[1].height = 26
    merge_text(ws, "A2", "K2",
               data["exampleLabel"] if filled else TAGLINE,
               font=F_LEAD, height=18)

    # ---- A · the workflow -------------------------------------------------
    put(ws, "A4", "A · The workflow", font=F_H2)
    fields = data["workflowInputs"]
    keys = ["cleanRun", "alwaysOnChecks", "peakTasksPerMinute", "quotaPerMinute"]
    refs = {}
    for i, (f, key) in enumerate(zip(fields, keys)):
        r = 5 + i
        ws.merge_cells(f"A{r}:B{r}")
        put(ws, f"A{r}", f["label"], font=F_BODY, align=WRAP_MID)
        inp(ws, f"C{r}", ex[key] if filled else None, fmt=NUMFMT)
        ws.merge_cells(f"D{r}:K{r}")
        put(ws, f"D{r}", f["hint"], font=F_NOTE, align=WRAP_MID)
        ws.row_dimensions[r].height = 20
        refs[key] = f"C{r}"

    clean, always = refs["cleanRun"], refs["alwaysOnChecks"]
    peak, quota = refs["peakTasksPerMinute"], refs["quotaPerMinute"]

    # ---- B · the paths ----------------------------------------------------
    put(ws, "A11", "B · Paths that send work back", font=F_H2)
    merge_text(ws, "A12", "K12",
               "A round costs the same shape of thing whatever sent it back: the steps that run "
               "again, the check that re-reads the new attempt, and the extra context the redo "
               "carries. Zero is a real answer. Blank is not.",
               font=F_NOTE, height=26)

    heads = ["Path", "Present?", "Billed?", "Avg rounds", "Work repeated", "Review call",
             "Context growth", "Round cost", "Max rounds", "After the last round…", "Owner"]
    hdr = ROW0 - 1
    ws.row_dimensions[hdr].height = 30
    for i, h in enumerate(heads):
        put(ws, f"{get_column_letter(1 + i)}{hdr}", h, font=F_HEAD, fill=FILL_NOIR,
            align=Alignment(horizontal="left", vertical="bottom", wrap_text=True))

    for i, meta in enumerate(paths):
        r = ROW0 + i
        kind = meta["kind"]
        given = ex_by_kind.get(kind) if filled else None
        ws.row_dimensions[r].height = 22

        if meta["custom"]:
            inp(ws, f"A{r}", None, align=LEFT, font=F_BODY)
            ws[f"A{r}"].font = F_GHOST
            ws[f"A{r}"].value = meta["label"]
        else:
            put(ws, f"A{r}", meta["label"], font=F_BODY, align=WRAP_MID, border=HAIR)

        inp(ws, f"B{r}", ("Yes" if given["present"] else "No") if given else "No", align=CENTER, font=F_BODY)
        inp(ws, f"C{r}", ("Yes" if given["billed"] else "No") if given
            else ("Yes" if meta["billedByDefault"] else "No"), align=CENTER, font=F_BODY)
        inp(ws, f"D{r}", given["averageRounds"] if given else None, fmt="0.0#")
        inp(ws, f"E{r}", given["repeated"] if given else None, fmt=NUMFMT)
        inp(ws, f"F{r}", given["review"] if given else None, fmt=NUMFMT)
        inp(ws, f"G{r}", given["growth"] if given else None, fmt=NUMFMT)
        # Round cost: blank unless the path is present AND all three are set.
        calc(ws, f"H{r}",
             f'=IF(B{r}<>"Yes","",IF(OR(E{r}="",F{r}="",G{r}=""),"",E{r}+F{r}+G{r}))',
             fmt=NUMFMT, font=F_BODY_B, border=HAIR)
        inp(ws, f"I{r}", given["maxRounds"] if given else None, fmt="0")
        inp(ws, f"J{r}", (given["afterLast"] or None) if given else None, align=LEFT, font=F_BODY)
        inp(ws, f"K{r}", (given["owner"] or None) if given else None, align=LEFT, font=F_BODY)

        # Hidden working, one rule per column.
        calc(ws, f"{H_TBLOCK}{r}", f'=IF(B{r}="Yes",IF(OR(D{r}="",H{r}=""),1,0),0)')
        calc(ws, f"{H_TCONTRIB}{r}", f'=IF(AND(B{r}="Yes",{H_TBLOCK}{r}=0),D{r}*H{r},0)')
        calc(ws, f"{H_WBLOCK}{r}", f'=IF(B{r}="Yes",IF(OR(I{r}="",H{r}=""),1,0),0)')
        calc(ws, f"{H_WCONTRIB}{r}", f'=IF(AND(B{r}="Yes",{H_WBLOCK}{r}=0),I{r}*H{r},0)')
        calc(ws, f"{H_UNCAP}{r}", f'=IF(AND(B{r}="Yes",I{r}=""),1,0)')

    validate(ws, ["Yes", "No"], f"B{ROW0}:C{ROW1}", slot=slot)
    validate(ws, list(data["afterLast"]), f"J{ROW0}:J{ROW1}", slot=slot)
    # A path the reader says is not present is dimmed, not removed: the row
    # still says what it would have been.
    ws.conditional_formatting.add(
        f"A{ROW0}:K{ROW1}",
        FormulaRule(formula=[f'$B{ROW0}<>"Yes"'], fill=FILL_MIST,
                    font=Font(name=BODY, size=11, color=rgb("text_quiet"))))

    P = lambda c: f"{c}{ROW0}:{c}{ROW1}"   # noqa: E731  - a range on the path block

    # ---- C · budget -------------------------------------------------------
    b0 = ROW1 + 2
    put(ws, f"A{b0}", "C · Budget", font=F_H2)
    ws.merge_cells(f"A{b0 + 1}:B{b0 + 1}")
    put(ws, f"A{b0 + 1}", "Per-task token budget", font=F_BODY, align=WRAP_MID)
    inp(ws, f"C{b0 + 1}", ex["budget"] if filled else None, fmt=NUMFMT)
    budget = f"C{b0 + 1}"
    ws.merge_cells(f"A{b0 + 2}:B{b0 + 2}")
    put(ws, f"A{b0 + 2}", "The budget is enforced…", font=F_BODY, align=WRAP_MID)
    ws.merge_cells(f"C{b0 + 2}:D{b0 + 2}")
    inp(ws, f"C{b0 + 2}", (ex["enforcement"] or None) if filled else None, align=LEFT, font=F_BODY)
    enf = f"C{b0 + 2}"
    validate(ws, list(data["budgetEnforcement"]), enf, slot=slot)
    ws.merge_cells(f"A{b0 + 3}:B{b0 + 3}")
    put(ws, f"A{b0 + 3}", "Blended price per million tokens (optional)", font=F_BODY, align=WRAP_MID)
    inp(ws, f"C{b0 + 3}", ex.get("pricePerMillion") if filled else None, fmt=MONEYFMT)
    price = f"C{b0 + 3}"
    ws.merge_cells(f"E{b0 + 3}:K{b0 + 3}")
    put(ws, f"E{b0 + 3}", "Your own currency. The workbook shows no symbol.", font=F_NOTE, align=WRAP_MID)

    # ---- the two headline numbers ----------------------------------------
    h0 = b0 + 5
    put(ws, f"A{h0}", "The two numbers", font=F_H2)

    base_blank = f'OR({clean}="",{always}="")'
    unbounded = f"(SUM({P(H_UNCAP)})>0)"
    t_blocked = f'OR({base_blank},SUM({P(H_TBLOCK)})>0)'
    w_blocked = f'OR({base_blank},{unbounded},SUM({P(H_WBLOCK)})>0)'

    rows = [
        ("Typical cost per task", f'=IF({t_blocked},"",{clean}+{always}+SUM({P(H_TCONTRIB)}))', NUMFMT,
         "Unit economics: what every task costs on a normal day."),
        ("…as a multiple of a clean run", None, MULFMT, ""),
        ("Typical demand at peak", None, NUMFMT, ""),
        ("…as a share of the quota", None, PCTFMT, ""),
        ("Worst case per task", None, NUMFMT,
         "Availability: every loop at its cap."),
        ("…as a multiple of a clean run", None, MULFMT, ""),
        ("Worst case at peak", None, NUMFMT, ""),
        ("…as a share of the quota", None, PCTFMT, ""),
        ("Cost per task", None, MONEYFMT, "Only when a price is entered."),
        ("Cost per 1,000 tasks", None, MONEYFMT, ""),
    ]
    for i, (label, _f, fmt, note) in enumerate(rows):
        r = h0 + 1 + i
        ws.merge_cells(f"A{r}:B{r}")
        put(ws, f"A{r}", label, font=F_BODY_B if not label.startswith("…") else F_BODY,
            align=WRAP_MID)
        # The value spans C:D. Column C is 10 wide for the paths table above,
        # and 170,600 set at 16pt does not fit in it — Excel renders that as
        # ###, which is a number the reader cannot read.
        ws.merge_cells(f"C{r}:D{r}")
        ws.row_dimensions[r].height = 20
        if note:
            ws.merge_cells(f"E{r}:K{r}")
            put(ws, f"E{r}", note, font=F_NOTE, align=WRAP_MID)

    typ = f"C{h0 + 1}"
    typ_mul = f"C{h0 + 2}"
    typ_peak = f"C{h0 + 3}"
    typ_share = f"C{h0 + 4}"
    wst = f"C{h0 + 5}"
    wst_mul = f"C{h0 + 6}"
    wst_peak = f"C{h0 + 7}"
    wst_share = f"C{h0 + 8}"
    per_task = f"C{h0 + 9}"
    per_k = f"C{h0 + 10}"

    calc(ws, typ, f'=IF({t_blocked},"",{clean}+{always}+SUM({P(H_TCONTRIB)}))', fmt=NUMFMT, font=F_BIG)
    calc(ws, typ_mul, f'=IF(OR({typ}="",{clean}="",{clean}<=0),"",{typ}/{clean})', fmt=MULFMT)
    calc(ws, typ_peak, f'=IF(OR({typ}="",{peak}=""),"",{peak}*{typ})', fmt=NUMFMT)
    calc(ws, typ_share, f'=IF(OR({typ_peak}="",{quota}="",{quota}<=0),"",{typ_peak}/{quota})', fmt=PCTFMT)
    # Unbounded is a word. Text and number share the cell, so the format is
    # General-safe: the number format only bites when a number lands there.
    calc(ws, wst, f'=IF({unbounded},"Unbounded",IF({w_blocked},"",{clean}+{always}+SUM({P(H_WCONTRIB)})))',
         fmt=NUMFMT, font=F_BIG)
    calc(ws, wst_mul, f'=IF(OR({unbounded},{wst}="",{clean}="",{clean}<=0),"",{wst}/{clean})', fmt=MULFMT)
    calc(ws, wst_peak, f'=IF({unbounded},"Unbounded",IF(OR({wst}="",{peak}=""),"",{peak}*{wst}))', fmt=NUMFMT)
    calc(ws, wst_share, f'=IF({unbounded},"Unbounded",IF(OR({wst_peak}="",{quota}="",{quota}<=0),"",{wst_peak}/{quota}))', fmt=PCTFMT)
    calc(ws, per_task, f'=IF(OR({typ}="",{price}=""),"",{typ}*{price}/1000000)', fmt=MONEYFMT)
    calc(ws, per_k, f'=IF({per_task}="","",{per_task}*1000)', fmt=MONEYFMT)

    return {
        "clean": clean, "always": always, "peak": peak, "quota": quota,
        "budget": budget, "enf": enf, "price": price,
        "typ": typ, "typ_mul": typ_mul, "typ_peak": typ_peak, "typ_share": typ_share,
        "wst": wst, "wst_mul": wst_mul, "wst_peak": wst_peak, "wst_share": wst_share,
        "per_task": per_task, "per_k": per_k,
        "unbounded": unbounded, "P": P, "checks_top": h0 + 12,
    }


def checks_block(ws: Worksheet, ref: dict) -> dict:
    """The five checks, as formulas that mirror src/lib/sendBackCost.ts.

    Each one returns a status word and a reason. The status is a WORD first;
    the conditional formatting colours it and never replaces it, which is what
    keeps the sheet readable for somebody who cannot tell the fills apart.
    """
    P = ref["P"]
    top = ref["checks_top"]
    put(ws, f"A{top}", "The five checks", font=F_H2)

    # The columns here are the paths table's, and they are the wrong widths for
    # a checklist: B is 11 characters, which clips every check title. The block
    # spans instead — # in A, the title across B:E, the status in F, the reason
    # across G:K — so no width has to serve two jobs.
    hdr = top + 1
    ws.merge_cells(f"B{hdr}:E{hdr}")
    ws.merge_cells(f"G{hdr}:K{hdr}")
    for col, h in (("A", "#"), ("B", "Check"), ("F", "Status"), ("G", "Why")):
        put(ws, f"{col}{hdr}", h, font=F_HEAD, fill=FILL_NOIR, align=LEFT)
    for col in ("C", "D", "E", "H", "I", "J", "K"):
        put(ws, f"{col}{hdr}", None, font=F_HEAD, fill=FILL_NOIR, align=LEFT)
    ws.row_dimensions[hdr].height = 20

    present = f'COUNTIF({P("B")},"Yes")'
    unb = ref["unbounded"]

    # --- the counts each rule needs ---------------------------------------
    miss_avg = f'COUNTIFS({P("B")},"Yes",{P("D")},"")'
    miss_cost = (f'COUNTIFS({P("B")},"Yes",{P("E")},"")'
                 f'+COUNTIFS({P("B")},"Yes",{P("F")},"")'
                 f'+COUNTIFS({P("B")},"Yes",{P("G")},"")')
    # Judge, validation, tool and human are rows 2..5 of the block: those are
    # the paths whose redo normally carries the previous attempt forward.
    g0, g1 = ROW0 + 1, ROW0 + 4
    zero_growth = f'COUNTIFS(B{g0}:B{g1},"Yes",G{g0}:G{g1},0)'
    uncapped = f'SUM({P(H_UNCAP)})'
    undecided = (f'COUNTIFS({P("B")},"Yes",{P("J")},"")'
                 f'+COUNTIFS({P("B")},"Yes",{P("J")},"Not decided")')
    unowned = f'COUNTIFS({P("B")},"Yes",{P("K")},"")'

    NA = q("Not answered")
    PASS, ATT, FAIL = q("Pass"), q("Attention"), q("Fail")

    rows = [
        # (title, status formula, reason formula)
        (
            "Every path that sends work back is listed",
            ifchain([(f"{present}=0", NA), (f"{miss_avg}>0", FAIL)], PASS),
            ifchain([
                (f"{present}=0", q("No path is marked as present yet.")),
                (f"{miss_avg}>0", q("A present path has no average rounds. A path you cannot count is a path you cannot cost.")),
            ], q("Every present path says how often it sends work back.")),
        ),
        (
            "Each round is costed as repeated, review and growth",
            ifchain([(f"{present}=0", NA), (f"{miss_cost}>0", FAIL), (f"{zero_growth}>0", ATT)], PASS),
            ifchain([
                (f"{present}=0", q("No path is marked as present yet.")),
                (f"{miss_cost}>0", q("A present path is missing repeated, review or growth. Zero is a valid answer; blank is not.")),
                (f"{zero_growth}>0", q("A judge, validation, tool or human path has no context growth. Redos usually carry more context. Check this.")),
            ], q("Every present path has all three parts of a round.")),
        ),
        (
            "Rounds are capped, and something happens after the last one",
            ifchain([(f"{present}=0", NA), (f"OR({uncapped}>0,{undecided}>0)", FAIL), (f"{unowned}>0", ATT)], PASS),
            ifchain([
                (f"{present}=0", q("No path is marked as present yet.")),
                (f"{uncapped}>0", q("A present path has no cap. Without one the worst case has no ceiling.")),
                (f"{undecided}>0", q("A present path has no decided outcome after the last round.")),
                (f"{unowned}>0", q("A path is capped with an outcome, but nobody owns it.")),
            ], q("Every present path is capped, with a decided outcome and a named owner.")),
        ),
        (
            "Two numbers: unit economics and availability",
            ifchain([
                (unb, FAIL),
                (f'OR({ref["typ_share"]}="",{ref["wst_share"]}="")', NA),
                (f'{ref["typ_share"]}>1', FAIL),
                (f'AND({ref["typ_share"]}<=0.8,{ref["wst_share"]}<=1)', PASS),
            ], ATT),
            ifchain([
                (unb, q("The worst case is unbounded, so there is no availability number to check.")),
                (f'OR({ref["typ_share"]}="",{ref["wst_share"]}="")', q("Peak tasks per minute and the quota are needed for this one.")),
                (f'{ref["typ_share"]}>1', q("Your normal load already exceeds the quota.")),
                (f'AND({ref["typ_share"]}<=0.8,{ref["wst_share"]}<=1)', q("Normal load has headroom, and the worst case still fits.")),
                (f'{ref["typ_share"]}<=0.8', q("Normal load has headroom, but the worst case goes over the quota.")),
            ], q("Normal load is above 80% of the quota. Any spike tips it over.")),
        ),
        (
            "The budget belongs to the task, across all its rounds",
            ifchain([
                (f'{ref["enf"]}=""', NA),
                (f'{ref["enf"]}<>"Per task, across all rounds"', FAIL),
                (f'OR({ref["budget"]}="",{ref["typ"]}="")', NA),
                (f'{ref["budget"]}<{ref["typ"]}', FAIL),
                (f'OR({unb},AND({ref["wst"]}<>"",{ref["wst"]}<>"Unbounded",{ref["wst"]}>{ref["budget"]}))', PASS),
                (f'{ref["wst"]}=""', NA),
            ], ATT),
            ifchain([
                (f'{ref["enf"]}=""', q("Say how the budget is enforced.")),
                (f'{ref["enf"]}="Per call only"', q("A per-call budget never sees the rounds add up. The task is what costs money.")),
                (f'{ref["enf"]}="Not enforced"', q("Nothing stops a task at a token total.")),
                (f'OR({ref["budget"]}="",{ref["typ"]}="")', q("Enter a per-task budget to check it against the two totals.")),
                (f'{ref["budget"]}<{ref["typ"]}', q("The budget will cut normal tasks short.")),
                (f'OR({unb},AND({ref["wst"]}<>"",{ref["wst"]}<>"Unbounded",{ref["wst"]}>{ref["budget"]}))',
                 q("The budget stops the worst case. Make sure check 3's outcome runs when it does.")),
                (f'{ref["wst"]}=""', q("The worst case is not worked out yet.")),
            ], q("The budget never binds; your caps are the only stop.")),
        ),
    ]

    first = hdr + 1
    for i, (title, status, reason) in enumerate(rows):
        r = first + i
        ws.row_dimensions[r].height = 30
        put(ws, f"A{r}", i + 1, font=F_BODY, align=CENTER, border=HAIR)
        ws.merge_cells(f"B{r}:E{r}")
        put(ws, f"B{r}", title, font=F_BODY, align=WRAP_MID, border=HAIR)
        calc(ws, f"F{r}", "=" + status, font=F_BODY_B, align=CENTER, border=HAIR)
        ws.merge_cells(f"G{r}:K{r}")
        calc(ws, f"G{r}", "=" + reason, font=F_NOTE, align=WRAP_MID)

    last = first + len(rows) - 1
    for word, fill, colour in (("Pass", FILL_OK, "success"),
                               ("Attention", FILL_WARN, "warning"),
                               ("Fail", FILL_BAD, "danger")):
        ws.conditional_formatting.add(
            f"F{first}:F{last}",
            FormulaRule(formula=[f'$F{first}={q(word)}'], fill=fill,
                        font=Font(name=BODY, size=11, bold=True, color=rgb(colour))))

    # Your next step: the first Fail in check order, else the first Attention.
    nr = last + 2
    ws.merge_cells(f"A{nr}:B{nr}")
    put(ws, f"A{nr}", "Your next step", font=F_LABEL, align=LEFT)
    ws.merge_cells(f"C{nr}:K{nr}")
    ws.row_dimensions[nr].height = 26
    sts, rsn = f"F{first}:F{last}", f"G{first}:G{last}"
    nums = f"A{first}:A{last}"
    calc(ws, f"C{nr}",
         f'=IF(COUNTIF({sts},"Fail")>0,'
         f'"Check "&INDEX({nums},MATCH("Fail",{sts},0))&" — "&INDEX({rsn},MATCH("Fail",{sts},0)),'
         f'IF(COUNTIF({sts},"Attention")>0,'
         f'"Check "&INDEX({nums},MATCH("Attention",{sts},0))&" — "&INDEX({rsn},MATCH("Attention",{sts},0)),'
         f'"Nothing outstanding."))',
         font=F_BODY, align=WRAP_MID)

    return {"first": first, "last": last, "next": f"C{nr}", "bottom": nr}


# ---------------------------------------------------------------------------
# Start here
# ---------------------------------------------------------------------------


def cover(ws: Worksheet):
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 2.4, "B": 20, "C": 64, "D": 24})

    lockup, lw, lh = brand_lockup()
    img = XlImage(str(lockup))
    img.width, img.height = lw, lh
    img.anchor = "B2"
    ws.add_image(img)
    ws.row_dimensions[2].height = 40
    ws.row_dimensions[3].height = 14
    put(ws, "B4", "RESOURCES · AGENTIC SYSTEM DESIGN", font=F_LABEL)
    ws.row_dimensions[4].height = 14

    ws.merge_cells("B6:D6")
    put(ws, "B6", TITLE, font=F_TITLE)
    ws.row_dimensions[6].height = 38
    ws.merge_cells("B7:D7")
    put(ws, "B7", TAGLINE, font=F_LEAD, align=WRAP)
    ws.row_dimensions[7].height = 32
    ws.merge_cells("B8:D8")
    put(ws, "B8", "Sunil Mathew · The Living Craft", font=F_NOTE)

    blocks = [
        ("What to bring",
         "One agent workflow. Token counts from a handful of traces, or estimates. How often each "
         "check or limit sends work back. Your peak tasks per minute and your tokens-per-minute quota."),
        ("What to do",
         "List the paths that send work back, cost one round of each, set a cap and a final outcome "
         "for each, and add your per-task budget."),
        ("What you leave with",
         "Two numbers — typical cost per task and worst case at peak — plus a Pass, Attention or "
         "Fail on five checks, and the next fix to make."),
        ("How to use it",
         "Fill in the Check tab, top to bottom. The Example tab is the same sheet filled in. "
         "How to find these numbers says where each figure lives in your own system."),
    ]
    r = 10
    for label, text in blocks:
        put(ws, f"B{r}", label, font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
        ws.merge_cells(f"C{r}:D{r}")
        put(ws, f"C{r}", text, font=F_BODY, align=WRAP)
        ws.row_dimensions[r].height = 46 if len(text) > 120 else 30
        r += 2

    put(ws, f"B{r}", "The five checks", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    for i, t in enumerate([
        "Every path that sends work back is listed",
        "Each round is costed as repeated, review and growth",
        "Rounds are capped, and something happens after the last one",
        "Two numbers: unit economics and availability",
        "The budget belongs to the task, across all its rounds",
    ]):
        put(ws, f"C{r + i}", f"{i + 1}.  {t}", font=F_BODY)
        ws.row_dimensions[r + i].height = 20
    r += 6

    put(ws, f"B{r}", "Cell legend", font=F_LABEL, align=Alignment(horizontal="left", vertical="top"))
    put(ws, f"C{r}", LEGEND, font=F_BODY)
    ws.row_dimensions[r].height = 22
    r += 1
    inp(ws, f"C{r}", "Yours to fill", align=LEFT, font=F_BODY)
    put(ws, f"D{r}", "Shaded, with a thin border.", font=F_NOTE)
    ws.row_dimensions[r].height = 22
    r += 1
    put(ws, f"C{r}", "Calculates for you", font=F_NUM)
    put(ws, f"D{r}", "No fill, and locked.", font=F_NOTE)
    ws.row_dimensions[r].height = 22
    r += 2
    ws.merge_cells(f"B{r}:D{r}")
    put(ws, f"B{r}", FOOTER, font=F_NOTE)


# ---------------------------------------------------------------------------
# How to find these numbers
# ---------------------------------------------------------------------------


def guide_tab(ws: Worksheet, data: dict):
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 2.4, "B": 34, "C": 92})
    put(ws, "A1", "How to find these numbers", font=F_H1)
    ws.row_dimensions[1].height = 26
    ws.merge_cells("B2:C2")
    put(ws, "B2",
        "Where each figure lives in your own system. Where a product or a setting is named, "
        "check your provider's docs; names change.",
        font=F_LEAD, align=WRAP)
    ws.row_dimensions[2].height = 30

    r = 4
    for section in data["howToFind"]:
        put(ws, f"B{r}", section["title"], font=F_BODY_B, align=WRAP)
        ws.row_dimensions[r].height = 20
        for point in section["points"]:
            # The page renders these with inline markup; a cell cannot, so the
            # tags are stripped rather than shown as angle brackets.
            text = point.replace("<b>", "").replace("</b>", "")
            ws.merge_cells(f"C{r}:C{r}")
            put(ws, f"C{r}", "•  " + text, font=F_NOTE, align=WRAP)
            ws.row_dimensions[r] = ws.row_dimensions[r]
            ws.row_dimensions[r].height = max(16, 14 * (1 + len(text) // 90))
            r += 1
        r += 1
    ws.freeze_panes = "A4"


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------


def protect(ws: Worksheet):
    """A guide rail, not security: no password, and the reader can turn it off."""
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


def page_setup(ws: Worksheet, fit_height: int = 1):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.4


def main() -> int:
    data = content()

    wb = Workbook()
    wb.remove(wb.active)
    ws_cover = wb.create_sheet("Start here")
    ws_check = wb.create_sheet("Check")
    ws_guide = wb.create_sheet("How to find these numbers")
    ws_example = wb.create_sheet("Example")

    cover(ws_cover)
    ws_cover.sheet_properties.tabColor = rgb("sun")[2:]

    ref_blank = check_tab(ws_check, data, filled=False)
    chk_blank = checks_block(ws_check, ref_blank)
    ws_check.freeze_panes = "A4"

    guide_tab(ws_guide, data)

    ref_ex = check_tab(ws_example, data, filled=True)
    chk_ex = checks_block(ws_example, ref_ex)
    ws_example.freeze_panes = "A4"

    for ws in wb.worksheets:
        page_setup(ws, fit_height=0 if ws is ws_guide else 1)
        protect(ws)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(REPO)}  ({OUT.stat().st_size:,} bytes)")

    refs = {
        "sheets": {"cover": "Start here", "check": "Check",
                   "guide": "How to find these numbers", "example": "Example"},
        "example": {**{k: v for k, v in ref_ex.items() if isinstance(v, str)},
                    "checks_first": chk_ex["first"], "checks_last": chk_ex["last"],
                    "next": chk_ex["next"],
                    "judge_round": f"H{ROW0 + 1}", "transport_round": f"H{ROW0}",
                    "judge_cap": f"I{ROW0 + 1}"},
        "blank": {**{k: v for k, v in ref_blank.items() if isinstance(v, str)},
                  "checks_first": chk_blank["first"], "checks_last": chk_blank["last"],
                  "next": chk_blank["next"]},
        "expected": data["reading"],
    }
    (HERE / "refs.json").write_text(json.dumps(refs, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {(HERE / 'refs.json').relative_to(REPO)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
