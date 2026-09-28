#!/usr/bin/env python3
"""Recalculate the workbook and assert it agrees with the TypeScript module.

openpyxl writes formulas and no cached values, so anything read back here was
computed by a real spreadsheet engine rather than copied from a script.

THREE THINGS ARE CHECKED, AND THE THIRD IS THE POINT:
  1. the Example tab reproduces the figures in the brief, section 6;
  2. the blank Check tab shows blanks, never errors;
  3. the workbook and src/lib/sendBackCost.ts agree — the expected values come
     from refs.json, which build_xlsx.py wrote from the module's own output.

Two engines run where both are available and they have to agree: LibreOffice
headless, the closest thing on hand to what a reader opens the file in, and the
`formulas` package, which parses the formulas out of the file itself. One
engine alone can be wrong in a way that looks like a pass.

Run:  python3 tools/send-back-cost-check/verify_xlsx.py
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
XLSX = REPO / "public" / "downloads" / "send-back-cost-check.xlsx"
REFS = json.loads((HERE / "refs.json").read_text(encoding="utf-8"))

SOFFICE_CANDIDATES = [
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/opt/homebrew/bin/soffice",
    "/usr/local/bin/soffice",
    "/usr/bin/soffice",
]
ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!")


def soffice() -> str | None:
    for path in SOFFICE_CANDIDATES:
        if Path(path).exists():
            return path
    return shutil.which("soffice") or shutil.which("libreoffice")


def values_libreoffice(src: Path) -> dict[str, dict[str, object]]:
    out = Path(tempfile.mkdtemp(prefix="sbc-lo-"))
    subprocess.run(
        [soffice(), "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
        check=True, capture_output=True, timeout=300,
    )
    got = out / (src.stem + ".xlsx")
    if not got.exists():
        raise SystemExit(f"LibreOffice produced nothing for {src}")
    wb = openpyxl.load_workbook(got, data_only=True)
    return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row}
            for ws in wb.worksheets}


def values_formulas(src: Path) -> dict[str, dict[str, object]]:
    import formulas  # slow to import and only needed here

    model = formulas.ExcelModel().loads(str(src)).finish()
    out: dict[str, dict[str, object]] = {}
    for key, cell in model.calculate().items():
        if "]" not in key or "!" not in key:
            continue
        sheet, _, coord = key.partition("!")
        sheet = sheet.split("]", 1)[1].rstrip("'").upper()
        if not coord or ":" in coord:
            continue
        try:
            out.setdefault(sheet, {})[coord] = cell.value[0, 0]
        except Exception:
            continue
    return out


def norm(v):
    if v is None:
        return ""
    if hasattr(v, "item"):
        try:
            v = v.item()
        except Exception:
            pass
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return round(float(v), 6)
    return v


def fmt(v) -> str:
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = f"{v:,}" if isinstance(v, int) else str(v)
    return s if len(s) <= 22 else s[:21] + "…"


class Report:
    def __init__(self, title: str) -> None:
        self.title, self.rows, self.ok = title, [], True

    def check(self, label, expected, a, b):
        e, x, y = norm(expected), norm(a), norm(b)
        passed = (x == e) and (x == y)
        self.rows.append((label, fmt(e), fmt(x), fmt(y), passed))
        if not passed:
            self.ok = False

    def print(self):
        print(f"\n{self.title}")
        print(f"  {'':2} {'check':<44} {'expected':>22}  {'engine A':>22}  {'engine B':>22}")
        for label, e, x, y, passed in self.rows:
            print(f"  {'ok' if passed else 'XX'} {label:<44} {e:>22}  {x:>22}  {y:>22}")


def error_like(value) -> bool:
    if not isinstance(value, str):
        return False
    v = value.strip()
    if v in ERRORS:
        return True
    return len(v) > 1 and v.startswith("#") and (v.endswith(("!", "?")) or v == "#N/A")


def main() -> int:
    if not XLSX.exists():
        raise SystemExit(f"{XLSX} is missing. Run build_xlsx.py first.")

    have_lo = soffice() is not None
    if not have_lo:
        print("NOTE  LibreOffice is not installed, so only the `formulas` engine ran.")
        print("      Install it and run this again before the file is published:")
        print("        brew install --cask libreoffice")

    print("recalculating the shipped file …")
    fo = values_formulas(XLSX)
    lo = values_libreoffice(XLSX) if have_lo else fo

    E, B, X = REFS["example"], REFS["blank"], REFS["expected"]
    ex_sheet, blank_sheet = REFS["sheets"]["example"], REFS["sheets"]["check"]

    def get(sheets, sheet, ref):
        for key in (sheet, sheet.upper()):
            if key in sheets:
                return sheets[key].get(ref)
        return None

    def pair(sheet, ref):
        return get(lo, sheet, ref), get(fo, sheet, ref)

    # --- 1 + 3. the example, against the module's own output ---------------
    r = Report("Example tab — the brief's section 6, and the TypeScript module's output")
    cases = [
        ("Judge round cost", 21_900, E["judge_round"]),
        ("Transport round cost", 16_700, E["transport_round"]),
        ("Typical cost per task", X["typical"], E["typ"]),
        ("…as a multiple of a clean run", round(X["typicalMultiple"], 6), E["typ_mul"]),
        ("Typical demand at peak", X["typicalPeak"], E["typ_peak"]),
        ("…as a share of the quota", round(X["typicalShare"], 6), E["typ_share"]),
        ("Worst case per task", X["worst"], E["wst"]),
        ("…as a multiple of a clean run", round(X["worstMultiple"], 6), E["wst_mul"]),
        ("Worst case at peak", X["worstPeak"], E["wst_peak"]),
        ("…as a share of the quota", round(X["worstShare"], 6), E["wst_share"]),
    ]
    for label, expected, ref in cases:
        r.check(label, expected, *pair(ex_sheet, ref))
    for c in X["checks"]:
        r.check(f"Check {c['n']} status", c["status"], *pair(ex_sheet, f"F{E['checks_first'] + c['n'] - 1}"))
    # The next-step cell reads the reason cell, which now carries the action on
    # a second line, so the expected string is composed the same way.
    nxt = X["nextStep"]
    expected_next = f"Check {nxt['n']} — {nxt['reason']}"
    if nxt.get("action"):
        expected_next += f"\nDo this: {nxt['action']}"
    r.check("Your next step names check 4", expected_next, *pair(ex_sheet, E["next"]))
    # And the action itself reached the sheet.
    r.check("Check 4 carries its action", True,
            *[bool(v and "Cut what a round costs" in str(v))
              for v in pair(ex_sheet, f"G{E['checks_first'] + 3}")])
    r.print()

    # --- the money line ----------------------------------------------------
    m = Report("Example tab with a price of 3 per million")
    tmp = Path(tempfile.mkdtemp(prefix="sbc-price-")) / "price.xlsx"
    shutil.copy(XLSX, tmp)
    wb = openpyxl.load_workbook(tmp)
    wb[ex_sheet][E["price"]] = 3
    wb.save(tmp)
    lo2 = values_libreoffice(tmp) if have_lo else values_formulas(tmp)
    fo2 = values_formulas(tmp)
    m.check("Cost per task", 0.2958, get(lo2, ex_sheet, E["per_task"]), get(fo2, ex_sheet, E["per_task"]))
    m.check("Cost per 1,000 tasks", 295.8, get(lo2, ex_sheet, E["per_k"]), get(fo2, ex_sheet, E["per_k"]))
    m.print()

    # --- unbounded ---------------------------------------------------------
    u = Report("Example tab with the judge cap cleared")
    tmp2 = Path(tempfile.mkdtemp(prefix="sbc-unb-")) / "unbounded.xlsx"
    shutil.copy(XLSX, tmp2)
    wb = openpyxl.load_workbook(tmp2)
    wb[ex_sheet][E["judge_cap"]] = None
    wb.save(tmp2)
    lo3 = values_libreoffice(tmp2) if have_lo else values_formulas(tmp2)
    fo3 = values_formulas(tmp2)
    u.check("Worst case per task", "Unbounded", get(lo3, ex_sheet, E["wst"]), get(fo3, ex_sheet, E["wst"]))
    u.check("Worst case at peak", "Unbounded", get(lo3, ex_sheet, E["wst_peak"]), get(fo3, ex_sheet, E["wst_peak"]))
    u.check("Typical cost is unaffected", X["typical"], get(lo3, ex_sheet, E["typ"]), get(fo3, ex_sheet, E["typ"]))
    u.check("Check 3", "Fail", get(lo3, ex_sheet, f"F{E['checks_first'] + 2}"), get(fo3, ex_sheet, f"F{E['checks_first'] + 2}"))
    u.check("Check 4", "Fail", get(lo3, ex_sheet, f"F{E['checks_first'] + 3}"), get(fo3, ex_sheet, f"F{E['checks_first'] + 3}"))
    u.print()

    # --- 2. the blank Check tab -------------------------------------------
    b = Report("Blank Check tab — blanks and 'Not answered', never an error")
    for label, ref in [("Typical cost per task", B["typ"]), ("Worst case per task", B["wst"]),
                       ("Typical at peak", B["typ_peak"]), ("Cost per task", B["per_task"])]:
        b.check(label, "", *pair(blank_sheet, ref))
    for n in (1, 2, 3):
        b.check(f"Check {n} status", "Not answered", *pair(blank_sheet, f"F{B['checks_first'] + n - 1}"))
    b.print()

    # --- no error values anywhere -----------------------------------------
    problems = sorted({
        f"{eng}: '{name}'!{coord} = {v.strip()}"
        for eng, sheets in (("formulas", fo), *((("LibreOffice", lo),) if have_lo else ()))
        for name, cells in sheets.items()
        for coord, v in cells.items()
        if error_like(v)
    })
    print("\nError scan — #REF! #DIV/0! #VALUE! #NAME? and friends, across every cell")
    if problems:
        for p in problems:
            print(f"  XX {p}")
    else:
        print("  ok no error values in any cell of any tab, in either engine")

    passed = r.ok and m.ok and u.ok and b.ok and not problems
    engines = "2 engines (LibreOffice + formulas)" if have_lo else "1 engine (formulas only)"
    total = len(r.rows) + len(m.rows) + len(u.rows) + len(b.rows)
    print(f"\nEngines: {engines}")
    print(f"{'PASS' if passed else 'FAIL'} — {total} assertions, "
          f"{sum(len(c) for c in lo.values()):,} cells scanned")
    return 0 if passed else 1


if __name__ == "__main__":
    sys.exit(main())
