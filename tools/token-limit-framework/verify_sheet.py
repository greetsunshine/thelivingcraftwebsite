#!/usr/bin/env python3
"""Recalculate the workbook and assert every acceptance figure.

openpyxl writes formulas and no cached values, so anything read back here was
computed by a real spreadsheet engine rather than copied from this script.

Two engines run, and they have to agree:
  * LibreOffice headless, which is the closest thing on hand to what a reader
    opens the .xlsx in;
  * the `formulas` package, which parses the formulas out of the file itself.

One engine alone can be wrong in a way that looks like a pass. Disagreement is
a failure even when the number is the one expected.

Run:  python3 tools/token-limit-framework/verify_sheet.py
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
XLSX = REPO / "public" / "downloads" / "token-limit-framework.xlsx"
REFS = json.loads((HERE / "refs.json").read_text(encoding="utf-8"))
SOFFICE_CANDIDATES = [
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/opt/homebrew/bin/soffice",
    "/usr/local/bin/soffice",
    "/usr/bin/soffice",
]


def soffice() -> str | None:
    for path in SOFFICE_CANDIDATES:
        if Path(path).exists():
            return path
    return shutil.which("soffice") or shutil.which("libreoffice")

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!")


# ------------------------------------------------------------------ engines


def recalc_libreoffice(src: Path) -> Path:
    out = Path(tempfile.mkdtemp(prefix="tlf-lo-"))
    subprocess.run(
        [soffice(), "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
        check=True, capture_output=True, timeout=300,
    )
    got = out / (src.stem + ".xlsx")
    if not got.exists():
        raise SystemExit(f"LibreOffice produced nothing for {src}")
    return got


def values_libreoffice(src: Path) -> dict[str, dict[str, object]]:
    wb = openpyxl.load_workbook(recalc_libreoffice(src), data_only=True)
    return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row}
            for ws in wb.worksheets}


def values_formulas(src: Path) -> dict[str, dict[str, object]]:
    import formulas  # imported late: it is slow to load and only needed here

    model = formulas.ExcelModel().loads(str(src)).finish()
    out: dict[str, dict[str, object]] = {}
    for key, cell in model.calculate().items():
        # keys look like "'[FILE.XLSX]SHEET NAME'!D12"
        if "]" not in key or "!" not in key:
            continue
        sheet, _, coord = key.partition("!")
        sheet = sheet.split("]", 1)[1].rstrip("'").upper()
        if not coord or ":" in coord:
            continue
        try:
            value = cell.value[0, 0]
        except Exception:
            continue
        out.setdefault(sheet, {})[coord] = value
    return out


# ------------------------------------------------------------------ helpers


def norm(v):
    if v is None:
        return ""
    if hasattr(v, "item"):
        try:
            v = v.item()
        except Exception:
            pass
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return round(float(v), 6)
    return v


class Report:
    def __init__(self) -> None:
        self.rows: list[tuple[str, str, str, str, bool]] = []
        self.ok = True

    def check(self, label: str, expected, lo, fo) -> None:
        e, l, f = norm(expected), norm(lo), norm(fo)
        agree = l == f
        passed = (l == e) and agree
        self.rows.append((label, fmt(e), fmt(l), fmt(f), passed))
        if not passed:
            self.ok = False

    def print(self, title: str) -> None:
        print(f"\n{title}")
        print(f"  {'':2} {'check':<44} {'expected':>16}  {'engine A':>16}  {'engine B':>16}")
        for label, e, l, f, passed in self.rows:
            print(f"  {'ok' if passed else 'XX'} {label:<44} {e:>16}  {l:>16}  {f:>16}")


def fmt(v) -> str:
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = f"{v:,}" if isinstance(v, int) else str(v)
    return s if len(s) <= 16 else s[:15] + "…"


def scan_errors(sheets: dict[str, dict[str, object]], where: str) -> list[str]:
    bad = []
    for name, cells in sheets.items():
        for coord, value in cells.items():
            if error_like(value):
                bad.append(f"{where}: '{name}'!{coord} = {value.strip()}")
    return bad


def error_like(value) -> bool:
    """Catches the listed error values, and anything else shaped like one.

    A bare "#" is a column header on the decision tab, not an error, so the
    test is a "#" followed by something and ending in ! or ?, plus #N/A.
    """
    if not isinstance(value, str):
        return False
    v = value.strip()
    if v in ERRORS:
        return True
    return len(v) > 1 and v.startswith("#") and (v.endswith(("!", "?")) or v == "#N/A")


def blank_or_zero(sheets: dict[str, dict[str, object]], where: str) -> list[str]:
    """A formula on the blank template may show a blank or a 0, never an error."""
    bad = []
    for name, cells in sheets.items():
        if name.upper() == REFS["sheets"]["example"].upper():
            continue
        for coord, value in cells.items():
            if error_like(value):
                bad.append(f"{where}: '{name}'!{coord} = {value.strip()}")
    return bad


# -------------------------------------------------------------------- suite


def main() -> int:
    if not XLSX.exists():
        raise SystemExit(f"{XLSX} is missing. Run build_sheet.py first.")
    have_lo = soffice() is not None
    if not have_lo:
        print("NOTE  LibreOffice is not installed, so only the `formulas` engine ran.")
        print("      Every figure below was still recalculated from the file's own")
        print("      formulas — but the cross-engine agreement check did not happen.")
        print("      Install it and run this again before the file is published:")
        print("        brew install --cask libreoffice")

    ex_sheet = REFS["sheets"]["example"]
    E = REFS["example"]
    B = REFS["blank"]

    print("recalculating the shipped file …")
    fo = values_formulas(XLSX)
    lo = values_libreoffice(XLSX) if have_lo else fo

    def get(sheets, sheet, ref):
        # Engine B upper-cases sheet names; engine A keeps them as written.
        for key in (sheet, sheet.upper()):
            if key in sheets:
                return sheets[key].get(ref)
        return None

    def pair(sheet, ref):
        return get(lo, sheet, ref), get(fo, sheet, ref)

    # --- the worked example: section 3, tab 6 -------------------------------
    r = Report()
    cases = [
        ("Tokens per clean run", 43300, E["clean_run"]),
        ("Unsafe steps (side effect, not idempotent)", 2, E["unsafe"]),
        ("Tokens repeated per failed attempt", 16700, E["repeated"]),
        ("Worst case, eventually succeeds", 93400, E["succeeds"]),
        ("Worst case, gives up", 66800, E["gives_up"]),
        ("Amplification factor", 2.16, E["amp"]),
        ("Peak demand, normal", 8660000, E["peak_norm"]),
        ("Quota used, normal", 0.866, E["quota_norm"]),
        ("Band, normal", "Tight: any spike tips it over", E["band_norm"]),
        ("Peak demand, worst case", 18680000, E["peak_worst"]),
        ("Quota used, worst case", 1.868, E["quota_worst"]),
        ("Verdict", "Retries alone can exhaust the quota", E["verdict"]),
        ("Stop states with no clean-up plan", 0, E["missing_stop"]),
        ("Decision total", 16, E["total"]),
        ("Critical checks answered No", 1, E["critical"]),
        ("Decision", "Hold", E["decision"]),
    ]
    for label, expected, ref in cases:
        r.check(label, expected, *pair(ex_sheet, ref))
    r.print("Worked example — the acceptance cases in section 3, tab 6")

    # --- the checkpoint variant --------------------------------------------
    # Same inputs, but retries resume from a checkpoint and step 3 saves one.
    # Nothing before the refused step is repeated, so the amplifier reads 1.00×.
    tmp = Path(tempfile.mkdtemp(prefix="tlf-variant-")) / "variant.xlsx"
    shutil.copy(XLSX, tmp)
    wbv = openpyxl.load_workbook(tmp)
    wsv = wbv[ex_sheet]
    wsv[E["scope"]] = "Resume from last checkpoint"
    wsv[f'{E["cp_col"]}{E["cp_row_step3"]}'] = "Yes"
    wbv.save(tmp)
    fo2 = values_formulas(tmp)
    lo2 = values_libreoffice(tmp) if have_lo else fo2

    def pair2(ref):
        # `or` would be wrong here: a legitimate 0 is falsy, and the fallback
        # would turn it into None. This cost a false failure once already.
        return get(lo2, ex_sheet, ref), get(fo2, ex_sheet, ref)

    v = Report()
    v.check("Tokens repeated per failed attempt", 0, *pair2(E["repeated"]))
    v.check("Worst case, eventually succeeds", 43300, *pair2(E["succeeds"]))
    v.check("Amplification factor", 1.0, *pair2(E["amp"]))
    v.print("Variant — retry scope 'Resume from last checkpoint', checkpoint after step 3")

    # --- the blank template -------------------------------------------------
    b = Report()
    b.check("Map · tokens per clean run", 0, *pair(REFS["sheets"]["map"], B["map_clean_run"]))
    b.check("Retry · tokens repeated", "", *pair(REFS["sheets"]["retry"], B["retry_repeated"]))
    b.check("Retry · verdict", "", *pair(REFS["sheets"]["retry"], B["retry_verdict"]))
    b.check("Stop · states with no clean-up plan", 0, *pair(REFS["sheets"]["stop"], B["stop_missing"]))
    b.check("Decision · total", 0, *pair(REFS["sheets"]["decision"], B["total"]))
    b.check("Decision", "Not scored yet", *pair(REFS["sheets"]["decision"], B["decision"]))
    b.print("Blank template — a formula shows a blank or a 0, never an error")

    # --- no error values anywhere -------------------------------------------
    problems = scan_errors(fo, "formulas") + blank_or_zero(fo, "formulas/blank")
    if have_lo:
        problems += scan_errors(lo, "LibreOffice") + blank_or_zero(lo, "LibreOffice/blank")
    problems = sorted(set(problems))
    print("\nError scan — #REF! #DIV/0! #VALUE! #NAME? and friends, across every cell")
    if problems:
        for p in problems:
            print(f"  XX {p}")
    else:
        print("  ok no error values in any cell of any tab, in either engine")

    passed = r.ok and v.ok and b.ok and not problems
    engines = "2 engines (LibreOffice + formulas)" if have_lo else "1 engine (formulas only)"
    print(f"\nEngines: {engines}")
    print(f"{'PASS' if passed else 'FAIL'} — "
          f"{len(cases)} acceptance cases, 3 variant cases, 6 blank-template cases, "
          f"{sum(len(c) for c in lo.values()):,} cells scanned")
    return 0 if passed else 1


if __name__ == "__main__":
    sys.exit(main())
